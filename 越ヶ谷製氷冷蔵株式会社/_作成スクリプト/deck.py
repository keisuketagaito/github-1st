# -*- coding: utf-8 -*-
"""概要書・評価書・ノンネームシート共通の描画ヘルパー（M&A承継機構テンプレート準拠）"""
import copy
import json
import os
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

NAVY = RGBColor(0x0B, 0x30, 0x41)
BLUE = RGBColor(0x08, 0x89, 0xC9)
LBLUE = RGBColor(0xE1, 0xF3, 0xFB)
GRAY = RGBColor(0xE3, 0xE7, 0xE9)
LINE = RGBColor(0xA6, 0xA6, 0xA6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0, 0, 0)
DGRAY = RGBColor(0x40, 0x40, 0x40)
RED = RGBColor(0xC0, 0x00, 0x00)
JP, NUMF, MED = '游ゴシック', 'Arial', '游ゴシック Medium'
BLANK = '　'


# ------------------------------------------------------------------ master
class Master:
    def __init__(self, path, ref_path=None):
        self.wb = load_workbook(path, data_only=True)
        ref_path = ref_path or os.environ.get('MASTER_REF') or path + '.ref.json'
        self.ref = json.load(open(ref_path))

    def v(self, sheet, cell):
        return self.wb[sheet][cell].value

    def row(self, sheet, key):
        return self.ref[sheet][key]

    def rv(self, sheet, key, col):
        return self.wb[sheet][f'{col}{self.ref[sheet][key]}'].value


def rk(v):
    """円→千円（四捨五入、Excelの #,##0, 表示と同じ）"""
    if v is None:
        return None
    return int(Decimal(v / 1000).quantize(Decimal('1'), rounding=ROUND_HALF_UP)) if v >= 0 else -int(Decimal(-v / 1000).quantize(Decimal('1'), rounding=ROUND_HALF_UP))


def fk(v, dash='-'):
    """千円表示文字列（△表記）"""
    if v is None:
        return dash
    k = rk(v)
    if k == 0:
        return dash
    return f'△{abs(k):,}' if k < 0 else f'{k:,}'


def fy(v):
    return f'△{abs(v):,}' if v < 0 else f'{v:,}'


def fp(v, digits=1):
    if v is None:
        return '-'
    s = f'{abs(v) * 100:.{digits}f}%'
    return ('△' + s) if v < 0 else s


# ------------------------------------------------------------------ deck
class Deck:
    def __init__(self, template):
        self.prs = Presentation(template)
        self.src = list(self.prs.slides)
        self.layouts = {l.name: l for l in self.prs.slide_layouts}
        # 右上セクション名テキストボックスの原本（基本情報ページ）
        self.section_proto = None
        for sh in self.src[4].shapes:
            if sh.has_text_frame and sh.name == 'タイトル 1' and not sh.is_placeholder:
                self.section_proto = copy.deepcopy(sh._element)
        self.order = []

    # 構造 ---------------------------------------------------------------
    def use(self, idx):
        """テンプレートの既存スライド（1始まり）をそのまま使う"""
        s = self.src[idx - 1]
        self.order.append(s)
        return s

    def new(self, title, section, subtitle=None):
        s = self.prs.slides.add_slide(self.layouts['本編1'])
        for ph in s.placeholders:
            if ph.placeholder_format.idx == 0:
                ph.text_frame.text = title
            elif ph.placeholder_format.idx == 1:
                if subtitle:
                    ph.text_frame.text = subtitle
                else:
                    ph._element.getparent().remove(ph._element)
        el = copy.deepcopy(self.section_proto)
        s.shapes._spTree.append(el)
        for sh in s.shapes:
            if sh._element is el:
                set_runs_text(sh.text_frame, section)
        self.order.append(s)
        return s

    def divider(self, title):
        s = self.prs.slides.add_slide(self.layouts['中扉'])
        for ph in s.placeholders:
            if ph.placeholder_format.idx == 0:
                ph.text_frame.text = title
        self.order.append(s)
        return s

    def finalize(self, out):
        lst = self.prs.slides._sldIdLst
        ids = {s.slide_id: s for s in self.prs.slides}
        keep = [s.slide_id for s in self.order]
        for sid in list(lst):
            sid_val = int(sid.get('id'))
            if sid_val not in keep:
                self.prs.part.drop_rel(sid.get(qn('r:id')))
                lst.remove(sid)
        # 並べ替え
        elems = {int(e.get('id')): e for e in lst}
        for e in list(lst):
            lst.remove(e)
        for sid in keep:
            lst.append(elems[sid])
        cp = self.prs.core_properties
        cp.title = ''
        cp.subject = ''
        cp.keywords = ''
        cp.comments = ''
        cp.last_modified_by = ''
        cp.author = '株式会社M&A承継機構'
        self.prs.save(out)
        _scrub_app_xml(out)


def _scrub_app_xml(path):
    """docProps/app.xml のスライド見出し一覧（テンプレート案件の情報）を消去する"""
    import re
    import shutil
    import zipfile
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zi, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename == 'docProps/app.xml':
                t = data.decode('utf-8')
                t = re.sub(r'<HeadingPairs>.*?</HeadingPairs>', '', t, flags=re.S)
                t = re.sub(r'<TitlesOfParts>.*?</TitlesOfParts>', '', t, flags=re.S)
                t = re.sub(r'<Company>.*?</Company>', '<Company>株式会社M&amp;A承継機構</Company>', t, flags=re.S)
                data = t.encode('utf-8')
            zo.writestr(item, data)
    shutil.move(tmp, path)


def set_runs_text(tf, text):
    """書式を保ったまま最初の段落・ランのテキストを差し替える"""
    p0 = tf.paragraphs[0]
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    runs = p0.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p0.add_run().text = text


def clear_shapes(slide, keep_names=()):
    for sh in list(slide.shapes):
        if sh.is_placeholder or sh.name in keep_names:
            continue
        sh._element.getparent().remove(sh._element)


def set_subtitle(slide, text):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 1:
            set_runs_text(ph.text_frame, text)


def set_title(slide, text):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0:
            set_runs_text(ph.text_frame, text)


# ------------------------------------------------------------------ primitives
def _font(run, size, bold=False, color=BLACK, font=JP, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', JP if font == NUMF else font)


def text_box(slide, x, y, w, h, paras, size=10, color=BLACK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, font=JP, line_spacing=None, margin=0.02):
    """paras: str | list[str | list[(text, {opts})]]"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, Inches(margin))
    if isinstance(paras, str):
        paras = paras.split('\n')
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if line_spacing:
            para.line_spacing = line_spacing
        segs = p if isinstance(p, list) else [(p, {})]
        for t, o in segs:
            r = para.add_run()
            r.text = t
            _font(r, o.get('size', size), o.get('bold', bold), o.get('color', color), o.get('font', font))
        if isinstance(p, list) and p and 'space_before' in p[0][1]:
            para.space_before = Pt(p[0][1]['space_before'])
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE, text=None,
         size=10, color=BLACK, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, dash=False, font=JP):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
        if dash:
            from pptx.enum.dml import MSO_LINE
            sp.line.dash_style = MSO_LINE.DASH
    tf = sp.text_frame
    tf.word_wrap = True
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, Inches(0.04))
    tf.vertical_anchor = anchor
    if text is not None:
        lines = text if isinstance(text, list) else text.split('\n')
        for i, t in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            segs = t if isinstance(t, list) else [(t, {})]
            for tt, o in segs:
                r = p.add_run()
                r.text = tt
                _font(r, o.get('size', size), o.get('bold', bold), o.get('color', color), o.get('font', font))
    return sp


def line(slide, x1, y1, x2, y2, color=NAVY, width=1.0):
    from pptx.enum.shapes import MSO_CONNECTOR
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def arrow(slide, x, y, w, h, fill=BLUE, direction='right'):
    shp = {'right': MSO_SHAPE.RIGHT_ARROW, 'down': MSO_SHAPE.DOWN_ARROW, 'left': MSO_SHAPE.LEFT_ARROW,
           'chevron': MSO_SHAPE.CHEVRON, 'pentagon': MSO_SHAPE.PENTAGON, 'tri': MSO_SHAPE.ISOSCELES_TRIANGLE}[direction]
    s = rect(slide, x, y, w, h, fill=fill, shape=shp)
    if direction == 'tri':
        s.rotation = 90
    return s


def section_bar(slide, x, y, w, text, h=0.32, size=12):
    """紺色の見出しバー（会社概要などの帯）"""
    return rect(slide, x, y, w, h, fill=NAVY, text=text, size=size, color=WHITE, bold=True)


def sub_head(slide, x, y, w, text, size=11):
    """■ 付きの小見出し（左に青い縦棒）"""
    rect(slide, x, y + 0.03, 0.06, 0.2, fill=BLUE)
    return text_box(slide, x + 0.12, y, w - 0.12, 0.26, [[(text, {'bold': True, 'size': size, 'color': NAVY})]])


def unit(slide, x, y, w, text='（単位：千円）'):
    return text_box(slide, x, y, w, 0.2, text, size=8, color=DGRAY, align=PP_ALIGN.RIGHT)


def footnote(slide, lines, y=7.18, x=0.37, w=10.9, size=8):
    if isinstance(lines, str):
        lines = [lines]
    return text_box(slide, x, y, w, 0.2 * len(lines), lines, size=size, color=DGRAY, font=MED)


def comment_box(slide, x, y, w, h, sections, title='Comment', size=9, title_size=14):
    """sections: [(見出し, [本文行,...]), ...] 見出しは≪≫で囲う"""
    rect(slide, x, y, w, h, fill=WHITE, line=LINE, line_w=1.0)
    text_box(slide, x, y + 0.08, w, 0.32, [[(title, {'bold': True, 'size': title_size, 'color': NAVY, 'font': NUMF})]], align=PP_ALIGN.CENTER)
    paras = []
    for i, (head, body) in enumerate(sections):
        paras.append([(f'≪{head}≫', {'bold': True, 'size': size + 0.5, 'space_before': 6 if i else 0})])
        for b in body:
            paras.append([('・' + b, {'size': size})])
    text_box(slide, x + 0.14, y + 0.45, w - 0.28, h - 0.5, paras, size=size, line_spacing=1.05)


def photo(slide, key, x, y, w, h, caption, img_dir=None):
    """画像があれば貼付、なければ差替え用の枠を描く"""
    path = None
    if img_dir:
        for ext in ('.jpg', '.jpeg', '.png'):
            p = os.path.join(img_dir, key + ext)
            if os.path.exists(p):
                path = p
    if path:
        from PIL import Image
        im = Image.open(path)
        iw, ih = im.size
        pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
        # アスペクト比を保ったまま中央トリミング
        target = w / h
        cur = iw / ih
        if cur > target:
            crop = (1 - target / cur) / 2
            pic.crop_left = pic.crop_right = crop
        else:
            crop = (1 - cur / target) / 2
            pic.crop_top = pic.crop_bottom = crop
        text_box(slide, x, y + h + 0.02, w, 0.18, caption, size=7.5, color=DGRAY, align=PP_ALIGN.LEFT)
        return pic
    r = rect(slide, x, y, w, h, fill=RGBColor(0xF4, 0xF6, 0xF7), line=LINE, dash=True,
             text=[[('写真差替え枠', {'bold': True, 'size': 10, 'color': DGRAY})], [(caption, {'size': 7.5, 'color': DGRAY})],
                   [(f'画像/{key}.jpg を置いて再生成', {'size': 7, 'color': BLUE})]])
    return r


# ------------------------------------------------------------------ table
def table(slide, x, y, col_w, rows, row_h=0.22, size=9, header_rows=1, label_cols=1,
          num_cols=None, bold_rows=(), fill_rows=None, header_fill=NAVY, label_fill=GRAY,
          align_cols=None, wrap=True, hdr_size=None, italic_rows=(), small_rows=(), merge=None, col_fonts=None):
    """rows: list[list[str]]。num_cols: 右寄せ・Arialにする列番号の集合"""
    n_r, n_c = len(rows), len(col_w)
    gs = slide.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(sum(col_w)), Inches(row_h * n_r))
    tbl = gs.table
    # 既定スタイルを無効化
    tblPr = tbl._tbl.tblPr
    for a in ('firstRow', 'bandRow'):
        tblPr.set(a, '0')
    style = tblPr.find(qn('a:tableStyleId'))
    if style is not None:
        style.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'  # No Style, Table Grid
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Inches(w)
    num_cols = set(num_cols or [])
    fill_rows = fill_rows or {}
    for i in range(n_r):
        tbl.rows[i].height = Inches(row_h)
        for j in range(n_c):
            cell = tbl.cell(i, j)
            val = rows[i][j] if j < len(rows[i]) else ''
            if val is None:
                continue
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.01)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = wrap
            p = tf.paragraphs[0]
            lines_ = (val if val != '' else BLANK).split('\n')
            r = p.add_run()
            r.text = lines_[0]
            extra_runs = []
            for ln_ in lines_[1:]:
                p2 = tf.add_paragraph()
                r2 = p2.add_run()
                r2.text = ln_
                extra_runs.append((p2, r2))
            is_hdr = i < header_rows
            is_num = (j in num_cols) and not is_hdr
            fnt = NUMF if is_num else JP
            if col_fonts and j in col_fonts and not is_hdr:
                fnt = col_fonts[j]
            sz = (hdr_size or size) if is_hdr else (size - 1.5 if i in small_rows else size)
            col = WHITE if is_hdr else (DGRAY if i in small_rows else BLACK)
            _font(r, sz, bold=is_hdr or i in bold_rows, color=col, font=fnt, italic=i in italic_rows)
            if is_hdr:
                al = PP_ALIGN.CENTER
            elif align_cols and j in align_cols:
                al = align_cols[j]
            elif is_num:
                al = PP_ALIGN.RIGHT
            else:
                al = PP_ALIGN.LEFT
            p.alignment = al
            for p2, r2 in extra_runs:
                _font(r2, sz, bold=is_hdr or i in bold_rows, color=col, font=fnt)
                p2.alignment = al
            # 塗り
            if is_hdr:
                fc = header_fill
            elif i in fill_rows:
                fc = fill_rows[i]
            elif j < label_cols:
                fc = label_fill
            else:
                fc = WHITE
            cell.fill.solid()
            cell.fill.fore_color.rgb = fc
            _borders(cell)
    if merge:
        for (r1, c1, r2, c2) in merge:
            tbl.cell(r1, c1).merge(tbl.cell(r2, c2))
    return gs


def _borders(cell, color='A6A6A6', w=6350):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {'w': str(w), 'cap': 'flat', 'cmpd': 'sng', 'algn': 'ctr'})
        if tag in ('a:lnL', 'a:lnR'):
            ln.append(ln.makeelement(qn('a:noFill'), {}))
        else:
            sf = ln.makeelement(qn('a:solidFill'), {})
            c = sf.makeelement(qn('a:srgbClr'), {'val': color})
            sf.append(c)
            ln.append(sf)
        # tcPr 内の順序：lnL lnR lnT lnB ... solidFill
        tcPr.insert(['a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'].index(tag), ln)


# ------------------------------------------------------------------ charts
def bar_line_chart(slide, x, y, w, h, cats, bar_series, line_series=None, number_format='#,##0',
                   legend=True, bar_colors=(NAVY, BLUE), font_size=8, gap=80, labels=True):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in bar_series:
        cd.add_series(name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_title = False
    ch.font.size = Pt(font_size)
    ch.font.name = JP
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
    plot = ch.plots[0]
    plot.gap_width = gap
    for i, s in enumerate(plot.series):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = bar_colors[i % len(bar_colors)]
        if labels:
            s.data_labels.show_value = True
            s.data_labels.number_format = number_format
            s.data_labels.number_format_is_linked = False
            s.data_labels.font.size = Pt(font_size - 0.5)
            s.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = GRAY
    va.tick_labels.number_format = number_format
    va.tick_labels.number_format_is_linked = False
    va.format.line.fill.background()
    if min(min(v) for _, v in bar_series) >= 0:
        va.minimum_scale = 0
    ch.category_axis.format.line.color.rgb = LINE
    if line_series:
        _add_line_overlay(ch, cats, line_series)
    return gf


def _add_line_overlay(chart, cats, line_series):
    """棒グラフに第2軸の折れ線を重ねる（XML直接編集）"""
    from lxml import etree
    plotArea = chart._chartSpace.find('.//' + qn('c:plotArea'))
    bar = plotArea.find(qn('c:barChart'))
    ax_ids = [a.get('val') for a in bar.findall(qn('c:axId'))]
    n_exist = len(bar.findall(qn('c:ser')))
    C = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
    A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    lc = etree.SubElement(plotArea, qn('c:lineChart'))
    etree.SubElement(lc, qn('c:grouping')).set('val', 'standard')
    etree.SubElement(lc, qn('c:varyColors')).set('val', '0')
    for k, (name, vals, color, fmt) in enumerate(line_series):
        idx = n_exist + k
        ser = etree.SubElement(lc, qn('c:ser'))
        etree.SubElement(ser, qn('c:idx')).set('val', str(idx))
        etree.SubElement(ser, qn('c:order')).set('val', str(idx))
        tx = etree.SubElement(ser, qn('c:tx'))
        v = etree.SubElement(tx, qn('c:v'))
        v.text = name
        spPr = etree.SubElement(ser, qn('c:spPr'))
        ln = etree.SubElement(spPr, qn('a:ln'), w='28575')
        sf = etree.SubElement(ln, qn('a:solidFill'))
        etree.SubElement(sf, qn('a:srgbClr')).set('val', str(color))
        mk = etree.SubElement(ser, qn('c:marker'))
        etree.SubElement(mk, qn('c:symbol')).set('val', 'circle')
        etree.SubElement(mk, qn('c:size')).set('val', '6')
        mspPr = etree.SubElement(mk, qn('c:spPr'))
        msf = etree.SubElement(mspPr, qn('a:solidFill'))
        etree.SubElement(msf, qn('a:srgbClr')).set('val', str(color))
        dl = etree.SubElement(ser, qn('c:dLbls'))
        etree.SubElement(dl, qn('c:numFmt'), formatCode=fmt, sourceLinked='0')
        etree.SubElement(dl, qn('c:spPr')).append(etree.Element(qn('a:noFill')))
        txPr = etree.SubElement(dl, qn('c:txPr'))
        bp = etree.SubElement(txPr, qn('a:bodyPr'))
        etree.SubElement(txPr, qn('a:lstStyle'))
        p = etree.SubElement(txPr, qn('a:p'))
        pPr = etree.SubElement(p, qn('a:pPr'))
        dr = etree.SubElement(pPr, qn('a:defRPr'), sz='750', b='1')
        sfx = etree.SubElement(dr, qn('a:solidFill'))
        etree.SubElement(sfx, qn('a:srgbClr')).set('val', str(color))
        etree.SubElement(p, qn('a:endParaRPr'), lang='ja-JP')
        etree.SubElement(dl, qn('c:dLblPos')).set('val', 't')
        for tag, val in (('c:showLegendKey', '0'), ('c:showVal', '1'), ('c:showCatName', '0'), ('c:showSerName', '0'), ('c:showPercent', '0'), ('c:showBubbleSize', '0')):
            etree.SubElement(dl, qn(tag)).set('val', val)
        cat = etree.SubElement(ser, qn('c:cat'))
        sr = etree.SubElement(cat, qn('c:strLit'))
        etree.SubElement(sr, qn('c:ptCount')).set('val', str(len(cats)))
        for i, c in enumerate(cats):
            pt = etree.SubElement(sr, qn('c:pt'), idx=str(i))
            etree.SubElement(pt, qn('c:v')).text = str(c)
        val = etree.SubElement(ser, qn('c:val'))
        nl = etree.SubElement(val, qn('c:numLit'))
        etree.SubElement(nl, qn('c:formatCode')).text = 'General'
        etree.SubElement(nl, qn('c:ptCount')).set('val', str(len(vals)))
        for i, vv in enumerate(vals):
            pt = etree.SubElement(nl, qn('c:pt'), idx=str(i))
            etree.SubElement(pt, qn('c:v')).text = str(vv)
        etree.SubElement(ser, qn('c:smooth')).set('val', '0')
    etree.SubElement(lc, qn('c:marker')).set('val', '1')
    etree.SubElement(lc, qn('c:axId')).set('val', '90001')
    etree.SubElement(lc, qn('c:axId')).set('val', '90002')
    # 第2軸（カテゴリ軸は非表示、値軸は右）
    valAx0 = plotArea.find(qn('c:valAx'))
    catAx = etree.SubElement(plotArea, qn('c:catAx'))
    for tag, val in (('c:axId', '90001'), ):
        etree.SubElement(catAx, qn(tag)).set('val', val)
    sc = etree.SubElement(catAx, qn('c:scaling'))
    etree.SubElement(sc, qn('c:orientation')).set('val', 'minMax')
    etree.SubElement(catAx, qn('c:delete')).set('val', '1')
    etree.SubElement(catAx, qn('c:axPos')).set('val', 'b')
    etree.SubElement(catAx, qn('c:majorTickMark')).set('val', 'none')
    etree.SubElement(catAx, qn('c:minorTickMark')).set('val', 'none')
    etree.SubElement(catAx, qn('c:tickLblPos')).set('val', 'nextTo')
    etree.SubElement(catAx, qn('c:crossAx')).set('val', '90002')
    etree.SubElement(catAx, qn('c:crosses')).set('val', 'autoZero')
    etree.SubElement(catAx, qn('c:auto')).set('val', '1')
    etree.SubElement(catAx, qn('c:lblAlgn')).set('val', 'ctr')
    etree.SubElement(catAx, qn('c:lblOffset')).set('val', '100')
    etree.SubElement(catAx, qn('c:noMultiLvlLbl')).set('val', '0')
    valAx = etree.SubElement(plotArea, qn('c:valAx'))
    etree.SubElement(valAx, qn('c:axId')).set('val', '90002')
    sc = etree.SubElement(valAx, qn('c:scaling'))
    etree.SubElement(sc, qn('c:orientation')).set('val', 'minMax')
    etree.SubElement(valAx, qn('c:delete')).set('val', '0')
    etree.SubElement(valAx, qn('c:axPos')).set('val', 'r')
    etree.SubElement(valAx, qn('c:numFmt'), formatCode=line_series[0][3], sourceLinked='0')
    etree.SubElement(valAx, qn('c:majorTickMark')).set('val', 'none')
    etree.SubElement(valAx, qn('c:minorTickMark')).set('val', 'none')
    etree.SubElement(valAx, qn('c:tickLblPos')).set('val', 'nextTo')
    sp = etree.SubElement(valAx, qn('c:spPr'))
    ln = etree.SubElement(sp, qn('a:ln'))
    etree.SubElement(ln, qn('a:noFill'))
    etree.SubElement(valAx, qn('c:crossAx')).set('val', '90001')
    etree.SubElement(valAx, qn('c:crosses')).set('val', 'max')
    etree.SubElement(valAx, qn('c:crossBetween')).set('val', 'between')
    # plotArea 内で軸要素は chart 要素の後に並べる必要がある
    axes = [e for e in plotArea if e.tag in (qn('c:catAx'), qn('c:valAx'), qn('c:dateAx'), qn('c:serAx'))]
    others = [e for e in plotArea if e not in axes and e.tag not in (qn('c:spPr'), qn('c:extLst'))]
    tail = [e for e in plotArea if e.tag in (qn('c:spPr'), qn('c:extLst'))]
    for e in list(plotArea):
        plotArea.remove(e)
    for e in others + axes + tail:
        plotArea.append(e)


def doughnut(slide, x, y, w, h, cats, vals, colors, font_size=8):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series('構成比', vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.font.size = Pt(font_size)
    ch.font.name = JP
    pl = ch.plots[0]
    for i, pt in enumerate(pl.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = colors[i % len(colors)]
    return gf

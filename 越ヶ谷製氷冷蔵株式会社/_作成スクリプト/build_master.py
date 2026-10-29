# -*- coding: utf-8 -*-
"""案件マスター（越ヶ谷製氷冷蔵株式会社）を openpyxl で構築する。
数値は円単位で入力し、表示形式で千円表示する。計算は全て数式で持たせる。"""
import sys
sys.path.insert(0, '.')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from data import *

OUT = sys.argv[1]

NAVY, BLUE, LBLUE, GRAY, WHITE = '0B3041', '0889C9', 'E1F3FB', 'E3E7E9', 'FFFFFF'
FMT = '#,##0,;"△"#,##0,;"-"'          # 円入力→千円表示
FMT_Y = '#,##0;"△"#,##0;"-"'           # 円表示
FMT_P = '0.0%;"△"0.0%;"-"'
JP, NUM, MED = '游ゴシック', 'Arial', '游ゴシック Medium'
thin = Side(style='thin', color='A6A6A6')
BORDER = Border(bottom=thin)

wb = Workbook()
wb.remove(wb.active)
REF = {}   # REF[sheet][key] = row


def ws_new(name, title, widths=None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 2
    ws['B2'] = title
    ws['B2'].font = Font(name=JP, size=14, bold=True, color=NAVY)
    if widths:
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
    REF[name] = {}
    return ws


def sec(name):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = NAVY
    ws['B2'] = name.replace('→', '') + ' セクション'
    ws['B2'].font = Font(name=JP, size=14, bold=True, color=NAVY)
    return ws


def hdr(ws, row, labels, start_col=2, fill=NAVY):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=t)
        c.font = Font(name=JP, size=9, bold=True, color=WHITE)
        c.fill = PatternFill('solid', fgColor=fill)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def label(ws, row, text, bold=False, indent=0, col=2, fill=None, size=9):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(name=JP, size=size, bold=bold)
    c.alignment = Alignment(indent=indent, vertical='center', wrap_text=False)
    if fill:
        c.fill = PatternFill('solid', fgColor=fill)
    return c


def num(ws, row, col, v, bold=False, fmt=FMT, fill=None):
    c = ws.cell(row=row, column=col, value=v)
    c.font = Font(name=NUM, size=9, bold=bold)
    c.number_format = fmt
    c.border = BORDER
    if fill:
        c.fill = PatternFill('solid', fgColor=fill)
    return c


def note(ws, row, text, col=2, size=8):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(name=MED, size=size, color='404040')
    return c


def text(ws, row, col, v, bold=False, wrap=True, size=9, fill=None):
    c = ws.cell(row=row, column=col, value=v)
    c.font = Font(name=JP, size=size, bold=bold)
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    c.border = BORDER
    if fill:
        c.fill = PatternFill('solid', fgColor=fill)
    return c


PCOLS = [3, 4, 5, 6]          # C..F = 4期
P3 = [4, 5, 6]                 # 直近3期（評価で使用）= D..F
BSD, BSC = 'BS(借方)', 'BS (貸方)'


def q(sheet):
    return "'" + sheet + "'"

# =====================================================================
ws = wb.create_sheet('ルール')
ws.sheet_view.showGridLines = False
rules = ['案件マスター 作成ルール',
         '1. 文字は「游ゴシック」、数字は「Arial」。備考欄などの細かい文章は「游ゴシック Medium」',
         '2. タイトル・合計欄は太字',
         '3. 数字は円単位で入力する。表示は千円単位に設定済み（千円で入力しないこと）',
         '4. 背景は紺・水色・グレー・白・青。線は紺・白・青',
         '5. 数値の原本は本ファイルのみ。概要書・評価書・ノンネームシートは本ファイルから転記する',
         '6. 未確認事項は推測で埋めず「要確認」と記載し、「その他」シートに集約する',
         '',
         '【本ファイルの出所】',
         '・第64〜66期（2023年3月期〜2025年3月期）の決算数値：受領した既存の企業概要書（財務表）から転記',
         '・第67期（2026年3月期）の決算数値：第67期決算報告書・勘定科目内訳書・法人税申告書から転記',
         '・減価償却：第67期 減価償却計算書（合計表・明細）',
         '・登記情報：履歴事項（2026年9月9日現在）',
         '・転記後、4期すべてで資産合計・負債純資産合計・営業利益〜当期純利益・純資産の期間連続性が決算書記載値と一致することを確認済み']
for i, r in enumerate(rules):
    c = ws.cell(row=2 + i, column=2, value=r)
    c.font = Font(name=JP, size=12 if i == 0 else 10, bold=(i == 0 or r.startswith('【')))
ws.column_dimensions['B'].width = 120

# =====================================================================
sec('財務入力→')

# ---------- BS(借方) ----------
ws = ws_new(BSD, '貸借対照表（資産の部）　単位：千円（入力は円）', {'B': 30, 'C': 15, 'D': 15, 'E': 15, 'F': 15, 'G': 15, 'H': 40})
hdr(ws, 4, ['科目'] + PERIODS + ['進行期TB', '備考'])
for i, k in enumerate(KI):
    note(ws, 3, k, col=3 + i)
groups = [('cash', '（現金・預金）'), ('quick', '（当座資産）'), ('inv', '（棚卸資産）'), ('oca', '（流動資産諸口）'),
          ('ppe', '（有形固定資産）'), ('intan', '（無形固定資産）'), ('invest', '（投資その他の資産）')]
r = 5
R = REF[BSD]
cur_groups, fix_groups = ['cash', 'quick', 'inv', 'oca'], ['ppe', 'intan', 'invest']
remarks = {'諸掛金（生命保険積立）': '67期内訳書：大同生命8,461,965円・第一生命900,529円。現預金に含めて表示されているが実態は保険積立金',
           '定期預金': '67期：預貯金内訳書の定期預金計19,601,032円と差異6,800,000円（受取利息内訳では栃木銀行定期10,400,000円）→要確認',
           '売掛金': '67期内訳書：上位7社＋その他。合計が残高を1,723,680円超過（(株)小野田商店の重複計上と推察）→要確認',
           '建物': '67期は償却計上なし（償却不足額21,745,406円）。主たる工場建物は1991年取得',
           '機械': '67期償却不足額3,092,403円',
           '構築物': '67期は償却計上なし（償却不足額447,749円）',
           '貸付金': '従業員への生活資金貸付',
           '未収金': '67期：消費税還付等',
           '出資金': '越谷市工業団地組合800,000円ほか'}
for key, gname in groups:
    if key == 'ppe':
        label(ws, r, '流動資産合計', bold=True, fill=GRAY)
        for c in PCOLS:
            num(ws, r, c, '=' + '+'.join(f'{CL(c)}{R["sub_" + g]}' for g in cur_groups), bold=True, fill=GRAY)
        R['流動資産合計'] = r
        r += 1
    label(ws, r, gname, bold=True)
    sub_row = r
    R['sub_' + key] = r
    r += 1
    first = r
    for name, g, vals in BS_ASSETS:
        if g != key:
            continue
        label(ws, r, name, indent=1)
        for i, c in enumerate(PCOLS):
            num(ws, r, c, vals[i])
        if name in remarks:
            note(ws, r, remarks[name], col=8)
        R[name] = r
        r += 1
    for c in PCOLS:
        num(ws, sub_row, c, f'=SUM({CL(c)}{first}:{CL(c)}{r - 1})', bold=True)
label(ws, r, '固定資産合計', bold=True, fill=GRAY)
for c in PCOLS:
    num(ws, r, c, '=' + '+'.join(f'{CL(c)}{R["sub_" + g]}' for g in fix_groups), bold=True, fill=GRAY)
R['固定資産合計'] = r
r += 1
label(ws, r, '資産合計', bold=True, fill=LBLUE)
for c in PCOLS:
    num(ws, r, c, f'={CL(c)}{R["流動資産合計"]}+{CL(c)}{R["固定資産合計"]}', bold=True, fill=LBLUE)
R['資産合計'] = r
r += 1
label(ws, r, '（参考）決算書記載の資産合計')
for i, c in enumerate(PCOLS):
    num(ws, r, c, BS_ASSET_TOTAL[i])
R['資産合計_決算書'] = r
r += 2
# 借貸一致チェック（63行）
r = max(r, 63)
label(ws, r, '借貸一致チェック（資産合計＝負債純資産合計）', bold=True)
R['check'] = r

# ---------- BS (貸方) ----------
ws2 = ws_new(BSC, '貸借対照表（負債・純資産の部）　単位：千円（入力は円）', {'B': 30, 'C': 15, 'D': 15, 'E': 15, 'F': 15, 'G': 15, 'H': 40})
hdr(ws2, 4, ['科目'] + [f'={q(BSD)}!{CL(c)}4' for c in PCOLS] + ['進行期TB', '備考'])
R2 = REF[BSC]
r = 5
remarks2 = {'短期借入金': '67期：借入金内訳書（短期・長期合計461,024千円、うち代表者借入14,000千円）',
            '長期借入金': 'うち日本政策金融公庫の資本性ローン105,420千円（67期）',
            '未払金': '67期内訳書計15,715,659円とBS残高に5,398円差異。3月分給与5,720,681円・未払消費税1,676,200円を含む',
            '資本金': '2026年5月15日付で10,000千円に減資（登記）'}
for title, keys in [('（流動負債）', 'cl'), ('（固定負債）', 'fl')]:
    label(ws2, r, title, bold=True)
    sub = r
    R2['sub_' + keys] = r
    r += 1
    first = r
    for name, g, vals in BS_LIAB:
        if g != keys:
            continue
        label(ws2, r, name, indent=1)
        for i, c in enumerate(PCOLS):
            num(ws2, r, c, vals[i])
        if name in remarks2:
            note(ws2, r, remarks2[name], col=8)
        R2[name] = r
        r += 1
    for c in PCOLS:
        num(ws2, sub, c, f'=SUM({CL(c)}{first}:{CL(c)}{r - 1})', bold=True)
label(ws2, r, '負債合計', bold=True, fill=GRAY)
for c in PCOLS:
    num(ws2, r, c, f'={CL(c)}{R2["sub_cl"]}+{CL(c)}{R2["sub_fl"]}', bold=True, fill=GRAY)
R2['負債合計'] = r
r += 1
label(ws2, r, '（純資産の部）', bold=True)
r += 1
first = r
for name, g, vals in BS_EQUITY:
    label(ws2, r, name, indent=1)
    for i, c in enumerate(PCOLS):
        num(ws2, r, c, vals[i])
    if name in remarks2:
        note(ws2, r, remarks2[name], col=8)
    R2[name] = r
    r += 1
label(ws2, r, '純資産合計', bold=True, fill=GRAY)
for c in PCOLS:
    num(ws2, r, c, f'=SUM({CL(c)}{first}:{CL(c)}{r - 1})', bold=True, fill=GRAY)
R2['純資産合計'] = r
r += 1
label(ws2, r, '負債・純資産合計', bold=True, fill=LBLUE)
for c in PCOLS:
    num(ws2, r, c, f'={CL(c)}{R2["負債合計"]}+{CL(c)}{R2["純資産合計"]}', bold=True, fill=LBLUE)
R2['負債純資産合計'] = r
r += 2
label(ws2, r, '有利子負債合計（短期＋長期借入金）', bold=True)
for c in PCOLS:
    num(ws2, r, c, f'={CL(c)}{R2["短期借入金"]}+{CL(c)}{R2["長期借入金"]}', bold=True)
R2['有利子負債'] = r
# チェック行（借方シート）
for c in PCOLS:
    cc = ws.cell(row=R['check'], column=c, value=f'=EXACT({CL(c)}{R["資産合計"]},{q(BSC)}!{CL(c)}{R2["負債純資産合計"]})')
    cc.font = Font(name=NUM, size=9, bold=True)
label(ws, R['check'] + 1, '決算書記載値との一致チェック')
for c in PCOLS:
    ws.cell(row=R['check'] + 1, column=c, value=f'=EXACT({CL(c)}{R["資産合計"]},{CL(c)}{R["資産合計_決算書"]})').font = Font(name=NUM, size=9)

# ---------- CR ----------
ws3 = ws_new('CR', '製造原価報告書　単位：千円（入力は円）', {'B': 30, 'C': 15, 'D': 15, 'E': 15, 'F': 15, 'G': 40})
hdr(ws3, 4, ['科目'] + [f'={q(BSD)}!{CL(c)}4' for c in PCOLS] + ['備考'])
R3 = REF['CR']
r = 5
for gname, items in CR:
    label(ws3, r, '【' + gname + '】', bold=True)
    R3['g_' + gname] = r
    r += 1
    first = r
    for name, vals in items:
        label(ws3, r, name, indent=1)
        for i, c in enumerate(PCOLS):
            num(ws3, r, c, vals[i])
        R3[gname + '/' + name] = r
        r += 1
    last = r - 1
    for c in PCOLS:
        L = CL(c)
        if gname == '原材料費':
            f = f'={L}{first}+{L}{first + 1}-{L}{first + 2}'
        elif gname == '製造勘定':
            f = f'={L}{first}-{L}{first + 1}'
        else:
            f = f'=SUM({L}{first}:{L}{last})'
        num(ws3, R3['g_' + gname], c, f, bold=True)
label(ws3, r, '当期製品製造原価', bold=True, fill=LBLUE)
for c in PCOLS:
    num(ws3, r, c, '=' + '+'.join(f'{CL(c)}{R3["g_" + g]}' for g, _ in CR), bold=True, fill=LBLUE)
R3['製造原価'] = r
note(ws3, R3['労務費/退職給与'], '67期6,177,985円は過去3期に発生のない臨時的支出と推察（退職者・規程は要確認）', col=7)
note(ws3, R3['製造経費/減価償却費'], '64期はほぼ償却計上なし（19,763円）', col=7)
note(ws3, R3['製造経費/動力費'], '電力料。64期は電気料金高騰期', col=7)

# ---------- SGA ----------
ws4 = ws_new('SGA', '販売費及び一般管理費　単位：千円（入力は円）', {'B': 30, 'C': 15, 'D': 15, 'E': 15, 'F': 15, 'G': 40})
hdr(ws4, 4, ['科目'] + [f'={q(BSD)}!{CL(c)}4' for c in PCOLS] + ['備考'])
R4 = REF['SGA']
r = 5
for gname, items in SGA:
    label(ws4, r, '（' + gname + '）', bold=True)
    R4['g_' + gname] = r
    r += 1
    first = r
    for name, vals in items:
        label(ws4, r, name, indent=1)
        for i, c in enumerate(PCOLS):
            num(ws4, r, c, vals[i])
        R4[gname + '/' + name] = r
        r += 1
    for c in PCOLS:
        num(ws4, R4['g_' + gname], c, f'=SUM({CL(c)}{first}:{CL(c)}{r - 1})', bold=True)
label(ws4, r, '販売費及び一般管理費 合計', bold=True, fill=LBLUE)
for c in PCOLS:
    num(ws4, r, c, '=' + '+'.join(f'{CL(c)}{R4["g_" + g]}' for g, _ in SGA), bold=True, fill=LBLUE)
R4['合計'] = r
note(ws4, R4['資産管理費/地代家賃'], '駐車場（代表者から賃借）545,460円＋松戸営業所店舗1,200,000円', col=7)
note(ws4, R4['給与関係費/役員報酬'], '代表取締役 小野田洋 月額600千円', col=7)
note(ws4, R4['輸送関係費/運賃'], '配送委託（倉橋運送・鴻池運輸等）', col=7)

# ---------- PL ----------
ws5 = ws_new('PL', '損益計算書　単位：千円（入力は円）　※2列1組（帳簿価額／売上高比）',
             {'B': 30, 'C': 14, 'D': 8, 'E': 14, 'F': 8, 'G': 14, 'H': 8, 'I': 14, 'J': 8, 'K': 40})
VC = [3, 5, 7, 9]        # 値の列
PC = [4, 6, 8, 10]       # 売上高比の列
for i, c in enumerate(VC):
    ws5.merge_cells(start_row=4, start_column=c, end_row=4, end_column=c + 1)
hdr(ws5, 4, ['科目'], 2)
for i, c in enumerate(VC):
    hdr(ws5, 4, [f'={q(BSD)}!{CL(PCOLS[i])}4'], c)
hdr(ws5, 4, ['備考'], 11)
R5 = REF['PL']
rows_pl = []


def pl_row(key, text_, formulas=None, values=None, bold=False, indent=0, fill=None):
    global r
    label(ws5, r, text_, bold=bold, indent=indent, fill=fill)
    for i, c in enumerate(VC):
        v = formulas[i] if formulas else values[i]
        num(ws5, r, c, v, bold=bold, fill=fill)
    R5[key] = r
    rows_pl.append(r)
    r += 1


r = 5
pl_row('売上高', '売上高', formulas=[None] * 4, bold=True, fill=GRAY)
pl_row('売上', '売上（製氷・加工氷・商品）', values=PL['売上'], indent=1)
pl_row('保管料', '保管料（冷蔵保管）', values=PL['保管料'], indent=1)
pl_row('受入手数料', '受入手数料', values=PL['受入手数料'], indent=1)
for i, c in enumerate(VC):
    ws5.cell(row=R5['売上高'], column=c, value=f'=SUM({CL(c)}{R5["売上"]}:{CL(c)}{R5["受入手数料"]})')
pl_row('売上原価', '売上原価', formulas=[None] * 4, bold=True)
pl_row('期首商品', '期首商品棚卸高', values=PL['期首商品棚卸高'], indent=1)
pl_row('商品仕入', '商品仕入高', values=PL['商品仕入高'], indent=1)
pl_row('期末商品', '期末商品棚卸高', formulas=[f'={q(BSD)}!{CL(PCOLS[i])}{R["棚卸商品"]}' for i in range(4)], indent=1)
pl_row('商品原価', '商品売上原価', formulas=[f'={CL(c)}{R5["期首商品"]}+{CL(c)}{R5["商品仕入"]}-{CL(c)}{R5["期末商品"]}' for c in VC], indent=1)
pl_row('製造原価', '当期製品製造原価', formulas=[f'=CR!{CL(PCOLS[i])}{R3["製造原価"]}' for i in range(4)], indent=1)
for c in VC:
    ws5.cell(row=R5['売上原価'], column=c, value=f'={CL(c)}{R5["商品原価"]}+{CL(c)}{R5["製造原価"]}')
pl_row('売上総利益', '売上総利益', formulas=[f'={CL(c)}{R5["売上高"]}-{CL(c)}{R5["売上原価"]}' for c in VC], bold=True, fill=GRAY)
pl_row('販管費', '販売費及び一般管理費', formulas=[None] * 4, bold=True)
sga_first = r
for gname, _ in SGA:
    pl_row('販管_' + gname, gname, formulas=[f'=SGA!{CL(PCOLS[i])}{R4["g_" + gname]}' for i in range(4)], indent=1)
pl_row('支払手数料', '支払手数料（一般業務費）', values=PL['支払手数料'], indent=1)
for c in VC:
    ws5.cell(row=R5['販管費'], column=c, value=f'=SUM({CL(c)}{sga_first}:{CL(c)}{r - 1})')
pl_row('営業利益', '営業利益', formulas=[f'={CL(c)}{R5["売上総利益"]}-{CL(c)}{R5["販管費"]}' for c in VC], bold=True, fill=LBLUE)
pl_row('営業外収益', '営業外収益', formulas=[None] * 4, bold=True)
f0 = r
for k in ['受取利息', '受取配当金', '雑収入', '有価証券売却益']:
    pl_row(k, k, values=PL[k], indent=1)
for c in VC:
    ws5.cell(row=R5['営業外収益'], column=c, value=f'=SUM({CL(c)}{f0}:{CL(c)}{r - 1})')
pl_row('営業外費用', '営業外費用', formulas=[None] * 4, bold=True)
f0 = r
for k in ['支払利息', '雑損', '有価証券売却損']:
    pl_row(k, k, values=PL[k], indent=1)
for c in VC:
    ws5.cell(row=R5['営業外費用'], column=c, value=f'=SUM({CL(c)}{f0}:{CL(c)}{r - 1})')
pl_row('経常利益', '経常利益', formulas=[f'={CL(c)}{R5["営業利益"]}+{CL(c)}{R5["営業外収益"]}-{CL(c)}{R5["営業外費用"]}' for c in VC], bold=True, fill=LBLUE)
pl_row('特別利益', '特別利益', values=PL['特別利益'])
pl_row('特別損失', '特別損失', values=PL['特別損失'])
pl_row('税引前', '税引前当期純利益', formulas=[f'={CL(c)}{R5["経常利益"]}+{CL(c)}{R5["特別利益"]}-{CL(c)}{R5["特別損失"]}' for c in VC], bold=True)
pl_row('法人税等', '法人税等', values=PL['法人税等'])
pl_row('当期純利益', '当期純利益', formulas=[f'={CL(c)}{R5["税引前"]}-{CL(c)}{R5["法人税等"]}' for c in VC], bold=True, fill=LBLUE)
r += 1
pl_row('減価償却費', '（参考）減価償却費（製造原価＋販管費）', formulas=[f'=CR!{CL(PCOLS[i])}{R3["製造経費/減価償却費"]}+SGA!{CL(PCOLS[i])}{R4["その他/減価償却費"]}' for i in range(4)])
pl_row('EBITDA', '（参考）EBITDA（営業利益＋減価償却費）', formulas=[f'={CL(c)}{R5["営業利益"]}+{CL(c)}{R5["減価償却費"]}' for c in VC], bold=True)
# 売上高比
for row in rows_pl:
    for i, c in enumerate(VC):
        cc = ws5.cell(row=row, column=PC[i], value=f'=IFERROR({CL(c)}{row}/{CL(c)}${R5["売上高"]},0)')
        cc.number_format = FMT_P
        cc.font = Font(name=NUM, size=8, color='595959')
        cc.border = BORDER
r = max(r + 1, 63)
label(ws5, r, '当期純利益の整合チェック（純資産増減＝当期純利益）', bold=True)
for i, c in enumerate(VC):
    if i == 0:
        ws5.cell(row=r, column=c, value=f'=EXACT({CL(c)}{R5["当期純利益"]},{PL_CHECK["当期純利益"][0]})').font = Font(name=NUM, size=9, bold=True)
    else:
        pc, cc_ = CL(PCOLS[i - 1]), CL(PCOLS[i])
        ws5.cell(row=r, column=c, value=f'=EXACT({CL(c)}{R5["当期純利益"]},{q(BSC)}!{cc_}{R2["純資産合計"]}-{q(BSC)}!{pc}{R2["純資産合計"]})').font = Font(name=NUM, size=9, bold=True)
R5['check'] = r
label(ws5, r + 1, '営業利益の決算書記載値チェック')
for i, c in enumerate(VC):
    ws5.cell(row=r + 1, column=c, value=f'=EXACT({CL(c)}{R5["営業利益"]},{PL_CHECK["営業利益"][i]})').font = Font(name=NUM, size=9)
note(ws5, R5['雑収入'], '67期：無配当保険満期返戻金2,295,725円を含む（非経常）。65期8,482,408円の内容は要確認', col=11)
note(ws5, R5['雑損'], '65期5,192,049円の内容は要確認。67期は資産廃棄', col=11)
note(ws5, R5['支払手数料'], '66期から計上。内容（一過性の専門家報酬等か）は要確認', col=11)
note(ws5, R5['法人税等'], '均等割中心。67期は繰越欠損金1,356,059円を使用し所得517,708円', col=11)

# =====================================================================
sec('VR→')
YRS = ['2024年3月期', '2025年3月期', '2026年3月期']
W = [0.2, 0.3, 0.5]
TAX = 0.34

# ---------- 減価償却 ----------
wsd = ws_new('減価償却', '減価償却資産の検討（評価基準日：2026年3月31日）', {'B': 16, 'C': 8, 'D': 15, 'E': 15, 'F': 15, 'G': 15, 'H': 15, 'I': 15, 'J': 50})
hdr(wsd, 4, ['科目', '件数', '取得価額', '期末簿価', '当期償却額', '当期償却不足額', '償却累計額', '修正額', '評価の考え方'])
RD = REF['減価償却']
r = 5
policy = {'建物': '「土地」シートで固定資産税評価額へ評価替え（本シートでは修正しない）',
          '構築物': '当期償却不足額を減額（累積不足額は資料不足のため未反映・要確認）',
          '機械': '当期償却不足額を減額（累積不足額は資料不足のため未反映・要確認）',
          '車両運搬具': '償却不足なし。簿価評価',
          '器具備品': '償却不足なし。簿価評価'}
first = r
for name, n_, acq, bv, dep, short, acc in DEP_67:
    label(wsd, r, name)
    num(wsd, r, 3, n_, fmt='0')
    for j, v in enumerate([acq, bv, dep, short, acc]):
        num(wsd, r, 4 + j, v)
    adj = f'=-G{r}' if name in ('構築物', '機械') else 0
    num(wsd, r, 9, adj, bold=True)
    text(wsd, r, 10, policy[name])
    RD[name] = r
    r += 1
label(wsd, r, '合計', bold=True, fill=GRAY)
for j in range(4, 10):
    num(wsd, r, j, f'=SUM({CL(j)}{first}:{CL(j)}{r - 1})', bold=True, fill=GRAY)
RD['合計'] = r
r += 2
note(wsd, r, '※主たる工場建物（越谷工場本体、1991年3月取得、取得価額302,052千円、旧定率法・耐用年数26年）は償却累計額55,762千円にとどまり、税務上の償却可能限度に対し大幅な償却不足の状態。')
note(wsd, r + 1, '※機械の当期償却額は管理費108,503円＋製造費18,980,507円。当期償却額合計22,335,325円はPLの減価償却費（製造19,820,169円＋販管2,515,156円）と一致。')

# ---------- 土地（土地・建物） ----------
wsl = ws_new('土地', '不動産（土地・建物）の時価評価', {'B': 10, 'C': 14, 'D': 30, 'E': 12, 'F': 15, 'G': 15, 'H': 10, 'I': 15, 'J': 15, 'K': 50})
hdr(wsl, 4, ['種類', '区分', '所在地', '面積（㎡）', '簿価', '固定資産税評価額', '割戻率', '時価', '修正額', '評価の考え方'])
RL = REF['土地']
r = 5
for kind, area_, addr, m2, bv, ev, rem in REAL_ESTATE_OWNED:
    label(wsl, r, kind)
    text(wsl, r, 3, area_)
    text(wsl, r, 4, addr)
    num(wsl, r, 5, m2, fmt='#,##0.00')
    bv_f = f'={q(BSD)}!F{R[kind]}'
    num(wsl, r, 6, bv_f)
    num(wsl, r, 7, ev)
    if kind == '土地':
        num(wsl, r, 8, 0.7, fmt='0.0')
        num(wsl, r, 9, f'=ROUND(G{r}/H{r},0)')
        text(wsl, r, 11, '固定資産税評価額は公示価格水準の概ね7割とされるため÷0.7で時価を推計。路線価による検証は要確認')
    else:
        num(wsl, r, 8, 1.0, fmt='0.0')
        num(wsl, r, 9, f'=ROUND(G{r}/H{r},0)')
        text(wsl, r, 11, '建物は固定資産税評価額を時価相当額とみなす（再調達原価法による検証は未実施）。簿価は会計上の償却停止により実態を大きく上回る')
    num(wsl, r, 10, f'=I{r}-F{r}', bold=True)
    RL[kind] = r
    r += 1
label(wsl, r, '合計', bold=True, fill=GRAY)
for j in [6, 7, 9, 10]:
    num(wsl, r, j, f'=SUM({CL(j)}5:{CL(j)}{r - 1})', bold=True, fill=GRAY)
RL['合計'] = r
note(wsl, r + 2, '※固定資産税評価額は受領資料（2025年3月期時点の所有物件一覧）記載値。最新年度の評価証明書で更新要。建物の面積は登記2,093.87㎡／内訳書2,065.00㎡で差異あり（要確認）。')
note(wsl, r + 3, '※代表者個人所有の土地（栃木銀行・商工中金の担保提供）は会社資産ではないため評価対象外。所在地・面積・譲渡後の取扱いは要確認。')

# ---------- 保険 ----------
wsi = ws_new('保険', '保険積立金の評価', {'B': 22, 'C': 15, 'D': 15, 'E': 15, 'F': 55})
hdr(wsi, 4, ['保険会社', '簿価', '解約返戻金', '修正額', '備考'])
RI = REF['保険']
r = 5
for nm, v in [('大同生命保険', 8461965), ('第一生命保険', 900529)]:
    label(wsi, r, nm)
    num(wsi, r, 3, v)
    num(wsi, r, 4, f'=C{r}')
    num(wsi, r, 5, f'=D{r}-C{r}', bold=True)
    text(wsi, r, 6, '解約返戻金額の資料未受領のため簿価評価（要確認）。被保険者・契約形態も要確認')
    r += 1
label(wsi, r, '合計', bold=True, fill=GRAY)
for j in [3, 4, 5]:
    num(wsi, r, j, f'=SUM({CL(j)}5:{CL(j)}{r - 1})', bold=True, fill=GRAY)
RI['合計'] = r
note(wsi, r + 2, '※BS上は現預金の「諸掛金」に含めて表示。評価上は保険積立金として区分する。67期に無配当満期定期保険の返戻金2,295,725円を雑収入計上。')

# ---------- 有価証券 / 売掛・棚卸 / 未払人件費・賞与 / 退職金 ----------
for nm, title, lines in [
    ('有価証券', '有価証券の評価', ['該当なし：上場有価証券の保有なし。', '出資金820,000円（越谷市工業団地組合800,000円、埼玉縣信用金庫10,000円ほか）は換金性に乏しいが少額のため簿価評価。']),
    ('売掛・棚卸', '売掛金・棚卸資産の評価', ['売掛金：年齢表・補助残高一覧が未受領のため簿価評価（要確認）。月末締め翌月末回収（法人事業概況説明書）。',
                                      '　67期内訳書の上位先合計が残高を1,723,680円超過しており、(株)小野田商店の計上区分を確認する（要確認）。',
                                      '　(株)小野田商店は代表者と同姓であり、関連当事者取引か否かを確認する（要確認）。',
                                      '棚卸資産：実地棚卸（2026年3月31日実施）に基づく計上。商品20,661千円（加工氷・袋氷等）、原材料12,953千円（ロール・段ボール・ビニール袋等）。滞留・陳腐化の有無は要確認のため簿価評価。']),
    ('未払人件費・賞与', '未払人件費・賞与引当の検討', ['給与：月末締め翌月20日払い。3月分給与5,720,681円は未払金に計上済みのため追加計上なし。',
                                            '賞与：賞与引当金の計上なし。支給時期・支給対象期間が未確認のため評価上も未計上（要確認：支給規程・夏季賞与の算定期間）。']),
    ('退職金', '退職給付債務の検討', ['退職給付引当金の計上なし。販管費に退職共済金120,000円／年（中小企業退職金共済と推察）。',
                                 '67期に製造原価の退職給与6,177,985円が発生（過去3期は発生なし）。退職金規程の有無・社内積立部分の要支給額は未確認のため、評価上は債務計上せず（要確認）。'])]:
    w_ = ws_new(nm, title, {'B': 120})
    for i, t in enumerate(lines):
        text(w_, 4 + i, 2, t, wrap=True)

# ---------- 修正BS(借方) ----------
wsb = ws_new('修正BS(借方)', '修正貸借対照表（資産の部）　評価基準日：2026年3月31日　単位：千円', {'B': 30, 'C': 15, 'D': 15, 'E': 15, 'F': 55})
hdr(wsb, 4, ['科目', '簿価', '修正額', '修正後', '修正内容'])
RB = REF['修正BS(借方)']
r = 5
items_b = [('現金及び預金（保険積立除く）', '+'.join(f'{q(BSD)}!F{R[k]}' for k in ['現金', '当座預金', '普通預金', '定期預金', '定期積金']), None, '定期預金の内訳差異6,800千円は要確認'),
           ('売掛金', f'{q(BSD)}!F{R["売掛金"]}', None, '年齢表未受領のため簿価'),
           ('棚卸資産', f'{q(BSD)}!F{R["sub_inv"]}', None, '滞留在庫は要確認のため簿価'),
           ('その他流動資産', f'{q(BSD)}!F{R["sub_oca"]}', None, '従業員貸付金・未収金・前払費用等'),
           ('保険積立金', f'{q(BSD)}!F{R["諸掛金（生命保険積立）"]}', f"='保険'!E{RI['合計']}", '解約返戻金未受領のため簿価（要確認）'),
           ('建物', f'{q(BSD)}!F{R["建物"]}', f"='土地'!J{RL['建物']}", '固定資産税評価額へ評価替え'),
           ('構築物', f'{q(BSD)}!F{R["構築物"]}', f"='減価償却'!I{RD['構築物']}", '当期償却不足額を減額'),
           ('機械', f'{q(BSD)}!F{R["機械"]}', f"='減価償却'!I{RD['機械']}", '当期償却不足額を減額'),
           ('車両運搬具・器具備品', f'{q(BSD)}!F{R["車両運搬具"]}+{q(BSD)}!F{R["器具備品"]}', None, '償却不足なし'),
           ('土地', f'{q(BSD)}!F{R["土地"]}', f"='土地'!J{RL['土地']}", '固定資産税評価額÷0.7'),
           ('特許商標権', f'{q(BSD)}!F{R["特許商標権"]}', None, '簿価（商標の内容は要確認）'),
           ('出資金', f'{q(BSD)}!F{R["出資金"]}', None, '簿価')]
first = r
for nm, bf, adj, rem in items_b:
    label(wsb, r, nm)
    num(wsb, r, 3, '=' + bf)
    num(wsb, r, 4, adj if adj else 0)
    num(wsb, r, 5, f'=C{r}+D{r}', bold=True)
    text(wsb, r, 6, rem)
    RB[nm] = r
    r += 1
label(wsb, r, '資産合計', bold=True, fill=LBLUE)
for j in [3, 4, 5]:
    num(wsb, r, j, f'=SUM({CL(j)}{first}:{CL(j)}{r - 1})', bold=True, fill=LBLUE)
RB['資産合計'] = r
label(wsb, r + 1, '簿価の検算（BS資産合計との一致）')
wsb.cell(row=r + 1, column=3, value=f'=EXACT(C{r},{q(BSD)}!F{R["資産合計"]})').font = Font(name=NUM, size=9)

# ---------- 修正BS(貸方) ----------
wsc = ws_new('修正BS(貸方) ', '修正貸借対照表（負債・純資産の部）　評価基準日：2026年3月31日　単位：千円', {'B': 34, 'C': 15, 'D': 15, 'E': 15, 'F': 55})
hdr(wsc, 4, ['科目', '簿価', '修正額', '修正後', '修正内容'])
RC = REF['修正BS(貸方) ']
r = 5
items_c = [('支払手形・買掛金', f'{q(BSC)}!F{R2["支払手形"]}+{q(BSC)}!F{R2["買掛金"]}', '—'),
           ('借入金（短期・長期）', f'{q(BSC)}!F{R2["有利子負債"]}', '代表者借入14,000千円・公庫資本性ローン105,420千円を含む'),
           ('未払金', f'{q(BSC)}!F{R2["未払金"]}', '3月分給与・未払消費税を含む'),
           ('未払法人税等', f'{q(BSC)}!F{R2["税金引当金（未払法人税等）"]}', '—'),
           ('その他流動負債', f'{q(BSC)}!F{R2["預り金"]}+{q(BSC)}!F{R2["仮受金"]}', '—'),
           ('賞与引当金', '0', '支給規程未確認のため未計上（要確認）'),
           ('退職給付引当金', '0', '退職金規程未確認のため未計上（要確認）')]
first = r
for nm, bf, rem in items_c:
    label(wsc, r, nm)
    num(wsc, r, 3, '=' + bf)
    num(wsc, r, 4, 0)
    num(wsc, r, 5, f'=C{r}+D{r}', bold=True)
    text(wsc, r, 6, rem)
    RC[nm] = r
    r += 1
label(wsc, r, '負債合計', bold=True, fill=GRAY)
for j in [3, 4, 5]:
    num(wsc, r, j, f'=SUM({CL(j)}{first}:{CL(j)}{r - 1})', bold=True, fill=GRAY)
RC['負債合計'] = r
r += 1
label(wsc, r, '評価差額に係る税効果')
num(wsc, r, 3, 0)
num(wsc, r, 4, f"=\'税効果\'!C8")
num(wsc, r, 5, f'=C{r}+D{r}', bold=True)
text(wsc, r, 6, '「税効果」シート参照')
RC['税効果'] = r
r += 1
label(wsc, r, '純資産（時価純資産）', bold=True, fill=LBLUE)
for j, col in [(3, 'C'), (4, 'D'), (5, 'E')]:
    num(wsc, r, j, f"={q('修正BS(借方)')}!{col}{RB['資産合計']}-{col}{RC['負債合計']}-{col}{RC['税効果']}", bold=True, fill=LBLUE)
RC['純資産'] = r
label(wsc, r + 1, '簿価純資産の検算（BS純資産との一致）')
wsc.cell(row=r + 1, column=3, value=f'=EXACT(C{r},{q(BSC)}!F{R2["純資産合計"]})').font = Font(name=NUM, size=9)

# ---------- 税効果 ----------
wst = ws_new('税効果', '評価差額に係る税効果', {'B': 40, 'C': 16, 'D': 60})
label(wst, 5, '評価差額（資産修正額合計－負債修正額合計）')
num(wst, 5, 3, f"={q('修正BS(借方)')}!D{RB['資産合計']}-{q('修正BS(貸方) ')}!D{RC['負債合計']}")
label(wst, 6, '実効税率')
num(wst, 6, 3, TAX, fmt='0%')
label(wst, 7, '税効果（評価益の場合のみ繰延税金負債を認識）')
num(wst, 7, 3, '=IF(C5>0,ROUND(C5*C6,0),0)')
label(wst, 8, '修正BSへ反映する税効果（負債）', bold=True)
num(wst, 8, 3, '=C7', bold=True)
text(wst, 10, 2, '評価差額が純額で評価損（含み損）となる場合、繰延税金資産の回収可能性（課税所得の見込み）が不確実なため保守的に認識しない。土地の評価益に対応する税効果は建物の評価損と相殺して判定。', wrap=True)
wst.merge_cells('B10:D11')

# ---------- ネットキャッシュ ----------
wsn = ws_new('ネットキャッシュ', 'ネットキャッシュの算定（評価基準日：2026年3月31日）', {'B': 40, 'C': 16, 'D': 60})
RN = REF['ネットキャッシュ']
hdr(wsn, 4, ['項目', '金額', '判断根拠'])
r = 5
for nm, f, rem, key in [
        ('現金及び預金（保険積立除く）', f"={q('修正BS(借方)')}!E{RB['現金及び預金（保険積立除く）']}", 'キャッシュライク', 'cash'),
        ('保険積立金', f"={q('修正BS(借方)')}!E{RB['保険積立金']}", '換金可能な非事業用資産としてキャッシュライク（解約返戻金で要更新）', 'ins'),
        ('キャッシュライク計', '=C5+C6', '', 'cl')]:
    label(wsn, r, nm, bold=(key == 'cl'))
    num(wsn, r, 3, f, bold=(key == 'cl'))
    text(wsn, r, 4, rem)
    RN[key] = r
    r += 1
for nm, f, rem, key in [
        ('借入金（金融機関）', f"={q('修正BS(貸方) ')}!E{RC['借入金（短期・長期）']}-14000000", '有利子負債。うち日本公庫資本性ローン105,420千円を含む', 'bank'),
        ('代表者借入金', '=14000000', '代表者からの借入。クロージング時の返済／放棄の取扱いは要協議', 'owner'),
        ('未払法人税等', f"={q('修正BS(貸方) ')}!E{RC['未払法人税等']}", 'デットライク', 'tax'),
        ('デットライク計', '=SUM(C8:C10)', '', 'dl')]:
    label(wsn, r, nm, bold=(key == 'dl'))
    num(wsn, r, 3, f, bold=(key == 'dl'))
    text(wsn, r, 4, rem)
    RN[key] = r
    r += 1
label(wsn, r, 'ネットキャッシュ（△はネットデット）', bold=True, fill=LBLUE)
num(wsn, r, 3, f'=C{RN["cl"]}-C{RN["dl"]}', bold=True, fill=LBLUE)
RN['nc'] = r
r += 2
label(wsn, r, '（参考）資本性ローンを資本とみなした場合のネットキャッシュ')
num(wsn, r, 3, f'=C{RN["nc"]}+{CAPITAL_LOAN_67}')
RN['nc_cap'] = r
note(wsn, r + 1, '※日本政策金融公庫の資本性ローンは金融機関の債務者区分上は自己資本とみなされるが、期限一括返済の債務であり評価上は有利子負債として扱う。')

# ---------- 修正PL / 修正SGA / 修正CR ----------
# 調整項目（正常収益力）：直近3期（2024/3〜2026/3）
ADJ = [  # key, 科目, 3期の修正額（＋は利益加算）, 理由
    ('退職給与', '退職給与（製造原価・販管費）',
     [f'=SGA!D{R4["給与関係費/退職給与"]}+CR!D{R3["労務費/退職給与"]}',
      f'=SGA!E{R4["給与関係費/退職給与"]}+CR!E{R3["労務費/退職給与"]}',
      f'=SGA!F{R4["給与関係費/退職給与"]}+CR!F{R3["労務費/退職給与"]}'],
     '退職者に対する臨時的な支出と推察し加算（退職金規程・対象者は要確認）'),
    ('役員報酬', '役員報酬（新体制の標準報酬へ置換）', [0, 0, 0],
     '現行7,200千円／年は経営者報酬として過大ではなく、後任経営者の標準報酬と同水準と想定し調整なし'),
]
wsp = ws_new('修正PL', '正常収益力の算定（修正損益計算書）　単位：千円', {'B': 36, 'C': 13, 'D': 13, 'E': 13, 'F': 13, 'G': 13, 'H': 13, 'I': 13, 'J': 13, 'K': 13, 'L': 50})
RP = REF['修正PL']
for i, y in enumerate(YRS):
    c0 = 3 + i * 3
    wsp.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 2)
    hdr(wsp, 4, [y], c0)
    hdr(wsp, 5, ['帳簿価額', '修正額', '修正後'], c0)
hdr(wsp, 4, ['科目'], 2)
hdr(wsp, 5, [''], 2)
hdr(wsp, 4, ['備考'], 12)
r = 6
src = {'売上高': '売上高', '売上原価': '売上原価', '売上総利益': '売上総利益', '販管費': '販管費', '営業利益': '営業利益', '減価償却費': '減価償却費'}
PLV = {1: 'E', 2: 'G', 3: 'I'}  # PLシートの値列（24/3=E,25/3=G,26/3=I）
lines = [('売上高', '売上高', None), ('売上原価', '売上原価', 'CR'), ('売上総利益', '売上総利益', None),
         ('販管費', '販売費及び一般管理費', 'SGA'), ('営業利益', '営業利益', None), ('減価償却費', '減価償却費', None), ('EBITDA', 'EBITDA', None)]
for key, nm, kind in lines:
    label(wsp, r, nm, bold=key in ('売上総利益', '営業利益', 'EBITDA'), fill=LBLUE if key in ('営業利益', 'EBITDA') else None)
    RP[key] = r
    r += 1
for i in range(3):
    b, m, a = 3 + i * 3, 4 + i * 3, 5 + i * 3
    pcol = PLV[i + 1]
    for key in ['売上高', '売上原価', '売上総利益', '販管費', '営業利益', '減価償却費']:
        num(wsp, RP[key], b, f"=PL!{pcol}{R5[key]}")
    # 修正額：売上原価＝修正CR、販管費＝修正SGA
    num(wsp, RP['売上高'], m, 0)
    num(wsp, RP['売上原価'], m, f"={q('修正CR')}!{CL(4 + i * 3)}{{CRADJ}}")
    num(wsp, RP['販管費'], m, f"={q('修正SGA')}!{CL(4 + i * 3)}{{SGAADJ}}")
    num(wsp, RP['売上総利益'], m, f'={CL(m)}{RP["売上高"]}-{CL(m)}{RP["売上原価"]}')
    num(wsp, RP['営業利益'], m, f'={CL(m)}{RP["売上総利益"]}-{CL(m)}{RP["販管費"]}')
    num(wsp, RP['減価償却費'], m, 0)
    for key in ['売上高', '売上原価', '売上総利益', '販管費', '営業利益', '減価償却費']:
        num(wsp, RP[key], a, f'={CL(b)}{RP[key]}+{CL(m)}{RP[key]}', bold=key in ('売上総利益', '営業利益'))
    for col in (b, m, a):
        num(wsp, RP['EBITDA'], col, f'={CL(col)}{RP["営業利益"]}+{CL(col)}{RP["減価償却費"]}', bold=True, fill=LBLUE)
r += 1
label(wsp, r, '【調整項目】（＋は利益に加算）', bold=True)
r += 1
for key, nm, fs, why in ADJ:
    label(wsp, r, nm, indent=1)
    for i in range(3):
        num(wsp, r, 4 + i * 3, fs[i])
    text(wsp, r, 12, why)
    RP['adj_' + key] = r
    r += 1
label(wsp, r, '調整額合計', bold=True)
for i in range(3):
    num(wsp, r, 4 + i * 3, f'=SUM({CL(4 + i * 3)}{RP["adj_退職給与"]}:{CL(4 + i * 3)}{r - 1})', bold=True)
RP['adj_total'] = r
r += 1
label(wsp, r, '検算（調整額合計＝営業利益の修正額）')
for i in range(3):
    wsp.cell(row=r, column=4 + i * 3, value=f'=EXACT({CL(4 + i * 3)}{RP["adj_total"]},{CL(4 + i * 3)}{RP["営業利益"]})').font = Font(name=NUM, size=9)
r += 2
for t in ['【調整しなかった項目と理由】',
          '・減価償却費：建物は会計上償却を停止しており、営業利益ベースの収益力は償却負担を過少に織り込んでいる可能性がある。このためEBITDAベースのマーケットアプローチを重視する。',
          '・支払手数料（66期3,155,869円・67期1,764,661円）：内容未確認のため調整せず。一過性の専門家報酬等であれば加算余地あり（要確認）。',
          '・地代家賃（代表者からの駐車場賃借545,460円／年）：譲渡後も継続利用を前提に調整せず。条件は要協議。',
          '・雑収入の保険返戻金（67期2,295,725円）：営業外項目のため営業利益・EBITDAに影響なし。']:
    note(wsp, r, t, size=9)
    r += 1

for nm, src_sheet, groups_, adj_key, adj_label in [('修正SGA', 'SGA', SGA, '給与関係費/退職給与', '退職給与'), ('修正CR', 'CR', CR, '労務費/退職給与', '退職給与')]:
    w_ = ws_new(nm, f'{"販売費及び一般管理費" if nm == "修正SGA" else "製造原価"}の修正　単位：千円', {'B': 30, 'C': 13, 'D': 13, 'E': 13, 'F': 13, 'G': 13, 'H': 13, 'I': 13, 'J': 13, 'K': 13, 'L': 40})
    RR = REF[nm]
    for i, y in enumerate(YRS):
        c0 = 3 + i * 3
        w_.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 2)
        hdr(w_, 4, [y], c0)
        hdr(w_, 5, ['帳簿価額', '修正額', '修正後'], c0)
    hdr(w_, 4, ['科目'], 2)
    hdr(w_, 5, [''], 2)
    hdr(w_, 4, ['修正理由'], 12)
    src_ref = REF[src_sheet]
    r = 6
    item_rows = []
    for gname, items in groups_:
        for iname, _ in items:
            k = gname + '/' + iname
            label(w_, r, f'{gname}／{iname}')
            for i in range(3):
                b, m, a = 3 + i * 3, 4 + i * 3, 5 + i * 3
                num(w_, r, b, f'={src_sheet}!{CL(P3[i])}{src_ref[k]}')
                num(w_, r, m, f'=-{CL(b)}{r}' if k == adj_key else 0)
                num(w_, r, a, f'={CL(b)}{r}+{CL(m)}{r}')
            if k == adj_key:
                text(w_, r, 12, '臨時的な退職金支出として控除（要確認）')
            item_rows.append(r)
            r += 1
    label(w_, r, '合計' if nm == '修正SGA' else '製造原価合計', bold=True, fill=LBLUE)
    for i in range(3):
        for j in range(3):
            col = 3 + i * 3 + j
            if nm == '修正CR':
                # 原材料期末・期末仕掛品は控除項目
                minus = {'原材料費/原材料期末棚卸高', '製造勘定/期末仕掛品'}
                keys = [g + '/' + it for g, its in groups_ for it, _ in its]
                f = '=' + '+'.join(('-' if kk in minus else '') + f'{CL(col)}{rr}' for kk, rr in zip(keys, item_rows)).replace('+-', '-')
            else:
                f = f'=SUM({CL(col)}{item_rows[0]}:{CL(col)}{item_rows[-1]})'
            num(w_, r, col, f, bold=True, fill=LBLUE)
    RR['合計'] = r
    label(w_, r + 1, '検算（帳簿価額合計＝原本シート）')
    for i in range(3):
        tgt = f'SGA!{CL(P3[i])}{R4["合計"]}' if nm == '修正SGA' else f'CR!{CL(P3[i])}{R3["製造原価"]}'
        w_.cell(row=r + 1, column=3 + i * 3, value=f'=EXACT({CL(3 + i * 3)}{r},{tgt})').font = Font(name=NUM, size=9)

# 修正PLの修正額参照を確定
for i in range(3):
    m = 4 + i * 3
    c1 = wsp.cell(row=RP['売上原価'], column=m)
    c1.value = c1.value.replace('{CRADJ}', str(REF['修正CR']['合計']))
    c2 = wsp.cell(row=RP['販管費'], column=m)
    c2.value = c2.value.replace('{SGAADJ}', str(REF['修正SGA']['合計']))

# ---------- 年買法 ----------
wsy = ws_new('年買法', '年買法（時価純資産＋営業権）　単位：千円', {'B': 36, 'C': 15, 'D': 15, 'E': 15, 'F': 15, 'G': 50})
RY = REF['年買法']
hdr(wsy, 4, ['項目'] + YRS + ['合計／備考'])
label(wsy, 5, '修正後営業利益')
label(wsy, 6, '税考慮額（△34%）')
label(wsy, 7, '税考慮後営業利益', bold=True)
label(wsy, 8, '加重')
label(wsy, 9, '加重後', bold=True)
for i in range(3):
    c = 3 + i
    num(wsy, 5, c, f"={q('修正PL')}!{CL(5 + i * 3)}{RP['営業利益']}")
    num(wsy, 6, c, f'=-ROUND(MAX({CL(c)}5,0)*{TAX},0)')
    num(wsy, 7, c, f'={CL(c)}5+{CL(c)}6', bold=True)
    num(wsy, 8, c, W[i], fmt='0%')
    num(wsy, 9, c, f'=ROUND({CL(c)}7*{CL(c)}8,0)', bold=True)
label(wsy, 10, '基準営業利益（加重平均）', bold=True, fill=LBLUE)
num(wsy, 10, 6, '=SUM(C9:E9)', bold=True, fill=LBLUE)
text(wsy, 10, 7, '加重平均は加重後の行（9行）のみを合計（修正額行を含めない）')
hdr(wsy, 12, ['', '下限（1年）', '中央（2年）', '上限（3年）', ''])
label(wsy, 13, '営業権倍率（年）')
for j, k in enumerate([1, 2, 3]):
    num(wsy, 13, 3 + j, k, fmt='0')
label(wsy, 14, '営業権')
label(wsy, 15, '時価純資産')
label(wsy, 16, '想定株式価値（時価純資産＋営業権）', bold=True)
label(wsy, 17, '株式価値（マイナスは0円表示）', bold=True)
for j in range(3):
    c = 3 + j
    num(wsy, 14, c, f'=ROUND($F$10*{CL(c)}13,0)')
    num(wsy, 15, c, f"={q('修正BS(貸方) ')}!E{RC['純資産']}")
    num(wsy, 16, c, f'={CL(c)}14+{CL(c)}15', bold=True, fill=LBLUE)
    num(wsy, 17, c, f'=MAX({CL(c)}16,0)', bold=True)
label(wsy, 21, '（感応度）建物を簿価で評価した場合の時価純資産')
label(wsy, 22, '（感応度）同 想定株式価値', bold=True)
for j in range(3):
    c = 3 + j
    num(wsy, 21, c, f"={CL(c)}15-'土地'!J{RL['建物']}")
    num(wsy, 22, c, f'={CL(c)}14+{CL(c)}21', bold=True)
note(wsy, 23, '※建物の評価方法（固定資産税評価額／簿価）が結論を左右するため感応度を併記。簿価は会計上の償却停止により過大であり、本評価の基本ケースは固定資産税評価額。')
RY.update({'base': 10, 'gw': 14, 'na': 15, 'val': 16, 'val0': 17, 'sens_na': 21, 'sens_val': 22})
note(wsy, 19, '※営業権倍率1〜3年は国内中小企業M&Aのボリュームゾーン。加重は直近期50%・前期30%・前々期20%。税考慮は営業利益がマイナスの期は0。')

# ---------- 参照企業 ----------
wsr = ws_new('参照企業', '類似上場会社（EV/EBITDA倍率の参照）', {'B': 8, 'C': 30, 'D': 30, 'E': 14, 'F': 40})
hdr(wsr, 4, ['コード', '会社名', '事業', 'EV/EBITDA', '備考'])
peers = [('2874', '(株)ヨコレイ', '冷蔵倉庫・食品販売'), ('2871', '(株)ニチレイ', '低温物流・加工食品'),
         ('9369', '(株)キユーソー流通システム', '食品物流（低温）'), ('9351', '東洋埠頭(株)', '冷蔵倉庫・港湾'),
         ('2884', 'ヨシムラ・フード・HD', '食品製造（中小企業M&A型）')]
for i, (cd, nm, biz) in enumerate(peers):
    r = 5 + i
    text(wsr, r, 2, cd)
    text(wsr, r, 3, nm)
    text(wsr, r, 4, biz)
    text(wsr, r, 5, '要更新')
    text(wsr, r, 6, '算定日時点の株価・直近決算で更新する')
note(wsr, 11, '※製氷専業の上場会社は存在しないため、低温物流・食品製造を参照。倍率は算定日時点で更新する前提であり、本試算では中小企業M&Aの実務レンジ3〜5倍を採用。')

# ---------- EV EBITDAマルチプル ----------
wse = ws_new('EV EBITDAマルチプル', 'EV/EBITDAマルチプル法　単位：千円', {'B': 40, 'C': 15, 'D': 15, 'E': 15, 'F': 15, 'G': 50})
RE = REF['EV EBITDAマルチプル']
hdr(wse, 4, ['項目'] + YRS + ['合計／備考'])
label(wse, 5, '修正後EBITDA')
label(wse, 6, '加重')
label(wse, 7, '加重後', bold=True)
for i in range(3):
    c = 3 + i
    num(wse, 5, c, f"={q('修正PL')}!{CL(5 + i * 3)}{RP['EBITDA']}")
    num(wse, 6, c, W[i], fmt='0%')
    num(wse, 7, c, f'=ROUND({CL(c)}5*{CL(c)}6,0)', bold=True)
label(wse, 8, '基準EBITDA（加重平均）', bold=True, fill=LBLUE)
num(wse, 8, 6, '=SUM(C7:E7)', bold=True, fill=LBLUE)
label(wse, 9, '（参考）修正後EBITDA 3期単純平均')
num(wse, 9, 6, '=ROUND(AVERAGE(C5:E5),0)')
hdr(wse, 11, ['', '下限（3倍）', '中央（4倍）', '上限（5倍）', ''])
label(wse, 12, 'EV/EBITDA倍率')
label(wse, 13, '事業価値（EV）')
label(wse, 14, '非流動性ディスカウント（△30%）')
label(wse, 15, 'ディスカウント後事業価値')
label(wse, 16, 'ネットキャッシュ（△はネットデット）')
label(wse, 17, '想定株式価値', bold=True)
label(wse, 18, '株式価値（マイナスは0円表示）', bold=True)
label(wse, 20, '（参考）資本性ローンを資本とみなした場合の株式価値')
for j, k in enumerate([3, 4, 5]):
    c = 3 + j
    num(wse, 12, c, k, fmt='0.0"倍"')
    num(wse, 13, c, f'=ROUND($F$8*{CL(c)}12,0)')
    num(wse, 14, c, f'=-ROUND({CL(c)}13*0.3,0)')
    num(wse, 15, c, f'={CL(c)}13+{CL(c)}14')
    num(wse, 16, c, f"='ネットキャッシュ'!C{RN['nc']}")
    num(wse, 17, c, f'={CL(c)}15+{CL(c)}16', bold=True, fill=LBLUE)
    num(wse, 18, c, f'=MAX({CL(c)}17,0)', bold=True)
    num(wse, 20, c, f"={CL(c)}15+'ネットキャッシュ'!C{RN['nc_cap']}")
RE.update({'base': 8, 'avg': 9, 'ev': 13, 'ev_d': 15, 'nc': 16, 'val': 17, 'val0': 18, 'val_cap': 20})
note(wse, 22, '※非流動性ディスカウント：上場会社の株価倍率には市場性（流動性）が織込まれているため、非上場株式への適用にあたり控除。コストアプローチには適用しない。')

# ---------- 株式価値 ----------
wsv = ws_new('株式価値 ', '株式価値の試算結果（サマリー）　単位：千円', {'B': 44, 'C': 16, 'D': 16, 'E': 16, 'F': 60})
RV = REF['株式価値 ']
hdr(wsv, 4, ['手法', '下限', '中央', '上限', '備考'])
label(wsv, 5, 'コストアプローチ（年買法）')
label(wsv, 6, 'マーケットアプローチ（EV/EBITDA法）')
label(wsv, 7, '（参考）時価純資産')
label(wsv, 8, '（参考）簿価純資産')
label(wsv, 9, '（参考）事業価値EV（ディスカウント後）')
label(wsv, 10, '（参考）ネットデット')
for j in range(3):
    c = 3 + j
    num(wsv, 5, c, f"='年買法'!{CL(c)}16", bold=True)
    num(wsv, 6, c, f"={q('EV EBITDAマルチプル')}!{CL(c)}17", bold=True)
    num(wsv, 7, c, f"={q('修正BS(貸方) ')}!E{RC['純資産']}")
    num(wsv, 8, c, f"={q(BSC)}!F{R2['純資産合計']}")
    num(wsv, 9, c, f"={q('EV EBITDAマルチプル')}!{CL(c)}15")
    num(wsv, 10, c, f"=-'ネットキャッシュ'!C{RN['nc']}")
text(wsv, 5, 6, '時価純資産がマイナスのため、営業権を加算しても株式価値はマイナス')
text(wsv, 6, 6, '有利子負債（461,024千円）が事業価値を上回るためマイナス')
label(wsv, 11, '（感応度）年買法・建物簿価評価の場合')
for j in range(3):
    num(wsv, 11, 3 + j, f"='年買法'!{CL(3 + j)}22")
text(wsv, 11, 6, '建物の評価替えを行わない場合。買手が建物の残存価値をどう見るかで大きく変動')
label(wsv, 12, '結論：算定上の株式価値', bold=True, fill=LBLUE)
num(wsv, 12, 3, '=MAX(0,MAX(C5:E6))', bold=True, fill=LBLUE)
text(wsv, 12, 6, '基本ケース（建物＝固定資産税評価額）では両手法とも算定上マイナスで、株式価値は備忘価額程度。ただし建物を簿価評価すると年買法はプラス（11行）に転じるため、建物時価の検証が最重要。実質的な交渉論点は有利子負債の承継・代表者保証の解除・担保差替え')
RV.update({'cost': 5, 'market': 6, 'sens': 11, 'concl': 12})

# =====================================================================
sec('IM→')

# ---------- 会社概要 ----------
CO = [
    ('会社名', '越ヶ谷製氷冷蔵株式会社', '登記'),
    ('ブランド', '純氷「キングアイス」', '公式サイト（king-ice.com）'),
    ('本店所在地', '埼玉県越谷市蒲生寿町19番15号', '登記'),
    ('工場・営業所', '越谷工場：埼玉県越谷市大字平方999-4（平方工業団地内）／松戸営業所：千葉県松戸市稔台5-14-3', '勘定科目内訳書'),
    ('設立', '1960年（昭和35年）3月18日', '登記'),
    ('資本金', '10,000千円（2026年5月15日に11,000千円から減資）', '登記'),
    ('代表者', '代表取締役 小野田 洋', '登記'),
    ('株主', '小野田 洋 100%（発行済株式数は登記22,000株／申告書220株で相違・要確認）', '登記・法人税申告書'),
    ('役員', '取締役1名（代表取締役のみ）／取締役会・監査役 設置なし', '登記'),
    ('従業員数', '25名（常勤役員1名・従業員24名、2026年3月期末）', '法人事業概況説明書'),
    ('事業内容', '純氷（角氷・加工氷・袋氷・特殊氷）の製造及び卸小売販売、ドライアイス等の商品販売、冷蔵保管', '登記・概況説明書'),
    ('主要取引先', '(株)日本アクセス／(株)日本橋冷凍手島商店／三菱食品(株)／金木屋商事(株)／(株)クールテックサガワ　他', '勘定科目内訳書'),
    ('主要設備', '製氷設備4台（日産40トン）、袋詰砕氷設備1台（日産2,500ケース）、冷蔵設備3台（F級1,000トン）', '法人事業概況説明書'),
    ('許認可', '営業許可（氷雪製造業）、食品営業許可（氷雪販売業）', '受領資料'),
    ('売上高', '=TEXT(ROUND(PL!I5/1000,0),"#,##0")&"千円（2026年3月期）"', 'マスターPL'),
    ('取引金融機関', '日本政策金融公庫、栃木銀行、商工組合中央金庫、足利銀行、埼玉縣信用金庫、埼玉りそな銀行', '借入金内訳書'),
    ('譲渡理由', '後継者不在', 'ヒアリング（受領資料）'),
    ('譲渡スキーム', '株式100%譲渡', '同上'),
]
wsk = ws_new('会社概要', '会社概要', {'B': 16, 'C': 90, 'D': 26})
hdr(wsk, 4, ['項目', '内容', '出所'])
RK = REF['会社概要']
for i, (k, v, s) in enumerate(CO):
    r = 5 + i
    text(wsk, r, 2, k, bold=True, fill=GRAY)
    text(wsk, r, 3, v)
    text(wsk, r, 4, s, size=8)
    RK[k] = r

# ---------- PLハイライト ----------
wsh = ws_new('PLハイライト', 'PLハイライト　単位：千円', {'B': 34, 'C': 14, 'D': 14, 'E': 14, 'F': 14, 'G': 50})
RH = REF['PLハイライト']
hdr(wsh, 4, ['科目'] + [f'={q(BSD)}!{CL(c)}4' for c in PCOLS] + ['備考'])
r = 5
for key, nm, kind in [('売上高', '売上高', 'v'), ('前期比', '　前期比', 'g'), ('売上総利益', '売上総利益', 'v'), ('粗利率', '　売上総利益率', 'p:売上総利益'),
                      ('営業利益', '営業利益', 'v'), ('営業利益率', '　営業利益率', 'p:営業利益'), ('経常利益', '経常利益', 'v'), ('経常利益率', '　経常利益率', 'p:経常利益'),
                      ('当期純利益', '当期純利益', 'v'), ('減価償却費', '減価償却費', 'v'), ('EBITDA', 'EBITDA（営業利益＋減価償却費）', 'v')]:
    label(wsh, r, nm, bold=(kind == 'v' and key in ('売上高', '営業利益', 'EBITDA')))
    for i, c in enumerate(PCOLS):
        vc = CL(VC[i])
        if kind == 'v':
            num(wsh, r, c, f'=PL!{vc}{R5[key]}')
        elif kind == 'g':
            num(wsh, r, c, '=0' if i == 0 else f'=IFERROR({CL(c)}{RH["売上高"]}/{CL(c - 1)}{RH["売上高"]},0)', fmt=FMT_P)
        else:
            base = kind.split(':')[1]
            num(wsh, r, c, f'=IFERROR({CL(c)}{RH[base]}/{CL(c)}{RH["売上高"]},0)', fmt=FMT_P)
    RH[key] = r
    r += 1
label(wsh, r, '【調整項目】', bold=True)
r += 1
label(wsh, r, '　退職給与（臨時的支出）')
num(wsh, r, 3, '=0')
for i in range(3):
    num(wsh, r, 4 + i, f"={q('修正PL')}!{CL(4 + i * 3)}{RP['adj_退職給与']}")
RH['adj_退職'] = r
r += 1
label(wsh, r, '　役員報酬')
num(wsh, r, 3, '=0')
for i in range(3):
    num(wsh, r, 4 + i, f"={q('修正PL')}!{CL(4 + i * 3)}{RP['adj_役員報酬']}")
RH['adj_役員'] = r
r += 1
label(wsh, r, '調整後EBITDA', bold=True, fill=LBLUE)
for i, c in enumerate(PCOLS):
    num(wsh, r, c, f'={CL(c)}{RH["EBITDA"]}+{CL(c)}{RH["adj_退職"]}+{CL(c)}{RH["adj_役員"]}', bold=True, fill=LBLUE)
RH['調整後EBITDA'] = r
r += 1
label(wsh, r, '調整後EBITDA 直近3期平均', bold=True)
num(wsh, r, 6, f'=ROUND(AVERAGE(D{RH["調整後EBITDA"]}:F{RH["調整後EBITDA"]}),0)', bold=True)
RH['avg3'] = r
r += 1
label(wsh, r, '調整後EBITDA 直近2期平均')
num(wsh, r, 6, f'=ROUND(AVERAGE(E{RH["調整後EBITDA"]}:F{RH["調整後EBITDA"]}),0)')
RH['avg2'] = r
r += 1
label(wsh, r, '検算（26/3期 調整後EBITDA＝修正PLの修正後EBITDA）')
wsh.cell(row=r, column=6, value=f"=EXACT(F{RH['調整後EBITDA']},{q('修正PL')}!K{RP['EBITDA']})").font = Font(name=NUM, size=9)
RH['check'] = r
note(wsh, r + 2, '※調整方針は評価書（修正PL）と同一。23/3期は評価の対象期間外のため調整なし。')

# ---------- BSハイライト ----------
wsg = ws_new('BSハイライト', 'BSハイライト　単位：千円', {'B': 34, 'C': 14, 'D': 14, 'E': 14, 'F': 14, 'G': 50})
RG = REF['BSハイライト']
hdr(wsg, 4, ['科目'] + [f'={q(BSD)}!{CL(c)}4' for c in PCOLS] + ['備考'])
r = 5
bs_lines = [('現預金（保険積立除く）', lambda L: '+'.join(f'{q(BSD)}!{L}{R[k]}' for k in ['現金', '当座預金', '普通預金', '定期預金', '定期積金'])),
            ('保険積立金', lambda L: f'{q(BSD)}!{L}{R["諸掛金（生命保険積立）"]}'),
            ('売掛金', lambda L: f'{q(BSD)}!{L}{R["sub_quick"]}'),
            ('棚卸資産', lambda L: f'{q(BSD)}!{L}{R["sub_inv"]}'),
            ('その他流動資産', lambda L: f'{q(BSD)}!{L}{R["sub_oca"]}'),
            ('流動資産合計', lambda L: f'{q(BSD)}!{L}{R["流動資産合計"]}'),
            ('建物・構築物', lambda L: f'{q(BSD)}!{L}{R["建物"]}+{q(BSD)}!{L}{R["構築物"]}'),
            ('機械・車両・器具', lambda L: f'{q(BSD)}!{L}{R["機械"]}+{q(BSD)}!{L}{R["車両運搬具"]}+{q(BSD)}!{L}{R["器具備品"]}'),
            ('土地', lambda L: f'{q(BSD)}!{L}{R["土地"]}'),
            ('無形固定資産・投資等', lambda L: f'{q(BSD)}!{L}{R["sub_intan"]}+{q(BSD)}!{L}{R["sub_invest"]}'),
            ('固定資産合計', lambda L: f'{q(BSD)}!{L}{R["固定資産合計"]}'),
            ('資産合計', lambda L: f'{q(BSD)}!{L}{R["資産合計"]}'),
            ('支払手形・買掛金', lambda L: f'{q(BSC)}!{L}{R2["支払手形"]}+{q(BSC)}!{L}{R2["買掛金"]}'),
            ('借入金（短期・長期）', lambda L: f'{q(BSC)}!{L}{R2["有利子負債"]}'),
            ('その他負債', lambda L: f'{q(BSC)}!{L}{R2["負債合計"]}-{q(BSC)}!{L}{R2["有利子負債"]}-{q(BSC)}!{L}{R2["支払手形"]}-{q(BSC)}!{L}{R2["買掛金"]}'),
            ('負債合計', lambda L: f'{q(BSC)}!{L}{R2["負債合計"]}'),
            ('純資産合計', lambda L: f'{q(BSC)}!{L}{R2["純資産合計"]}'),
            ('負債・純資産合計', lambda L: f'{q(BSC)}!{L}{R2["負債純資産合計"]}')]
for nm, fn in bs_lines:
    b = nm.endswith('合計')
    label(wsg, r, nm, bold=b, fill=GRAY if b else None)
    for c in PCOLS:
        num(wsg, r, c, '=' + fn(CL(c)), bold=b, fill=GRAY if b else None)
    RG[nm] = r
    r += 1
label(wsg, r, '自己資本比率')
for c in PCOLS:
    num(wsg, r, c, f'=IFERROR({CL(c)}{RG["純資産合計"]}/{CL(c)}{RG["資産合計"]},0)', fmt=FMT_P)
RG['自己資本比率'] = r
r += 1
label(wsg, r, 'ネット有利子負債（借入金－現預金）')
for c in PCOLS:
    num(wsg, r, c, f'={CL(c)}{RG["借入金（短期・長期）"]}-{CL(c)}{RG["現預金（保険積立除く）"]}')
RG['netdebt'] = r
r += 1
label(wsg, r, '時価純資産（評価基準日）', bold=True)
num(wsg, r, 6, f"={q('修正BS(貸方) ')}!E{RC['純資産']}", bold=True)
RG['時価純資産'] = r

# ---------- 取引先 ----------
wsc2 = ws_new('取引先', '主要販売先（期末売掛金残高ベース）　単位：千円', {'B': 30, 'C': 13, 'D': 8, 'E': 13, 'F': 8, 'G': 13, 'H': 8, 'I': 13, 'J': 8, 'K': 40})
RT = REF['取引先']
for i, c in enumerate(VC):
    wsc2.merge_cells(start_row=4, start_column=c, end_row=4, end_column=c + 1)
    hdr(wsc2, 4, [f'={q(BSD)}!{CL(PCOLS[i])}4'], c)
hdr(wsc2, 4, ['取引先'], 2)
hdr(wsc2, 4, ['備考'], 11)
r = 5
first = r
for nm, vals in AR_BY_CUSTOMER:
    label(wsc2, r, nm)
    for i, c in enumerate(VC):
        num(wsc2, r, c, vals[i])
    RT[nm] = r
    r += 1
label(wsc2, r, 'その他（複数社）')
for i, c in enumerate(VC):
    num(wsc2, r, c, f'={q(BSD)}!{CL(PCOLS[i])}{R["売掛金"]}-SUM({CL(c)}{first}:{CL(c)}{r - 1})')
RT['その他'] = r
r += 1
label(wsc2, r, '期末売掛金 合計', bold=True, fill=LBLUE)
for i, c in enumerate(VC):
    num(wsc2, r, c, f'=SUM({CL(c)}{first}:{CL(c)}{r - 1})', bold=True, fill=LBLUE)
RT['合計'] = r
for row in range(first, r + 1):
    for i, c in enumerate(VC):
        cc = wsc2.cell(row=row, column=c + 1, value=f'=IFERROR({CL(c)}{row}/{CL(c)}${r},0)')
        cc.number_format = FMT_P
        cc.font = Font(name=NUM, size=8)
note(wsc2, RT['(株)小野田商店'], '代表者と同姓。関連当事者か要確認', col=11)
note(wsc2, RT['(株)日本橋冷凍手島商店'], '資材の仕入先でもある', col=11)
note(wsc2, r + 2, '※取引先別の売上高データは未受領（要確認）。期末売掛金残高（月末締め翌月末回収）を取引規模の代理指標として表示。')
note(wsc2, r + 3, '※その他は売掛金残高との差額。67期内訳書は上位先合計が残高を1,723,680円超過しており、(株)小野田商店の重複計上と推察（要確認）。')

# ---------- 仕入先 ----------
wsu = ws_new('仕入先', '主要仕入先・主要経費支払先（2026年3月期末残高）　単位：千円', {'B': 36, 'C': 26, 'D': 14, 'E': 8, 'F': 40})
RU = REF['仕入先']
hdr(wsu, 4, ['仕入先（買掛金）', '内容', '期末残高', '構成比', '備考'])
r = 5
first = r
for nm, kind, v in AP_BY_SUPPLIER_67:
    text(wsu, r, 2, nm)
    text(wsu, r, 3, kind)
    num(wsu, r, 4, v)
    r += 1
label(wsu, r, '買掛金 合計', bold=True, fill=LBLUE)
num(wsu, r, 4, f'=SUM(D{first}:D{r - 1})', bold=True, fill=LBLUE)
RU['ap_total'] = r
for row in range(first, r + 1):
    cc = wsu.cell(row=row, column=5, value=f'=IFERROR(D{row}/D${r},0)')
    cc.number_format = FMT_P
    cc.font = Font(name=NUM, size=8)
wsu.cell(row=r, column=6, value=f'=IF(D{r}={q(BSC)}!F{R2["買掛金"]},"BS残高と一致","BS残高と不一致")').font = Font(name=MED, size=8)
r += 2
hdr(wsu, r, ['主な経費支払先（未払金）', '内容', '期末残高', '', '備考'])
r += 1
RU['payees'] = r
for nm, kind, v in OTHER_PAYEES_67:
    text(wsu, r, 2, nm)
    text(wsu, r, 3, kind)
    num(wsu, r, 4, v)
    r += 1
note(wsu, r + 1, '※原材料（ロール・段ボール・ビニール袋等の包装資材）の年間仕入高は77,113千円（67期）。電力（動力費 製造33,685千円）、配送運賃（27,570千円）が主要な費用。')

# ---------- 外注先 ----------
wso = ws_new('外注先', '外注先', {'B': 100})
text(wso, 4, 2, '外注作業費 11,434,787円（67期）。外注先の内訳は未受領（要確認）。人材派遣（(株)スタッフサービス、(株)ヒューマンアイ、ランスタッド(株)）の利用あり。')

# ---------- 株主 / 取締役 ----------
wss = ws_new('株主', '株主', {'B': 18, 'C': 14, 'D': 10, 'E': 14, 'F': 60})
hdr(wss, 4, ['氏名', '保有株式数', '比率', '役職', '備考'])
text(wss, 5, 2, '小野田 洋')
num(wss, 5, 3, 22000, fmt='#,##0"株"')
num(wss, 5, 4, 1, fmt='0%')
text(wss, 5, 5, '代表取締役')
text(wss, 5, 6, '発行済株式数：登記22,000株（2012年4月30日変更）／第67期株主資本等変動計算書注記200株／法人税申告書別表二220株と相違 → 株主名簿で要確認')
text(wss, 7, 2, '株券発行会社（定款の定め）。株式譲渡には株券の交付を要するため、株券の現存状況を要確認。譲渡制限あり（株主総会の承認）。', wrap=True)
wss.merge_cells('B7:F8')
wsdr = ws_new('取締役', '取締役', {'B': 14, 'C': 14, 'D': 14, 'E': 14, 'F': 20, 'G': 50})
hdr(wsdr, 4, ['氏名', '役職', '就任（重任）', '役員報酬（年）', 'M&A後の進退', '備考'])
text(wsdr, 5, 2, '小野田 洋')
text(wsdr, 5, 3, '代表取締役')
text(wsdr, 5, 4, '2020年5月25日重任')
num(wsdr, 5, 5, 7200000)
text(wsdr, 5, 6, '引継ぎ後に退任予定')
text(wsdr, 5, 7, '引継ぎ期間・条件は譲受企業と協議。年齢・経歴は要確認')

# ---------- 従業員 ----------
wsw = ws_new('従業員', '従業員・組織', {'B': 26, 'C': 12, 'D': 60})
hdr(wsw, 4, ['区分', '人数', '備考'])
emp = [('常勤役員', 1, '代表取締役'), ('従業員（越谷工場）', 19, '工場長（鈴木直則氏）が本社・工場・松戸営業所の責任者を兼務'),
       ('従業員（松戸営業所）', 5, ''), ('合計', 25, '2026年3月期末（法人事業概況説明書）。月別従事員数は23〜26名')]
for i, (k, n_, rem) in enumerate(emp):
    text(wsw, 5 + i, 2, k, bold=(k == '合計'))
    num(wsw, 5 + i, 3, n_, fmt='0"名"')
    text(wsw, 5 + i, 4, rem)
text(wsw, 10, 2, '組織（受領資料ベース）：代表取締役／工場長（鈴木）／財務・経理・総務（千葉・関根、一部外部委託）／営業（森谷）／越谷工場：製造（揚氷：林、加工：増渕）・配送（駒崎）・品質検査（社長）／松戸営業所（森）', wrap=True)
wsw.merge_cells('B10:D12')
text(wsw, 13, 2, '従業員一覧（年齢・勤続・資格・雇用形態・給与）は未受領（要確認）。人件費（67期）：役員報酬7,200千円、給料手当12,920千円、賃金手当78,160千円。', wrap=True)
wsw.merge_cells('B13:D14')

# ---------- 拠点 / 不動産 ----------
wsx = ws_new('拠点(本社等)', '拠点', {'B': 20, 'C': 34, 'D': 16, 'E': 14, 'F': 10, 'G': 10, 'H': 40})
hdr(wsx, 4, ['拠点', '所在地', '内容', '売上高（67期）', '従業員', '延床（㎡）', '備考'])
for i, (nm, addr, cont, s_, inv, e_, a_) in enumerate(SITE_SALES_67):
    r = 5 + i
    text(wsx, r, 2, nm)
    text(wsx, r, 3, addr)
    text(wsx, r, 4, cont)
    num(wsx, r, 5, s_)
    num(wsx, r, 6, e_, fmt='0"名"')
    num(wsx, r, 7, a_, fmt='#,##0')
text(wsx, 5, 8, '登記上の本店。代表者住所と同一（実務上の拠点は越谷工場）')
text(wsx, 6, 8, '自社所有（土地・建物）。製氷・加工・保管')
text(wsx, 7, 8, '賃借（年1,200千円）')
wsy2 = ws_new('不動産(賃借) ', '不動産（賃借）', {'B': 20, 'C': 34, 'D': 24, 'E': 14, 'F': 40})
hdr(wsy2, 4, ['区分', '所在地', '貸主', '年間賃料', '備考'])
for i, (k, addr, lessor, rent, rem) in enumerate(REAL_ESTATE_LEASED):
    r = 5 + i
    text(wsy2, r, 2, k)
    text(wsy2, r, 3, addr)
    text(wsy2, r, 4, lessor)
    num(wsy2, r, 5, rent)
    text(wsy2, r, 6, rem)
label(wsy2, 7, '合計', bold=True)
num(wsy2, 7, 5, '=SUM(E5:E6)', bold=True)
wsz = ws_new('不動産(所有)', '不動産（所有）', {'B': 8, 'C': 12, 'D': 30, 'E': 12, 'F': 15, 'G': 15, 'H': 50})
hdr(wsz, 4, ['種類', '区分', '所在地', '面積（㎡）', '簿価', '固定資産税評価額', '備考'])
for i, (kind, area_, addr, m2, bv, ev, rem) in enumerate(REAL_ESTATE_OWNED):
    r = 5 + i
    text(wsz, r, 2, kind)
    text(wsz, r, 3, area_)
    text(wsz, r, 4, addr)
    num(wsz, r, 5, m2, fmt='#,##0.00')
    num(wsz, r, 6, f'={q(BSD)}!F{R[kind]}')
    num(wsz, r, 7, ev)
    text(wsz, r, 8, rem)
label(wsz, 7, '合計', bold=True)
num(wsz, 7, 6, '=SUM(F5:F6)', bold=True)
num(wsz, 7, 7, '=SUM(G5:G6)', bold=True)

# ---------- 設備 ----------
wse2 = ws_new('設備', '保有設備', {'B': 18, 'C': 10, 'D': 16, 'E': 15, 'F': 15, 'G': 15, 'H': 40})
hdr(wse2, 4, ['主要設備', '台数', '能力', '', '', '', '出所'])
for i, (k, n_, cap) in enumerate(EQUIPMENT_MAIN):
    text(wse2, 5 + i, 2, k)
    text(wse2, 5 + i, 3, n_)
    text(wse2, 5 + i, 4, cap)
    text(wse2, 5 + i, 8, '法人事業概況説明書（67期）')
hdr(wse2, 9, ['償却資産（科目）', '件数', '取得価額', '期末簿価', '当期償却額', '償却不足額', '備考'])
for i, (nm, n_, acq, bv, dep, short, acc) in enumerate(DEP_67):
    r = 10 + i
    text(wse2, r, 2, nm)
    num(wse2, r, 3, n_, fmt='0"件"')
    num(wse2, r, 4, acq)
    num(wse2, r, 5, bv)
    num(wse2, r, 6, dep)
    num(wse2, r, 7, short)
label(wse2, 15, '合計', bold=True)
for j in range(3, 8):
    num(wse2, 15, j, f'=SUM({CL(j)}10:{CL(j)}14)', bold=True, fmt=FMT if j > 3 else '0"件"')
note(wse2, 17, '※減価償却計算書（185件）の明細は別途データルームで開示予定。主要設備の更新時期・修繕計画は要確認。')

# ---------- 借入金 ----------
wsl2 = ws_new('借入金', '借入金の状況（2026年3月期末）', {'B': 32, 'C': 15, 'D': 15, 'E': 44})
RLN = REF['借入金']
hdr(wsl2, 4, ['借入先', '期末残高', '支払利息（年）', '担保・備考'])
r = 5
for nm, bal, intr, col_ in LOANS_67:
    text(wsl2, r, 2, nm)
    num(wsl2, r, 3, bal)
    num(wsl2, r, 4, intr)
    text(wsl2, r, 5, col_)
    r += 1
label(wsl2, r, '合計', bold=True, fill=LBLUE)
num(wsl2, r, 3, '=SUM(C5:C%d)' % (r - 1), bold=True, fill=LBLUE)
num(wsl2, r, 4, '=SUM(D5:D%d)' % (r - 1), bold=True, fill=LBLUE)
RLN['合計'] = r
wsl2.cell(row=r + 1, column=3, value=f'=EXACT(C{r},{q(BSC)}!F{R2["有利子負債"]})').font = Font(name=NUM, size=9)
wsl2.cell(row=r + 1, column=4, value=f'=EXACT(D{r},PL!I{R5["支払利息"]})').font = Font(name=NUM, size=9)
label(wsl2, r + 1, '検算（BS有利子負債／PL支払利息との一致）')
note(wsl2, r + 3, '※代表者は金融機関借入について連帯保証・個人所有土地の担保提供を行っている（受領資料）。M&A実行後の保証解除・担保差替えが譲渡条件。')

# ---------- 月次売上 ----------
wsm = ws_new('月次売上', '月別売上高（2026年3月期）　単位：千円', {'B': 14} | {CL(c): 9 for c in range(3, 16)})
hdr(wsm, 4, ['項目'] + MONTHLY_67['month'] + ['合計'])
RM = REF['月次売上']
for i, k in enumerate(['凍氷', '保管料', '原材料仕入', '外注費', '人件費', '従事員数']):
    r = 5 + i
    label(wsm, r, k if k != '凍氷' else '売上（凍氷）')
    for j, v in enumerate(MONTHLY_67[k]):
        num(wsm, r, 3 + j, v, fmt='#,##0')
    if k != '従事員数':
        num(wsm, r, 15, f'=SUM(C{r}:N{r})', bold=True, fmt='#,##0')
    RM[k] = r
note(wsm, 12, '※出所：第67期 法人事業概況説明書。7〜8月の盛夏期に年間売上の約3割が集中する季節性の高い事業。')

# ---------- 許認可 / 沿革 ----------
wsp2 = ws_new('許認可', '許認可', {'B': 30, 'C': 20, 'D': 16, 'E': 50})
hdr(wsp2, 4, ['名称', '番号', '許可日', '備考'])
for i, (nm, no, d, rem) in enumerate(LICENSES):
    for j, v in enumerate([nm, no, d, rem]):
        text(wsp2, 5 + i, 2 + j, v)
HIST = [('1960年（昭和35年）3月', '越ヶ谷製氷冷蔵株式会社を設立。越谷工場で操業開始'),
        ('1964年（昭和39年）', '越谷工場に純水装置（イオン交換方式）を導入'),
        ('1972年（昭和47年）', '本社を越谷市蒲生寿町へ移転'),
        ('1990年（平成2年）', '逆浸透膜（RO）ろ過装置を導入'),
        ('1991年（平成3年）3月', '越谷工場（現工場棟）を建設'),
        ('2008年（平成20年）', '越谷工場 外壁塗装工事'),
        ('2012年（平成24年）4月', '増資により資本金11,000千円'),
        ('2020年（令和2年）5月', '株式の譲渡制限に関する規定を変更'),
        ('2026年（令和8年）5月', '減資により資本金10,000千円')]
wsh2 = ws_new('沿革', '沿革', {'B': 24, 'C': 70, 'D': 30})
hdr(wsh2, 4, ['年月', '主な出来事', '出所'])
for i, (d, e) in enumerate(HIST):
    text(wsh2, 5 + i, 2, d)
    text(wsh2, 5 + i, 3, e)
    text(wsh2, 5 + i, 4, '公式サイト・Web掲載情報' if i in (0, 1, 2, 3) else ('減価償却計算書' if i in (4, 5) else '登記'), size=8)

# ---------- 案件サマリー（概要書・ノンネームへの貼付元） ----------
wsa = ws_new('案件サマリー', '案件サマリー（概要書「基本情報」「財務ハイライト」の貼付元）', {'B': 18, 'C': 15, 'D': 15, 'E': 15, 'F': 15, 'G': 4, 'H': 16, 'I': 50})
RA = REF['案件サマリー']
hdr(wsa, 4, ['会社概要'], 8)
wsa.merge_cells('H4:I4')
summary_items = ['会社名', '本店所在地', '工場・営業所', '設立', '資本金', '代表者', '株主', '従業員数', '事業内容', '主要取引先', '主要設備', '売上高', '譲渡理由', '譲渡スキーム']
for i, k in enumerate(summary_items):
    r = 5 + i
    text(wsa, r, 8, k, bold=True, fill=GRAY)
    text(wsa, r, 9, f"='会社概要'!C{RK[k]}")
    RA['co_' + k] = r
hdr(wsa, 4, ['財務サマリー（千円）'] + [f'={q(BSD)}!{CL(c)}4' for c in PCOLS], 2)
fin = [('売上高', f"'PLハイライト'!{{c}}{RH['売上高']}"), ('売上総利益', f"'PLハイライト'!{{c}}{RH['売上総利益']}"),
       ('営業利益', f"'PLハイライト'!{{c}}{RH['営業利益']}"), ('経常利益', f"'PLハイライト'!{{c}}{RH['経常利益']}"),
       ('当期純利益', f"'PLハイライト'!{{c}}{RH['当期純利益']}"), ('減価償却費', f"'PLハイライト'!{{c}}{RH['減価償却費']}"),
       ('調整後EBITDA', f"'PLハイライト'!{{c}}{RH['調整後EBITDA']}"), ('総資産', f"'BSハイライト'!{{c}}{RG['資産合計']}"),
       ('純資産', f"'BSハイライト'!{{c}}{RG['純資産合計']}"), ('借入金', f"'BSハイライト'!{{c}}{RG['借入金（短期・長期）']}"),
       ('現預金', f"'BSハイライト'!{{c}}{RG['現預金（保険積立除く）']}")]
for i, (k, f) in enumerate(fin):
    r = 5 + i
    label(wsa, r, k, bold=k in ('売上高', '営業利益', '調整後EBITDA'))
    for c in PCOLS:
        num(wsa, r, c, '=' + f.format(c=CL(c)), bold=k in ('売上高', '営業利益', '調整後EBITDA'))
    RA['fin_' + k] = r
r = 5 + len(fin)
label(wsa, r, '調整後EBITDA 直近3期平均', bold=True, fill=LBLUE)
num(wsa, r, 6, f"='PLハイライト'!F{RH['avg3']}", bold=True, fill=LBLUE)
RA['avg3'] = r

# ---------- その他（要確認事項） ----------
wsq = ws_new('その他', '要確認事項・ネクストステップ（ヒアリング項目一覧の原本）', {'B': 6, 'C': 16, 'D': 90, 'E': 10})
hdr(wsq, 4, ['No.', '分野', '確認事項', '優先度'])
QS = [
    ('株主・資本', '発行済株式数の確定（登記22,000株／決算注記200株／別表二220株）と株主名簿・株券の現存状況（株券発行会社）', '高'),
    ('株主・資本', '2026年5月の減資（11,000千円→10,000千円）の目的・手続（欠損填補か払戻しか）', '中'),
    ('譲渡条件', '代表者の連帯保証・個人所有土地の担保提供（栃木銀行・商工中金）の解除方法、代表者借入金14,000千円の取扱い', '高'),
    ('譲渡条件', '代表者の引継ぎ期間・退任後の関与形態、代表者所有駐車場の賃貸借の継続', '高'),
    ('財務', '建物の会計上の償却停止の経緯（償却不足額 当期21,745千円）と建物簿価（BS259,242千円／内訳書280,673千円）の差異', '高'),
    ('財務', '定期預金残高の内訳差異6,800千円（預貯金内訳19,601千円／BS26,401千円）、担保差入の有無', '高'),
    ('財務', '売掛金内訳の合計超過1,723,680円（(株)小野田商店）と同社との関係（関連当事者か）', '高'),
    ('財務', '退職給与6,178千円（67期）の対象者・退職金規程・中退共加入状況、今後の退職予定者', '高'),
    ('財務', '生命保険（大同生命・第一生命）の解約返戻金額・被保険者', '中'),
    ('財務', '65期の雑収入8,482千円・雑損5,192千円の内容', '中'),
    ('財務', '支払手数料（66期3,156千円・67期1,765千円）の内容', '中'),
    ('財務', '賞与の支給時期・算定期間（賞与引当の要否）', '中'),
    ('事業', '取引先別・商品別（角氷／加工氷／袋氷／ドライアイス／保管）の売上高・数量・単価', '高'),
    ('事業', '電力契約（新電力：バンブーパワートレーディング）の条件と電気料金の感応度', '中'),
    ('事業', '主要設備（製氷設備・冷凍機）の更新時期・修繕計画、缶グリッド部品更新（大森工業 2,633千円）の位置づけ', '中'),
    ('組織', '従業員一覧（年齢・勤続・資格・雇用形態・給与）、工場長の継続意向', '高'),
    ('許認可', '食品営業許可（氷雪販売業、2021年1月15日・5年更新）の更新状況', '高'),
    ('不動産', '工場土地・建物の最新固定資産税評価額、建物面積（2,093.87㎡／2,065.00㎡）の差異、代表者個人所有土地の明細', '中'),
]
r = 5
for i, (area_, qtext, pr) in enumerate(QS):
    text(wsq, r, 2, i + 1)
    text(wsq, r, 3, area_)
    text(wsq, r, 4, qtext)
    text(wsq, r, 5, pr)
    r += 1
r += 1
hdr(wsq, r, ['', 'プロセス関連（社内確認事項）', '', ''])
r += 1
for t in ['※本項は社内確認事項であり、買手候補への開示資料には記載しない。',
          '・本件は先行して他の仲介会社が企業概要書を作成済み（2025年10月時点）。当該資料の買手候補への配布範囲・先行打診先の有無を確認し、重複打診を避ける。',
          '・第64〜66期の財務数値は当該資料から転記。原本（決算書）の受領と照合を行う。',
          '・Web掲載画像（公式サイト king-ice.com）の概要書への使用は売主了承済み（担当者確認）。',
          '・必要資料：決算書（64〜66期原本）、株主名簿、定款、従業員一覧、取引先別売上、設備台帳、賃貸借契約書、保険証券、退職金規程']:
    note(wsq, r, t, col=3, size=9)
    r += 1

# 期ヘッダー（原本）
for i, c in enumerate(PCOLS):
    ws.cell(row=4, column=c, value=PERIODS[i])

# シート順：セクション順に整列
order = ['ルール', '財務入力→', BSD, BSC, 'PL', 'CR', 'SGA', 'VR→', '株式価値 ', '修正BS(借方)', '修正BS(貸方) ',
         '売掛・棚卸', '減価償却', '土地', '保険', '有価証券', '未払人件費・賞与', '退職金', '税効果', 'ネットキャッシュ',
         '修正PL', '修正SGA', '修正CR', '年買法', 'EV EBITDAマルチプル', '参照企業', 'IM→', '案件サマリー', '会社概要',
         'PLハイライト', 'BSハイライト', '取引先', '仕入先', '外注先', '株主', '取締役', '従業員', '拠点(本社等)', '不動産(賃借) ',
         '不動産(所有)', '設備', '借入金', '月次売上', '許認可', '沿革', 'その他']
wb._sheets = [wb[n] for n in order]
for w_ in wb.worksheets:
    w_.sheet_view.zoomScale = 90
wb.save(OUT)

import json
REF_OUT = sys.argv[2] if len(sys.argv) > 2 else OUT + '.ref.json'
json.dump(REF, open(REF_OUT, 'w'), ensure_ascii=False, indent=1)
print('saved', OUT)

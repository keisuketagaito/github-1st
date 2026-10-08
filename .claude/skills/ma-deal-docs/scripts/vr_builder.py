# -*- coding: utf-8 -*-
"""株式価値評価書（Valuation Report）ハウスフォーマットの部品ライブラリ。

当社テンプレート `assets/株式価値評価書_テンプレート.pptx` を複製したうえで、
本モジュールと `im_builder` の部品だけでページを組み立てる。
座標・色・フォントはカナリア案件（Valuation Report 260729）から採取した実測値。

マスター・レイアウト・テーマは企業概要書テンプレートと同一であるため、
ページ枠（タイトル／字幕／セクション名／注記）や表は `im_builder` の部品をそのまま使う。
本モジュールは評価書に固有の部品（Point帯・前提条件表・採用アプローチ表・
試算結果レンジ・年買法／マルチプル法の計算ブロック）を提供する。

使い方::

    from im_builder import *
    from vr_builder import *

    prs = open_vr_template()
    keep_slides(prs, [...])
    strip_red_banners(prs); fix_broken_pictures(prs)
    set_cover_vr(prs.slides[0], '株式会社○○', '2026年10月7日')
"""
from __future__ import annotations

import os
import shutil
import tempfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from im_builder import (  # noqa: F401  （評価書でもそのまま使う）
    NAVY, BLUE, PALE, GREY, BODY, NOTE, WHITE, RED, DARK, JP, NUM,
    rect, textbox, rich, set_text, set_title, set_lead, set_section,
    note, house_table, comment_box, clear_body, keep_slides, delete_slide,
    duplicate_slide, strip_red_banners, fix_broken_pictures,
    fmt, pct, k, _ph, _apply_font, no_bullet,
)

_TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                         '..', 'assets', '株式価値評価書_テンプレート.pptx')

# ----------------------------------------------------------------- 評価書の定数
# 評価書の区画見出しは概要書より薄く・文字は大きい（実測）
BAR_H = 0.27
BAR_SIZE = 16

# 標準の区画幅
W_FULL = 10.71          # 全幅バー（x=0.49 起点）
W_HALF = 3.48           # 1/3幅バー
X0 = 0.49

TAX_RATE = 0.34         # 実効税率
ILLIQ_DISCOUNT = 0.30   # 非流動性ディスカウント
WEIGHTS = (0.2, 0.3, 0.5)   # 前々期・前期・直近期


def open_vr_template(path: str | None = None) -> Presentation:
    """評価書テンプレートを一時ファイルへ複製して開く（原本は書き換えない）。"""
    src = path or os.path.normpath(_TEMPLATE)
    tmp = tempfile.NamedTemporaryFile(suffix='.pptx', delete=False)
    tmp.close()
    shutil.copyfile(src, tmp.name)
    return Presentation(tmp.name)


# ----------------------------------------------------------------- 表紙
def set_cover_vr(slide, company: str, date_text: str):
    """評価書の表紙。タイトルは "Valuation Report" 固定、社名と算定日を差し替える。"""
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if '御中' in t:
            set_text(sh, f'{company}　御中', font=JP)
        elif '年' in t and '月' in t and '日' in t and len(t) < 20:
            set_text(sh, date_text, font=JP)
    return slide


# ----------------------------------------------------------------- 区画見出し
def vr_bar(slide, x, y, w, text, h=BAR_H, size=BAR_SIZE):
    """評価書の区画見出し（紺帯）。概要書より薄く、文字は16pt。"""
    sp = rect(slide, x, y, w, h, fill=NAVY)
    set_text(sp, text, size=size, bold=True, color=WHITE, font=JP, align=PP_ALIGN.LEFT)
    sp.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    sp.text_frame.margin_left = Inches(0.12)
    return y + h


def point_box(slide, x, y, w, lines, h=0.79, label='Point', label_w=1.57):
    """『Point』帯。評価手法の補足や注意喚起に使う（カナリア評価書P.6/P.20）。"""
    rect(slide, x, y, w, h, fill=PALE, line=GREY, lw=0.75)
    rect(slide, x, y, label_w, h, fill=NAVY)
    textbox(slide, x, y + h / 2 - 0.13, label_w, 0.26, label, size=14, bold=True,
            color=WHITE, font=JP, align=PP_ALIGN.CENTER)
    textbox(slide, x + label_w + 0.18, y + 0.10, w - label_w - 0.32, h - 0.20,
            lines, size=10.5, color=BODY, font=JP, line_sp=1.22, space=2,
            anchor=MSO_ANCHOR.MIDDLE)
    return y + h


def vr_comment(slide, x, y, w, lines, h=0.71, label='Comment', label_w=1.57, size=10):
    """横長の Comment 帯（カナリア評価書P.14/P.20）。箇条書きを横に並べる。"""
    rect(slide, x, y, w, h, fill=WHITE, line=NAVY, lw=1.0)
    textbox(slide, x + 0.12, y + h / 2 - 0.13, label_w, 0.26, label, size=14,
            bold=True, color=NAVY, font=JP)
    if isinstance(lines, (list, tuple)):
        n = len(lines)
        colw = (w - label_w - 0.24) / n
        for i, t in enumerate(lines):
            textbox(slide, x + label_w + 0.12 + i * colw, y + 0.10, colw - 0.12,
                    h - 0.20, t, size=size, color=BODY, font=JP, line_sp=1.18,
                    anchor=MSO_ANCHOR.MIDDLE)
    else:
        textbox(slide, x + label_w + 0.12, y + 0.10, w - label_w - 0.24, h - 0.20,
                lines, size=size, color=BODY, font=JP, line_sp=1.18, space=2,
                anchor=MSO_ANCHOR.MIDDLE)
    return y + h


# ----------------------------------------------------------------- 試算結果
def value_range(slide, x, y, w, label, lo, hi, note_text=None, h=0.44,
                adopted=True):
    """『想定レンジ ●● 〜 ●●（千円）』のバンド（カナリア評価書P.7）。"""
    textbox(slide, x, y - 0.26, w, 0.22, '■ ' + label, size=11, bold=True,
            color=NAVY, font=JP)
    if not adopted:
        rect(slide, x, y, w, h, fill=RGBColor(0xF2, 0xF4, 0xF5), line=GREY)
        textbox(slide, x, y, w, h, '※※※※※　不採用　※※※※※', size=12, bold=True,
                color=NOTE, font=JP, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        return y + h
    rect(slide, x, y, w, h, fill=PALE)
    rect(slide, x, y, 1.30, h, fill=NAVY)
    textbox(slide, x, y, 1.30, h, '想定レンジ', size=12, bold=True, color=WHITE,
            font=JP, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt = f'{fmt(lo)}　～　{fmt(hi)}（千円）'
    textbox(slide, x + 1.40, y, w - 1.50, h, txt, size=16, bold=True, color=NAVY,
            font=NUM, anchor=MSO_ANCHOR.MIDDLE)
    if note_text:
        textbox(slide, x, y + h + 0.04, w, 0.20, note_text, size=8, color=NOTE, font=JP)
    return y + h


def range_chart(slide, x, y, w, h, rows, unit='（単位：千円）'):
    """試算結果のレンジ図（横棒）。rows=[(ラベル, lo, hi), ...]  lo/hi は千円。"""
    rect(slide, x, y, w, h, fill=RGBColor(0xFB, 0xFC, 0xFD), line=GREY, lw=0.75)
    textbox(slide, x, y - 0.20, w, 0.18, unit, size=7.5, color=NOTE,
            align=PP_ALIGN.RIGHT, font=JP)
    vals = [v for _, lo, hi in rows for v in (lo, hi)]
    lo_all, hi_all = min(vals + [0]), max(vals + [0])
    span = (hi_all - lo_all) or 1
    pad = span * 0.08
    lo_all, hi_all = lo_all - pad, hi_all + pad
    span = hi_all - lo_all
    plot_x, plot_w = x + 1.70, w - 2.00
    rowh = (h - 0.70) / max(len(rows), 1)

    def px(v):
        return plot_x + plot_w * (v - lo_all) / span

    if lo_all < 0 < hi_all:
        rect(slide, px(0) - 0.008, y + 0.28, 0.016, h - 0.70, fill=RGBColor(0x9F, 0xAE, 0xB6))
    for i, (lab, lo, hi) in enumerate(rows):
        cy = y + 0.34 + i * rowh
        textbox(slide, x + 0.12, cy + rowh / 2 - 0.14, 1.52, 0.28, lab, size=8.5,
                color=BODY, font=JP, anchor=MSO_ANCHOR.MIDDLE)
        bx, bw = px(min(lo, hi)), abs(px(hi) - px(lo))
        rect(slide, bx, cy + rowh / 2 - 0.16, max(bw, 0.03), 0.32, fill=BLUE)
        textbox(slide, bx - 1.00, cy + rowh / 2 - 0.12, 0.96, 0.24, fmt(lo), size=8,
                color=NAVY, font=NUM, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        textbox(slide, bx + max(bw, 0.03) + 0.06, cy + rowh / 2 - 0.12, 1.00, 0.24,
                fmt(hi), size=8, color=NAVY, font=NUM, anchor=MSO_ANCHOR.MIDDLE)
    rect(slide, plot_x, y + h - 0.30, plot_w, 0.012, fill=GREY)
    return y + h


# ----------------------------------------------------------------- 計算ロジック
def weighted(v_old, v_mid, v_new, weights=WEIGHTS):
    """3期加重平均（前々期20％・前期30％・直近期50％）。"""
    w0, w1, w2 = weights
    return v_old * w0 + v_mid * w1 + v_new * w2


def goodwill_nenbai(base_op_after_tax, multiples=(1, 2, 3)):
    """年買法の営業権。基準営業利益（税考慮後）× 倍率。マイナスなら0。"""
    base = max(base_op_after_tax, 0)
    return {m: base * m for m in multiples}


def goodwill_excess(weighted_pretax, total_assets_fv, risk_premium=0.03,
                    jgb_yield=0.0, years=3, annuity_factor=None):
    """超過収益法の営業権（研修資料の手法）。

    期待利子率 ＝ 10年国債利回り ＋ リスクプレミアム（通常3.0％）
    期待利益   ＝ 時価総資産 × 期待利子率
    超過利益   ＝ 修正税引前純利益の3期加重平均 － 期待利益
    営業権     ＝ 超過利益 × 複利年金現価係数（持続年数・期待利子率）

    戻り値は (営業権, 期待利子率, 期待利益, 超過利益, 年金現価係数)。
    """
    rate = jgb_yield + risk_premium
    expected = total_assets_fv * rate
    excess = weighted_pretax - expected
    # 複利年金現価係数は係数表から読む運用のため、明示指定を優先する。
    # 研修資料（日本の川ちゃん）は期待利子率3.027％に対し係数2.7771を用いており、
    # 単純な現在価値計算（2.8271）とは一致しない。係数表の値を使うこと。
    annuity = (annuity_factor if annuity_factor is not None
               else sum(1 / (1 + rate) ** t for t in range(1, years + 1)))
    return max(excess, 0) * annuity, rate, expected, excess, annuity


def ev_multiple_value(base_ebitda, net_cash, multiples=(3, 4, 5),
                      discount=ILLIQ_DISCOUNT):
    """EV/EBITDAマルチプル法。戻り値は {倍率: (EV, ディスカウント, 株式価値)}。"""
    out = {}
    for m in multiples:
        ev = base_ebitda * m
        dc = -ev * discount
        out[m] = (ev, dc, ev + dc + net_cash)
    return out

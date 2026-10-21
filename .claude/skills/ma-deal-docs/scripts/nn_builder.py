# -*- coding: utf-8 -*-
"""ノンネームシート（買手候補への初期打診資料）の部品ライブラリ。

当社テンプレート `assets/ノンネームシート_テンプレート.pptx` を複製したうえで、
表のラベル（業種／エリア／ビジネス面 …）を手掛かりに値だけを差し替える。
python-pptx で白紙から組み立てることは禁止（他の資料と同じ鉄則）。

仕様（鶴岡中央青果 261015 から採取）::

    用紙      A4縦 8.27 × 11.69 インチ
    レイアウト 「タイトル スライド」1種のみ。ロゴ・Confidential・著作権表記を保持
    テーマ     企業概要書／評価書と同一（dk2=0B3041 / lt2=E3E7E9 / accent4=0889C9）
    構成      タイトル → 事業内容（4行表）→ 特徴・強み（3行表）
              → 補足事項（3行表）→ 本情報の取扱について（定型文・改変禁止）

使い方::

    from nn_builder import *
    prs = open_nn_template()
    fill_nonname(prs, {
        '業種': 'そば店（飲食業）',
        'エリア': '北関東',
        '売上高': '1億円未満',
        '従業員': '5名未満',
        'ビジネス面': ['…', '…'],
        '財務面':     ['…'],
        '組織面':     ['…'],
        '譲渡理由': '後継者不在',
        'スキーム': '株式100％譲渡',
        '希望譲渡価額': '応相談',
    })
    prs.save('【案件名】ノンネームシート.pptx')

**ノンネームである以上、社名が特定される情報は書かない。**
店名・商標名・代表者名・受賞名・市区町村名・固有の取引先名は載せない。
売上高・従業員数はレンジで示す（「1億円未満」「5名未満」）。
詳細は `references/nonname.md`。
"""
from __future__ import annotations

import copy
import os
import shutil
import tempfile

from pptx import Presentation
from pptx.oxml.ns import qn

_TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                         '..', 'assets', 'ノンネームシート_テンプレート.pptx')

# 本文テキストの差替え対象外（定型文・改変禁止）
FIXED_BLOCKS = ('本情報の取扱について',)


def open_nn_template(path: str | None = None) -> Presentation:
    """テンプレートを一時ファイルへ複製して開く（原本は書き換えない）。"""
    src = path or os.path.normpath(_TEMPLATE)
    tmp = tempfile.NamedTemporaryFile(suffix='.pptx', delete=False)
    tmp.close()
    shutil.copyfile(src, tmp.name)
    return Presentation(tmp.name)


def _norm(t: str) -> str:
    """ラベル照合用。全角スペース・半角スペース・改行を落とす（「売 上 高」対策）。"""
    return ''.join(str(t).split()).replace('　', '')


def set_cell_lines(cell, lines) -> None:
    """セルの文字を差し替える。1段落目の書式（箇条書き・フォント）を複製して使う。

    テンプレートの箇条書きは `buChar ▪` を段落プロパティに持つため、
    段落ごと複製しないと記号が消える。
    """
    if isinstance(lines, str):
        lines = [lines]
    tf = cell.text_frame
    proto = copy.deepcopy(tf.paragraphs[0]._p)
    # 複製元の余分なランを落とし、1ラン構成にしておく
    runs = proto.findall(qn('a:r'))
    for r in runs[1:]:
        proto.remove(r)
    txbody = tf._txBody
    for p in txbody.findall(qn('a:p')):
        txbody.remove(p)
    for ln in lines:
        p = copy.deepcopy(proto)
        r = p.find(qn('a:r'))
        if r is None:                      # 空セルを複製元にした場合の保険
            continue
        r.find(qn('a:t')).text = str(ln)
        txbody.append(p)


def iter_label_cells(prs):
    """全スライドの表から (ラベル文字列, 値セル) を列挙する。"""
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_table:
                continue
            t = sh.table
            for row in t.rows:
                cells = list(row.cells)
                if len(cells) >= 2:
                    yield _norm(cells[0].text), cells[1]


def fill_nonname(prs, data: dict) -> list:
    """ラベルを手掛かりに値を差し替える。戻り値は当たらなかったラベルの一覧。"""
    want = {_norm(k): v for k, v in data.items()}
    hit = set()
    for label, cell in iter_label_cells(prs):
        if label in want:
            set_cell_lines(cell, want[label])
            hit.add(label)
    return [k for k in want if k not in hit]


def set_title(prs, text: str) -> None:
    """表紙タイトル（既定は「M＆A案件 ノンネームシート」）。"""
    for sh in prs.slides[0].shapes:
        if sh.is_placeholder and sh.has_text_frame:
            set_cell_lines(sh, text)
            return


def find_table(prs, label):
    """先頭列に `label` を持つ表（のGraphicFrame）を返す。"""
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_table:
                continue
            for row in sh.table.rows:
                if _norm(list(row.cells)[0].text) == _norm(label):
                    return sh
    return None


def place_table(frame, y=None, heights=None, x=None):
    """表の位置と行高を調整する。

    テンプレートの行高は前案件の文字量に合わせて固定されているため、
    文章量を変えたら**必ず行高を取り直す**。放っておくと行が重なり、
    下の見出し（補足事項）まで押し潰す。
    """
    from pptx.util import Inches
    if x is not None:
        frame.left = Inches(x)
    if y is not None:
        frame.top = Inches(y)
    if heights:
        for r, h in zip(frame.table.rows, heights):
            r.height = Inches(h)
        frame.height = Inches(sum(heights))
    return frame


def move_text(prs, needle, y):
    """見出しテキストボックス（「事業内容」「特徴・強み」「補足事項」）を動かす。"""
    from pptx.util import Inches
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame and _norm(sh.text_frame.text) == _norm(needle):
                sh.top = Inches(y)
                return sh
    return None


def check_anonymity(prs, banned) -> list:
    """社名が割れる語が残っていないか検査する。納品前に必ず通す。

    banned には店名・商標・代表者名・市区町村名・受賞名・取引先名を入れる。
    """
    found = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []

        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == 6:
                    walk(sh.shapes)
                    continue
                if sh.has_text_frame:
                    texts.append(sh.text_frame.text)
                if sh.has_table:
                    for row in sh.table.rows:
                        for c in row.cells:
                            texts.append(c.text)
        walk(slide.shapes)
        blob = '\n'.join(texts)
        for w in banned:
            if w and w in blob:
                found.append((i, w))
    return found

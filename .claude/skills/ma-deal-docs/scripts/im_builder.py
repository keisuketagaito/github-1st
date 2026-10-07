# -*- coding: utf-8 -*-
"""企業概要書（IM）ハウスフォーマットの部品ライブラリ。

当社テンプレート `assets/企業概要書_テンプレート.pptx` を複製したうえで、
本モジュールの部品だけを使ってページを組み立てる。
座標・色・フォントは実案件（鶴岡中央青果 260923／鳥塚 260924）から採取した実測値。

使い方::

    from im_builder import *
    prs = open_template()                      # テンプレートを複製して開く
    keep_slides(prs, [1,2,3,4,5,6,7,8,9, ...]) # 使う頁だけ残す（1始まり）
    s = prs.slides[4]
    clear_body(s)                              # 枠（タイトル/字幕/セクション名）を残して中身を消す
    set_title(s, '基本情報'); set_lead(s, 'リード文'); set_section(s, 'エグゼクティブサマリー')
    section_bar(s, 0.34, 1.66, 6.41, '会社概要')
    house_table(s, 0.34, 2.05, 6.41, [1.35, 5.06], rows)
    note(s, '※ 出所…')
"""
from __future__ import annotations

import copy
import os
import shutil
import tempfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ----------------------------------------------------------------- 定数
W, H = 11.69, 8.27                     # A4横

NAVY  = RGBColor(0x0B, 0x30, 0x41)     # dk2／見出し帯・表ヘッダ
BLUE  = RGBColor(0x08, 0x89, 0xC9)     # accent4／小見出し・強調
PALE  = RGBColor(0xE1, 0xF3, 0xFB)     # accent2／淡い塗り
GREY  = RGBColor(0xE3, 0xE7, 0xE9)     # lt2／罫線・薄い塗り
BODY  = RGBColor(0x1A, 0x1A, 0x1A)     # 本文
NOTE  = RGBColor(0x5A, 0x66, 0x70)     # 注記
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD  = RGBColor(0x8C, 0x6D, 0x1F)     # 比較ページの系統色A
STEEL = RGBColor(0x45, 0x68, 0x8F)     # 比較ページの系統色B
RED   = RGBColor(0xC0, 0x3A, 0x2B)     # 減額・マイナス
DARK  = RGBColor(0x20, 0x20, 0x20)

JP  = '游ゴシック'
NUM = 'Arial'

# ページ枠（テンプレートのプレースホルダ実測値）
TITLE_POS   = (0.34, 0.20, 8.53, 0.37)
LEAD_POS    = (0.50, 0.91, 10.70, 0.26)
SECTION_POS = (9.15, 0.28, 2.17, 0.20)
NOTE_POS    = (0.34, 7.25, 10.85, 0.35)

L, R = 0.34, 0.34                      # 左右マージン
CW = W - L - R                         # コンテンツ幅 11.01
BODY_TOP = 1.51                        # 本文開始Y（字幕の下）
BODY_BOTTOM = 7.20                     # 本文下限（注記の上）

COL_L_W = 6.61                         # 2カラム時の左幅
COL_R_X = 7.06                         # 2カラム時の右X
COL_R_W = 4.29                         # 2カラム時の右幅

TABLE_STYLE_ID = '{5940675A-B579-460E-94D1-54222C63F5DA}'

_TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                         '..', 'assets', '企業概要書_テンプレート.pptx')


# ----------------------------------------------------------------- 基盤
def open_template(path: str | None = None) -> Presentation:
    """テンプレートを一時ファイルへ複製して開く（原本は書き換えない）。"""
    src = path or os.path.normpath(_TEMPLATE)
    tmp = tempfile.NamedTemporaryFile(suffix='.pptx', delete=False)
    tmp.close()
    shutil.copyfile(src, tmp.name)
    return Presentation(tmp.name)


def _sldIdLst(prs):
    return prs.slides._sldIdLst


def delete_slide(prs, index: int) -> None:
    """0始まりのindexのスライドを削除する。"""
    lst = _sldIdLst(prs)
    ids = list(lst)
    rid = ids[index].get(qn('r:id'))
    prs.part.drop_rel(rid)
    lst.remove(ids[index])


def keep_slides(prs, keep_1based: list[int]) -> None:
    """指定した頁番号（1始まり）だけ残して他を削除する。"""
    keep = {i - 1 for i in keep_1based}
    for i in sorted(set(range(len(prs.slides._sldIdLst))) - keep, reverse=True):
        delete_slide(prs, i)


def _unique_slide_partname(prs):
    """既存パート名と衝突しないスライドのパート名を返す。

    スライドを削除してから add_slide() すると python-pptx が既存と同じ
    partname を振ることがあり、保存時に zip 内で重複して壊れる。
    """
    from pptx.opc.packuri import PackURI
    used = {p.partname for p in prs.part.package.iter_parts()}
    i = 1
    while True:
        cand = PackURI('/ppt/slides/slide%d.xml' % i)
        if cand not in used:
            return cand
        i += 1


_R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def _relink(src_part, dst_part, element):
    """複製した図形XMLが参照する rId を、複製先のリレーションに張り替える。

    これをしないと画像を含むスライドを複製した瞬間にファイルが壊れる
    （<blip r:embed> が存在しないrIdを指す）。
    """
    attrs = ('{%s}embed' % _R_NS, '{%s}id' % _R_NS, '{%s}link' % _R_NS)
    mapping = {}
    for el in element.iter():
        for a in attrs:
            rid = el.get(a)
            if not rid:
                continue
            if rid not in mapping:
                try:
                    rel = src_part.rels[rid]
                except KeyError:
                    continue
                if rel.is_external:
                    mapping[rid] = dst_part.rels.get_or_add_ext_rel(
                        rel.reltype, rel.target_ref)
                else:
                    mapping[rid] = dst_part.relate_to(rel.target_part, rel.reltype)
            if rid in mapping:
                el.set(a, mapping[rid])


def duplicate_slide(prs, index: int, after: int | None = None):
    """index（0始まり）のスライドを複製し、after の直後へ挿入する。"""
    src = prs.slides[index]
    dst = prs.slides.add_slide(src.slide_layout)
    dst.part.partname = _unique_slide_partname(prs)
    for sh in list(dst.shapes):
        sh._element.getparent().remove(sh._element)
    for sh in src.shapes:
        el = copy.deepcopy(sh._element)
        _relink(src.part, dst.part, el)
        dst.shapes._spTree.append(el)
    lst = _sldIdLst(prs)
    ids = list(lst)
    new = ids[-1]
    lst.remove(new)
    pos = (after if after is not None else index) + 1
    lst.insert(pos, new)
    return prs.slides[pos]


def strip_red_banners(prs) -> int:
    """テンプレート由来の赤帯（社内向け指示）を全スライドから削除する。"""
    n = 0
    for s in prs.slides:
        for sh in list(s.shapes):
            if not sh.has_text_frame:
                continue
            red = False
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    try:
                        if run.font.color and run.font.color.rgb == RGBColor(0xFF, 0x00, 0x00):
                            red = True
                    except Exception:
                        pass
            if red:
                sh._element.getparent().remove(sh._element)
                n += 1
    return n


def fix_broken_pictures(prs) -> int:
    """解決できない画像参照を持つ図形を削除する。

    テンプレートの中扉（Section扉）には、リレーションが欠落した0×0の
    図形が2つ残っており、そのまま保存すると PowerPoint が
    「修復が必要」と判定する。テンプレート由来の既知の不具合なので、
    ビルドの最初に必ず通すこと。
    """
    n = 0
    for s in prs.slides:
        for sh in list(s.shapes):
            el = sh._element
            bad = False
            for e in el.iter():
                rid = e.get('{%s}embed' % _R_NS) or e.get('{%s}link' % _R_NS)
                if rid and rid not in s.part.rels:
                    bad = True
                    break
            if bad:
                el.getparent().remove(el)
                n += 1
    return n


FRAME_KEEP = ('タイトル', '字幕')


def clear_body(slide, keep_names=()) -> None:
    """タイトル／字幕／セクション名を残し、本文領域の図形を全て削除する。"""
    for sh in list(slide.shapes):
        if sh.is_placeholder and sh.placeholder_format.idx in (0, 1):
            continue
        if sh.name in keep_names:
            continue
        try:
            if sh.name.startswith('タイトル') and abs(sh.left / 914400 - SECTION_POS[0]) < 0.4:
                continue          # 右上のセクション名テキストボックス
        except Exception:
            pass
        sh._element.getparent().remove(sh._element)


# ----------------------------------------------------------------- 文字
def _apply_font(run, size, bold, color, font):
    f = run.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color is not None:
        f.color.rgb = color
    if font:
        f.name = font
        rPr = run._r.get_or_add_rPr()
        for tag in ('a:latin', 'a:ea', 'a:cs'):
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {})
                rPr.append(el)
            el.set('typeface', font)


def no_bullet(paragraph):
    """テンプレート由来の箇条書き記号を消す。"""
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ('a:buChar', 'a:buAutoNum', 'a:buBlip'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if pPr.find(qn('a:buNone')) is None:
        pPr.append(pPr.makeelement(qn('a:buNone'), {}))


def set_text(shape, text, size=None, bold=None, color=None, font=JP,
             align=None, line_sp=None, space=None, bullet=None):
    """図形のテキストを差し替える（最初の段落の書式を踏襲）。"""
    tf = shape.text_frame
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    lines = str(text).split('\n')
    p0 = tf.paragraphs[0]
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    for i, ln in enumerate(lines):
        p = p0 if i == 0 else tf.add_paragraph()
        if align is not None:
            p.alignment = align
        if line_sp:
            p.line_spacing = line_sp
        if space is not None:
            p.space_after = Pt(space)
        if bullet is False:
            no_bullet(p)
        run = p.add_run()
        run.text = ln
        _apply_font(run, size, bold, color, font)
    return shape


def _ph(slide, idx):
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.idx == idx:
            return sh
    return None


def set_title(slide, text):
    sh = _ph(slide, 0)
    if sh is not None:
        set_text(sh, text, font=JP)
    return sh


def set_cover_title(slide, industry: str, area: str):
    """表紙タイトル。「企業概要書」は既定サイズ、「<業種／地区>」は24ptで続ける。

    当社の表紙はこの2ラン構成が決まり。1ランで書くと折り返して
    下の罫線と重なる（鶴岡260923・鳥塚260924いずれもこの構成）。
    """
    sh = _ph(slide, 0)
    if sh is None:
        return None
    tf = sh.text_frame
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    p0 = tf.paragraphs[0]
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    r1 = p0.add_run()
    r1.text = '企業概要書  '
    r2 = p0.add_run()
    r2.text = f'<{industry}／{area}>'
    r2.font.size = Pt(24)
    return sh


def set_lead(slide, text):
    """字幕（リード文）。『だから何なのか』を1〜2行で断言する。全角110字目安。"""
    sh = _ph(slide, 1)
    if sh is not None:
        set_text(sh, text, font=JP)
    return sh


def set_section(slide, text):
    for sh in slide.shapes:
        if (not sh.is_placeholder and sh.has_text_frame
                and abs(sh.left / 914400 - SECTION_POS[0]) < 0.4
                and sh.top / 914400 < 0.5):
            set_text(sh, text, size=12, bold=True, font=JP)
            return sh
    return textbox(slide, *SECTION_POS, text, size=12, bold=True,
                   align=PP_ALIGN.RIGHT, font=JP)


def textbox(slide, x, y, w, h, text='', size=10, bold=False, color=BODY,
            font=JP, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            line_sp=None, space=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(str(text).split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        if line_sp:
            p.line_spacing = line_sp
        run = p.add_run()
        run.text = ln
        _apply_font(run, size, bold, color, font)
    return box


def rich(slide, x, y, w, h, parts, anchor=MSO_ANCHOR.TOP):
    """見出し＋本文が交互に続くブロック。parts=[(text, style_dict), ...]

    style_dict: newpara/size/bold/color/font/align/line_sp/space
    """
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    first = True
    for text, st in parts:
        if st.get('newpara') and not first:
            p = tf.add_paragraph()
        p.space_after = Pt(st.get('space', 2))
        if st.get('line_sp'):
            p.line_spacing = st['line_sp']
        p.alignment = st.get('align', PP_ALIGN.LEFT)
        if text != '':
            run = p.add_run()
            run.text = text
            _apply_font(run, st.get('size', 9), st.get('bold', False),
                        st.get('color', BODY), st.get('font', JP))
        first = False
    return box


def note(slide, text, y=None, size=7.5):
    """最下部の注記（※で始まる出所・補足・要確認事項）。"""
    x, y0, w, h = NOTE_POS
    return textbox(slide, x, y if y is not None else y0, w, h, text,
                   size=size, color=NOTE, font=JP, line_sp=1.12, space=1)


# ----------------------------------------------------------------- 図形
def _noshadow(sp):
    try:
        sp.shadow.inherit = False
    except Exception:
        pass


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75,
         shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    _noshadow(sp)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(lw)
    sp.text_frame.word_wrap = True
    return sp


def section_bar(slide, x, y, w, text, h=0.39, size=14):
    """本文内の区画見出し（紺帯）。会社概要／株主／役員 などに使う。"""
    sp = rect(slide, x, y, w, h, fill=NAVY)
    set_text(sp, text, size=size, bold=True, color=WHITE, font=JP,
             align=PP_ALIGN.LEFT)
    sp.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    sp.text_frame.margin_left = Inches(0.12)
    return y + h


def comment_box(slide, x, y, w, h, parts, title='Comment', title_size=16):
    """右カラムの Comment ボックス（枠＋見出し＋本文）。

    parts は rich() と同じ形式。見出しは ≪…≫ を BLUE 太字で。
    """
    frame = rect(slide, x, y, w, h, fill=WHITE, line=NAVY, lw=1.0)
    textbox(slide, x + 0.10, y + 0.08, w - 0.20, 0.30, title,
            size=title_size, bold=True, color=NAVY, font=JP)
    rich(slide, x + 0.12, y + 0.46, w - 0.24, h - 0.56, parts)
    return frame


def kv_boxes(slide, x, y, w, items, box_h=0.98, gap=0.06,
             head='本件のポイント', head_h=0.39):
    """『本件のポイント』の縦積みカード（見出し紺帯＋白カード×n）。

    items=[(見出し, 本文), ...]
    """
    if head:
        hb = rect(slide, x, y, w, head_h, fill=NAVY)
        set_text(hb, head, size=11, bold=True, color=WHITE, font=JP)
        hb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        hb.text_frame.margin_left = Inches(0.12)
        y += head_h + gap
    for ttl, txt in items:
        card = rect(slide, x, y, w, box_h, fill=WHITE, line=GREY, lw=0.75)
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.10)
        tf.margin_top = tf.margin_bottom = Inches(0.05)
        p0 = tf.paragraphs[0]
        r = p0.add_run()
        r.text = ttl
        _apply_font(r, 10, True, BLUE, JP)
        p1 = tf.add_paragraph()
        p1.line_spacing = 1.14
        r2 = p1.add_run()
        r2.text = txt
        _apply_font(r2, 9, False, BODY, JP)
        y += box_h + gap
    return y


def feature_bands(slide, items, x=0.49, y=1.23, w=10.71, h=1.81, gap=0.38):
    """『特徴／強み／成長性』の3バンド（鳥塚フォーマット）。

    items=[(ラベル, 見出し, 本文), ...]
    """
    for label, head, txt in items:
        rect(slide, x, y, w, h, fill=WHITE, line=GREY, lw=0.75)
        rect(slide, x, y, 0.055, h, fill=NAVY)
        lb = rect(slide, x + 0.22, y + 0.38, 1.30, 1.06, fill=NAVY)
        set_text(lb, label, size=13, bold=True, color=WHITE, font=JP,
                 align=PP_ALIGN.CENTER)
        lb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        textbox(slide, x + 1.82, y + 0.28, w - 2.05, 0.34, head,
                size=14, bold=True, color=NAVY, font=JP)
        textbox(slide, x + 1.82, y + 0.72, w - 2.05, h - 0.90, txt,
                size=10.5, bold=False, color=BODY, font=JP, line_sp=1.22)
        y += h + gap
    return y


def point_band(slide, x, y, w, items, h=0.92, label='ポイント'):
    """最下部の『ポイント』帯（✓×n）。鳥塚フォーマット。"""
    rect(slide, x, y, w, h, fill=PALE, line=GREY, lw=0.75,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(slide, x, y, 1.55, h, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    textbox(slide, x + 0.10, y + h / 2 - 0.12, 1.39, 0.24, label,
            size=13, bold=True, color=WHITE, font=JP, align=PP_ALIGN.CENTER)
    n = len(items)
    colw = (w - 1.72) / n
    for i, t in enumerate(items):
        cx = x + 1.72 + i * colw
        ov = rect(slide, cx, y + h / 2 - 0.13, 0.26, 0.26, fill=BLUE,
                  shape=MSO_SHAPE.OVAL)
        set_text(ov, '✓', size=12, bold=True, color=WHITE, font=JP,
                 align=PP_ALIGN.CENTER)
        ov.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        textbox(slide, cx + 0.34, y + 0.18, colw - 0.44, h - 0.30, t,
                size=9.3, bold=True, color=DARK, font=JP, line_sp=1.18,
                anchor=MSO_ANCHOR.MIDDLE)
    return y + h


# ----------------------------------------------------------------- 表
def _set_cell(cell, text, size, bold, color, font, align, fill):
    cell.margin_left = Inches(0.05)
    cell.margin_right = Inches(0.05)
    cell.margin_top = Inches(0.01)
    cell.margin_bottom = Inches(0.01)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    tf = cell.text_frame
    tf.word_wrap = True
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    txt = '　' if text in (None, '') else str(text)
    lines = txt.split('\n')
    p0 = tf.paragraphs[0]
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    for i, ln in enumerate(lines):
        p = p0 if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = ln
        _apply_font(run, size, bold, color, font)


def is_number(s) -> bool:
    s = str(s)
    return bool(s) and all(ch in '0123456789,.△▲%-−+ ()' for ch in s)


def house_table(slide, x, y, w, col_w, rows, row_h=0.23, head_h=0.23,
                size=8.5, head_rows=1, bold_rows=(), aligns=None,
                head_fill=NAVY, fills=None, red_minus=True):
    """当社標準の表。

    - ヘッダ：紺地・白文字・8.5pt太字・中央揃え
    - 本文　：白地・1A1A1A・8.5pt、文字は游ゴシック左寄せ／数字はArial右寄せ
    - col_w は相対幅でよい（w に合わせて正規化する）

    rows[0] の先頭セルには「（単位：千円）」を置くのが当社の慣行。
    """
    nr, nc = len(rows), len(col_w)
    scale = w / sum(col_w)
    gf = slide.shapes.add_table(
        nr, nc, Inches(x), Inches(y), Inches(w),
        Inches(head_h * head_rows + row_h * (nr - head_rows)))
    t = gf.table
    tbl = t._tbl
    pr = tbl.find(qn('a:tblPr'))
    if pr is not None:
        pr.set('firstRow', '0')
        pr.set('bandRow', '0')
        st = pr.find(qn('a:tableStyleId'))
        if st is None:
            st = pr.makeelement(qn('a:tableStyleId'), {})
            pr.append(st)
        st.text = TABLE_STYLE_ID
    for j, cw in enumerate(col_w):
        t.columns[j].width = Emu(int(Inches(cw * scale)))
    for i in range(nr):
        t.rows[i].height = Inches(head_h if i < head_rows else row_h)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            txt = '' if val is None else str(val)
            if i < head_rows:
                _set_cell(t.cell(i, j), txt, size, True, WHITE,
                          JP if j == 0 else NUM, PP_ALIGN.CENTER, head_fill)
                continue
            num = is_number(txt)
            if aligns and j < len(aligns):
                al = aligns[j]
            else:
                al = PP_ALIGN.LEFT if j == 0 else (
                    PP_ALIGN.RIGHT if num else PP_ALIGN.LEFT)
            col = RED if (red_minus and txt.startswith(('△', '▲'))) else BODY
            bg = WHITE
            if fills:
                bg = fills.get((i, j), fills.get(('row', i), WHITE))
            _set_cell(t.cell(i, j), txt, size, i in bold_rows, col,
                      NUM if num else JP, al, bg)
    for i in range(nr):
        for j in range(nc):
            tcPr = t.cell(i, j)._tc.get_or_add_tcPr()
            for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
                el = tcPr.makeelement(qn(tag), {'w': '6350', 'cap': 'flat',
                                                'cmpd': 'sng', 'algn': 'ctr'})
                sf = el.makeelement(qn('a:solidFill'), {})
                sf.append(sf.makeelement(qn('a:srgbClr'), {'val': 'B9C2C7'}))
                el.append(sf)
                tcPr.append(el)
    return gf


# ----------------------------------------------------------------- 図表
def bar_chart(slide, x, y, w, h, labels, values, title=None,
              highlight=None, unit='', value_fmt='{:,}', size=8):
    """簡易縦棒グラフ（図形で描画。後から数字を直せるように画像にしない）。"""
    rect(slide, x, y, w, h, fill=RGBColor(0xFB, 0xFC, 0xFD), line=GREY, lw=0.75)
    if title:
        textbox(slide, x + 0.12, y + 0.08, w - 0.24, 0.22, title,
                size=9, bold=True, color=NAVY, font=JP)
    n = len(values)
    mx = max(values) if values else 1
    base = y + h - 0.34
    maxh = h - (0.62 if title else 0.42) - 0.20
    colw = (w - 0.56) / n
    for i, (v, lb) in enumerate(zip(values, labels)):
        bh = maxh * (v / mx) if mx else 0
        bx = x + 0.34 + i * colw
        rect(slide, bx, base - bh, colw * 0.60, bh,
             fill=NAVY if (highlight is not None and i == highlight) else BLUE)
        textbox(slide, bx - 0.08, base - bh - 0.20, colw * 0.76, 0.18,
                value_fmt.format(v), size=size - 1.2, color=NAVY,
                align=PP_ALIGN.CENTER, font=NUM)
        textbox(slide, bx - 0.08, base + 0.05, colw * 0.76, 0.18, lb,
                size=size - 1, color=NOTE, align=PP_ALIGN.CENTER, font=NUM)
    rect(slide, x + 0.24, base, w - 0.44, 0.015, fill=GREY)
    if unit:
        textbox(slide, x, y - 0.20, w, 0.18, unit, size=7.5, color=NOTE,
                align=PP_ALIGN.RIGHT, font=JP)


def flow_block(slide, x, y, w, h, title, lines, color=BLUE,
               fill=RGBColor(0xF7, 0xFA, 0xFC)):
    """商流図のブロック（左に色帯＋見出し＋明細）。"""
    rect(slide, x, y, w, h, fill=fill, line=color, lw=1.0)
    rect(slide, x, y, 0.055, h, fill=color)
    textbox(slide, x + 0.16, y + 0.10, w - 0.30, 0.24, title,
            size=9.5, bold=True, color=NAVY, font=JP)
    textbox(slide, x + 0.16, y + 0.40, w - 0.30, h - 0.50, lines,
            size=8.3, color=NOTE, font=JP, line_sp=1.18, space=1)


def connector(slide, x, y, w, color=BLUE, thickness=0.024):
    rect(slide, x, y - thickness / 2, w, thickness, fill=color)


# ----------------------------------------------------------------- 書式
def fmt(v, dash='－'):
    """数値を 1,234 / △1,234 形式へ。None は dash。"""
    if v is None:
        return dash
    if isinstance(v, str):
        return v
    v = int(round(v))
    return f'△{abs(v):,}' if v < 0 else f'{v:,}'


def pct(v, dash='－', nd=1):
    if v is None or v == '':
        return dash
    s = f'{abs(v) * 100:.{nd}f}%'
    return ('△' + s) if v < 0 else s


def k(v):
    """円 → 千円（四捨五入）。案件マスターは円単位で持つため表示時に変換する。"""
    return None if v is None else int(round(v / 1000))

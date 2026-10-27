"""案件マスター IM→シートの「貼付範囲」を組むための書式キット（openpyxl）。

貼付範囲＝企業概要書の表と同じ行・列・書式でExcel上に組んだ範囲。概要書はこの範囲をコピーして貼るだけで作る。
書式は鳥塚案件の概要書に準拠（references/im-design.md）：

  見出し行  紺(0B3041)・白字・游ゴシック10pt・中央      → kind 'hdr'（左寄せ太字は 'hdrl'）
  科目列    グレー(E3E7E9)・游ゴシック10.5pt             → 'lab'（塗りなし 'labw'）
  数値      Arial10.5pt・右寄せ・千円表示                  → 'num'
  比率行    7pt斜体・右寄せ（前期比／利益率）             → 'rlab' / 'rnum'
  小計・合計 淡青(E1F3FB)・太字                           → 'sub' / 'subn'
  強調合計  グレー太字＋紺の上下罫線（調整後EBITDA等）   → 'gryb' / 'grybn' + border='box'
  総合計    紺・白字太字（BS詳細の資産合計等）           → 'tot' / 'totn'
  罫線      縦＝白(medium)、横＝グレー(A6A6A6)細線       → border='std'

使い方::

    from master_im_style import *
    ws = im_sheet(wb, 'PLハイライト', '財務ハイライト（P/L）・調整後EBITDA')
    cw(ws, {'B': 1.82, 'C': 1.2, 'D': 1.2, 'E': 1.2, 'F': 1.2})      # 列幅はインチ（＝パワポ上の幅）
    put(ws, 4, 6, '(単位：千円)', 'unit', border=None)
    row(ws, 5, 2, ['', '2024年3月期', '2025年3月期', '2026年3月期', '2026年4〜7月'], 'hdr', height=19)
    row(ws, 6, 2, ['売上高', '=PL!C10', '=PL!E10', '=PL!G10', '=PL!I10'], ['lab'] + ['num'] * 4, fmts=[None] + [KI] * 4, height=19)
    paste(ws, 'B4:F25', '財務ハイライト（P/L）', PASTE)                 # A2に貼付範囲を明記し一覧に登録
    ...
    write_paste_index(wb['ルール'], PASTE, start_row=13)               # ルールシートに貼付範囲一覧

注意：
  - 貼付範囲の行は他の表と共有しない（横並びの2表を同じ行に置くと、片方の行高がもう片方を押し広げる）。
    パワポで横に並べる表も、Excelでは縦に離して置く。
  - 貼付範囲の下に「参考（貼付範囲外）」欄を置き、リード文・Comment用の集計（平均・人数・比率等）はそこで計算する。
  - 数値は円で入力・千円表示（KI）。パワポ側は表示形式どおりに文字化される（△・千円丸め・%）。
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

NAVY, LBLUE, GRAY, WHITE = '0B3041', 'E1F3FB', 'E3E7E9', 'FFFFFF'
JP, NUM, NOTE = '游ゴシック', 'Arial', '游ゴシック Medium'

KI = '#,##0,;△#,##0,;"-"'          # 円入力・千円表示（概要書用：赤字表示なし）
PI = '0.0%;△0.0%;"-"'
YI = '#,##0;△#,##0;"-"'             # 円・株数・千円入力済みの値など
AI = '#,##0.00;△#,##0.00;"-"'       # 面積

WV = Side(style='medium', color=WHITE)
HB = Side(style='thin', color='A6A6A6')
NV = Side(style='thin', color=NAVY)

# kind -> (フォント, サイズ, 太字, 斜体, 文字色, 塗り, 横位置)
KINDS = {
    'hdr':   (JP, 10, False, False, WHITE, NAVY, 'center'),
    'hdrl':  (JP, 10, True, False, WHITE, NAVY, 'left'),
    'lab':   (JP, 10.5, False, False, '000000', GRAY, 'left'),
    'labw':  (JP, 10.5, False, False, '000000', None, 'left'),
    'num':   (NUM, 10.5, False, False, '000000', None, 'right'),
    'txt':   (JP, 10, False, False, '000000', None, 'left'),
    'txtc':  (JP, 10, False, False, '000000', None, 'center'),
    'rlab':  (JP, 7, False, True, '000000', GRAY, 'right'),
    'rnum':  (NUM, 7, False, True, '000000', None, 'right'),
    'sub':   (JP, 10.5, True, False, '000000', LBLUE, 'left'),
    'subn':  (NUM, 10.5, True, False, '000000', LBLUE, 'right'),
    'subt':  (JP, 10.5, True, False, '000000', LBLUE, 'left'),
    'gry':   (JP, 10.5, False, False, '000000', GRAY, 'left'),
    'gryn':  (NUM, 10.5, False, False, '000000', GRAY, 'right'),
    'gryb':  (JP, 10.5, True, False, '000000', GRAY, 'left'),
    'grybn': (NUM, 10.5, True, False, '000000', GRAY, 'right'),
    'tot':   (JP, 10.5, True, False, WHITE, NAVY, 'left'),
    'totn':  (NUM, 10.5, True, False, WHITE, NAVY, 'right'),
    'unit':  (JP, 9, False, False, '000000', None, 'right'),
    'blank': (JP, 10.5, False, False, '000000', None, 'left'),
}


def im_sheet(wb, name, title_):
    """IMシートを作る（目盛線なし、A1にタイトル）。"""
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws['A1'] = title_
    ws['A1'].font = Font(name=JP, size=14, bold=True, color=NAVY)
    return ws


def cw(ws, widths):
    """列幅をインチで指定（貼付時にパワポ上の幅になる）。A列は余白。"""
    ws.column_dimensions['A'].width = 2
    for col, inch in widths.items():
        ws.column_dimensions[col].width = round((inch * 96 - 5) / 7, 2)


def put(ws, r, c, v, kind, fmt=None, size=None, h=None, indent=0, border='std', wrap=True, v_='center'):
    """1セルを kind の書式で書く。border: 'std'（縦白＋下グレー）/'v'（縦白のみ）/'box'（紺の上下）/'bot'（紺の下）/None"""
    f, sz, b, it, colr, fill, al = KINDS[kind]
    x = ws.cell(r, c, v)
    x.font = Font(name=f, size=size or sz, bold=b, italic=it, color=colr)
    if fill:
        x.fill = PatternFill('solid', fgColor=fill)
    x.alignment = Alignment(horizontal=h or al, vertical=v_, wrap_text=wrap, indent=indent)
    if fmt:
        x.number_format = fmt
    if border == 'std':
        x.border = Border(left=WV, right=WV, bottom=HB)
    elif border == 'box':
        x.border = Border(left=WV, right=WV, top=NV, bottom=NV)
    elif border == 'bot':
        x.border = Border(left=WV, right=WV, bottom=NV)
    elif border == 'v':
        x.border = Border(left=WV, right=WV)
    return x


def row(ws, r, c0, vals, kinds, fmts=None, size=None, height=None, border='std', aligns=None, indent=None):
    """1行を書く。kinds/fmts は単一値または列ごとのリスト。height は pt（パワポ上の行高）。"""
    for j, v in enumerate(vals):
        k_ = kinds[j] if isinstance(kinds, (list, tuple)) else kinds
        put(ws, r, c0 + j, v, k_, fmt=(fmts[j] if isinstance(fmts, (list, tuple)) else fmts), size=size,
            h=(aligns[j] if aligns else None), indent=(indent[j] if indent else 0), border=border)
    if height:
        ws.row_dimensions[r].height = height


def paste(ws, rng, page, registry=None, key=None, cell='A2'):
    """貼付範囲を登録する。A2に赤字で「貼付範囲 → 概要書のページ」を明記。registry(list)に (シート, 範囲, ページ) を追加。
    公式マスターのテンプレートシート（行2が見出し）は cell='A1' に書く。"""
    ws[cell] = f'■ 貼付範囲 {rng} → 企業概要書「{page}」（範囲をコピーしてパワポに貼り付け。数字はここを直す）'
    ws[cell].font = Font(name=NOTE, size=9, color='C00000')
    ws[cell].alignment = Alignment(wrap_text=False)
    if registry is not None:
        registry.append((key or ws.title, f'{ws.title}!{rng}', page))
    return f'{ws.title}!{rng}'


def note(ws, r, t, c=2):
    """貼付範囲外の注記（游ゴシック Medium 9pt）。"""
    ws.cell(r, c, t).font = Font(name=NOTE, size=9)
    ws.cell(r, c).alignment = Alignment(wrap_text=False)


def write_paste_index(ws, registry, start_row=13):
    """ルールシートに「貼付範囲一覧」と作り方を書く。"""
    r = start_row
    ws.cell(r, 2, '■ 企業概要書への貼付範囲一覧（IM→セクション）').font = Font(name=JP, size=11, bold=True, color=NAVY); r += 1
    ws.cell(r, 2, '   作り方：①転記→②IMシートの貼付範囲を確認→③範囲をコピー→④概要書の該当ページへ貼り付け→⑤リード文・Commentを記入。'
                  '業種により不要なページは削除、必要なページ（商流・生産サイクル等）は追加する。').font = Font(name=JP, size=10); r += 1
    for key, ref, pg in registry:
        ws.cell(r, 2, f'   {ref}　→　{pg}').font = Font(name=NUM, size=10)
        r += 1
    return r


def ungroup_cols(ws, upto=40):
    """公式マスターの列幅定義は C:G のように範囲でまとまっていることがある。
    1列だけ非表示にすると範囲全体が消えるため、列ごとの定義に分解する（非表示にする前に呼ぶ）。"""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.dimensions import ColumnDimension
    for key, dim in list(ws.column_dimensions.items()):
        if dim.min and dim.max and dim.max > dim.min and dim.min <= upto:
            w, hid = dim.width, dim.hidden
            lo, hi = dim.min, min(dim.max, upto)
            del ws.column_dimensions[key]
            for i in range(lo, hi + 1):
                c = get_column_letter(i)
                ws.column_dimensions[c] = ColumnDimension(ws, index=c, width=w, hidden=hid)


def hide_cols(ws, *cols):
    """「IMに貼るときは非表示」の列（1期目・氏名等）を非表示にする。paste_range は非表示の列を貼らない。"""
    ungroup_cols(ws)
    for c in cols:
        ws.column_dimensions[c].hidden = True

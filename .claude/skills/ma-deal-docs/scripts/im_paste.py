"""企業概要書（IM）作成キット：案件マスターの「貼付範囲」をパワポへ貼り、テンプレートの枠に文章を書く。

使い方（詳細は references/im-design.md）::

    import sys; sys.path.insert(0, '.claude/skills/ma-deal-docs/scripts')
    from im_paste import *
    from openpyxl import load_workbook
    from pptx import Presentation

    wb  = load_workbook('案件マスター.xlsx', data_only=True)   # 再計算済み（キャッシュ値あり）のマスター
    prs = Presentation('概要書テンプレート.pptx')
    s = prs.slides[5]                                           # 財務ハイライト（P/L）
    set_lead(s, '直近3期の売上高は…')                           # 字幕（リード文）
    paste_range(s, wb, 'PLハイライト!B4:F25', 0.34, 1.6, 6.62)  # 表＝マスターの範囲をそのまま貼る
    fill_comment(shape_by_id(s, 9), ['≪売上高≫', '・…'])        # Comment枠（テンプレートの既存図形）

数字はパワポ上で打たない。直すときはマスターを直して貼り直す。
"""
import copy
import re
from decimal import Decimal, ROUND_HALF_UP

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

JP, NUM = '游ゴシック', 'Arial'


# ---------- テンプレート操作 ----------
def shape_by_id(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def remove(slide, *ids):
    for sid in ids:
        sh = shape_by_id(slide, sid)
        sh._element.getparent().remove(sh._element)


def remove_yellow(prs):
    for s in prs.slides:
        for sh in list(s.shapes):
            try:
                if sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == 'FFFF00':
                    sh._element.getparent().remove(sh._element)
            except Exception:
                pass


def set_text(sh, paras, size=None, bold=None, color=None):
    """既存の書式（先頭段落・先頭ラン）を保ったまま文字列を差し替える。paras: str | list[str]"""
    tf = sh.text_frame
    paras = paras if isinstance(paras, list) else [paras]
    p0 = tf.paragraphs[0]._p
    tmpl = copy.deepcopy(p0)
    txBody = p0.getparent()
    for p in list(txBody.findall(qn('a:p'))):
        txBody.remove(p)
    for t in paras:
        p = copy.deepcopy(tmpl)
        runs = p.findall(qn('a:r'))
        for r in runs[1:]:
            p.remove(r)
        for extra in p.findall(qn('a:br')) + p.findall(qn('a:fld')):
            p.remove(extra)
        if not runs:
            r = etree.SubElement(p, qn('a:r')); etree.SubElement(r, qn('a:t'))
            end = p.find(qn('a:endParaRPr'))
            if end is not None:
                p.remove(end); p.append(end)
                r.insert(0, copy.deepcopy(end)); r[0].tag = qn('a:rPr')
        else:
            r = runs[0]
        r.find(qn('a:t')).text = t
        rPr = r.find(qn('a:rPr'))
        if rPr is None:
            rPr = etree.SubElement(r, qn('a:rPr')); r.remove(rPr); r.insert(0, rPr)
        if size:
            rPr.set('sz', str(int(size * 100)))
        if bold is not None:
            rPr.set('b', '1' if bold else '0')
        txBody.append(p)
    return sh


def delete_slides(prs, numbers):
    """1始まりのスライド番号で削除（孤立パートは保存後にclean）。"""
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for n in sorted(numbers, reverse=True):
        sid = ids[n - 1]
        prs.part.drop_rel(sid.get(qn('r:id')))
        lst.remove(sid)


def remove_null_pics(prs):
    """テンプレート由来の、画像実体がない（Target=NULL）極小ピクチャを除去する。"""
    for s in prs.slides:
        rels = s.part.rels
        for pic in list(s.shapes._spTree.iter(qn('p:pic'))):
            blip = pic.find('.//' + qn('a:blip'))
            rid = blip.get(qn('r:embed')) if blip is not None else None
            rel = rels.get(rid) if rid else None
            bad = rel is None or (getattr(rel, 'target_ref', '') or '').endswith('NULL')
            if bad:
                pic.getparent().remove(pic)
        for rid, rel in list(rels.items()):
            if (getattr(rel, 'target_ref', '') or '').endswith('NULL'):
                s.part.drop_rel(rid) if hasattr(s.part, 'drop_rel') else None


# =====================================================================
# 案件マスターの貼付範囲 → パワポ表（Excelから「貼り付け」たのと同じ見た目で再現する）
# =====================================================================

_BORDER_W = {'hair': 3175, 'thin': 6350, 'medium': 12700, 'thick': 25400, 'double': 19050}


def _split_fmt(nf):
    out, cur, q = [], '', False
    for ch in nf:
        if ch == '"':
            q = not q
        if ch == ';' and not q:
            out.append(cur); cur = ''
        else:
            cur += ch
    out.append(cur)
    return out


def fmt_value(v, nf):
    """Excelの表示形式で値を文字列化（概要書で使う形式のみ対応）。"""
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, str):
        return v
    if not nf or nf == 'General':
        return f'{v:,.0f}' if float(v).is_integer() else f'{v:,.2f}'
    secs = _split_fmt(nf)
    if v < 0 and len(secs) > 1:
        sec, x, neg = secs[1], -v, False
    elif v == 0 and len(secs) > 2:
        sec, x, neg = secs[2], 0, False
    else:
        sec, x, neg = secs[0], v, v < 0
        x = abs(x)
    sec = re.sub(r'\[[^\]]*\]', '', sec)
    m_ = re.search(r'[#0][#0,\.]*%?', sec)
    if not m_:
        return sec.replace('"', '')
    tok = m_.group(0)
    pctf = tok.endswith('%')
    core = tok.rstrip('%')
    scale = len(core) - len(core.rstrip(','))
    core = core.rstrip(',')
    dec = len(core.split('.')[1]) if '.' in core else 0
    if pctf:
        x = x * 100
    x = x / (1000 ** scale)
    q = Decimal(str(x)).quantize(Decimal(1) / (Decimal(10) ** dec), rounding=ROUND_HALF_UP)
    s = f'{q:,.{dec}f}' if ',' in core else f'{q:.{dec}f}'
    if neg and float(q) != 0:
        s = '-' + s
    pre, post = sec[:m_.start()], sec[m_.end():]
    return (pre + s + ('%' if pctf else '') + post).replace('"', '')


def _rgb(color):
    if color is None or color.type != 'rgb' or not isinstance(color.rgb, str):
        return None
    return color.rgb[-6:]


def _ln(tcPr, tag, side):
    ln = etree.SubElement(tcPr, qn(tag))
    if side is not None and side.style:
        ln.set('w', str(_BORDER_W.get(side.style, 6350)))
        sf = etree.SubElement(ln, qn('a:solidFill'))
        etree.SubElement(sf, qn('a:srgbClr'), val=_rgb(side.color) or '000000')
    else:
        ln.set('w', '6350')
        etree.SubElement(ln, qn('a:noFill'))


def paste_range(slide, wb, ref, x, y, w=None, hscale=1.0, fscale=1.0, margin=0.04):
    """案件マスターの範囲（'シート!B4:F25'）を書式ごとパワポの表にする。
    列幅・行高・結合・塗り・罫線・フォント・表示形式・配置をExcelから引き継ぐ。w指定時は列幅を比例調整。"""
    from openpyxl.utils import range_boundaries, get_column_letter
    sheet, rng = ref.split('!')
    ws = wb[sheet.strip("'")]
    c1, r1, c2, r2 = range_boundaries(rng)
    cols = list(range(c1, c2 + 1)); rows_ = list(range(r1, r2 + 1))
    cw_ = [((ws.column_dimensions[get_column_letter(c)].width or 8.43) * 7 + 5) / 96 for c in cols]
    k_ = (w / sum(cw_)) if w else 1.0
    cw_ = [c * k_ for c in cw_]
    rh_ = [((ws.row_dimensions[r].height or 15) / 72) * hscale for r in rows_]
    gf = slide.shapes.add_table(len(rows_), len(cols), Inches(x), Inches(y), Inches(sum(cw_)), Inches(sum(rh_)))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for a in ('firstRow', 'bandRow'):
        tblPr.set(a, '0')
    st = tblPr.find(qn('a:tableStyleId'))
    if st is not None:
        st.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'   # スタイルなし・罫線なし
    for j, c in enumerate(cw_):
        tbl.columns[j].width = Inches(c)
    for i, h in enumerate(rh_):
        tbl.rows[i].height = Inches(h)
    for mr in ws.merged_cells.ranges:
        if mr.min_row >= r1 and mr.max_row <= r2 and mr.min_col >= c1 and mr.max_col <= c2:
            tbl.cell(mr.min_row - r1, mr.min_col - c1).merge(tbl.cell(mr.max_row - r1, mr.max_col - c1))
    for i, r in enumerate(rows_):
        for j, c in enumerate(cols):
            xc = ws.cell(r, c)
            cell = tbl.cell(i, j)
            tc = cell._tc
            fnt = xc.font
            size = (fnt.sz or 10.5) * fscale
            txt_ = fmt_value(xc.value, xc.number_format)
            ind = int(xc.alignment.indent or 0)
            if txt_ and ind:
                txt_ = '　' * ind + txt_
            hal = xc.alignment.horizontal or ('right' if isinstance(xc.value, (int, float)) and not isinstance(xc.value, bool) else 'left')
            algn = {'left': 'l', 'center': 'ctr', 'right': 'r', 'centerContinuous': 'ctr'}.get(hal, 'l')
            txBody = tc.find(qn('a:txBody'))
            for p in txBody.findall(qn('a:p')):
                txBody.remove(p)
            latin = fnt.name or JP
            ea = '游ゴシック Medium' if latin == NUM else JP
            color = _rgb(fnt.color) or '000000'
            if not txt_:   # 空セルは行高を押し広げないよう小さいフォントで
                size = min(size, (ws.row_dimensions[r].height or 15) * 0.6)
            for line_ in (txt_.split('\n') if txt_ else ['']):
                p = etree.SubElement(txBody, qn('a:p'))
                etree.SubElement(p, qn('a:pPr'), algn=algn, fontAlgn='ctr')
                def _rpr(tag):
                    e = etree.SubElement(p, qn(tag), lang='ja-JP', altLang='en-US', sz=str(int(round(size * 100))),
                                         b='1' if fnt.b else '0', i='1' if fnt.i else '0', dirty='0')
                    sf = etree.SubElement(e, qn('a:solidFill')); etree.SubElement(sf, qn('a:srgbClr'), val=color)
                    etree.SubElement(e, qn('a:latin'), typeface=latin)
                    etree.SubElement(e, qn('a:ea'), typeface=ea)
                    return e
                if line_:
                    rr = etree.SubElement(p, qn('a:r'))
                    rp = _rpr('a:rPr'); p.remove(rp); rr.append(rp)
                    t = etree.SubElement(rr, qn('a:t')); t.text = line_
                _rpr('a:endParaRPr')
            tcPr = tc.get_or_add_tcPr()
            for a_ in list(tcPr):
                tcPr.remove(a_)
            mi = Inches(margin)
            tcPr.set('marL', str(int(mi))); tcPr.set('marR', str(int(mi))); tcPr.set('marT', '0'); tcPr.set('marB', '0')
            tcPr.set('anchor', 'ctr')
            b = xc.border
            _ln(tcPr, 'a:lnL', b.left); _ln(tcPr, 'a:lnR', b.right); _ln(tcPr, 'a:lnT', b.top); _ln(tcPr, 'a:lnB', b.bottom)
            fill = _rgb(xc.fill.fgColor) if xc.fill is not None and xc.fill.fill_type == 'solid' else None
            if fill:
                sf = etree.SubElement(tcPr, qn('a:solidFill')); etree.SubElement(sf, qn('a:srgbClr'), val=fill)
            else:
                etree.SubElement(tcPr, qn('a:noFill'))
    gf.name = f'表（{ref}）'
    return gf


# ---------- 鳥塚の概要書と同じ文字書式 ----------
def _run(p, t, size, bold, color_sys=True, name=JP):
    r = etree.SubElement(p, qn('a:r'))
    rPr = etree.SubElement(r, qn('a:rPr'), lang='ja-JP', altLang='en-US', sz=str(int(size * 100)), b='1' if bold else '0', dirty='0')
    sf = etree.SubElement(rPr, qn('a:solidFill'))
    etree.SubElement(sf, qn('a:sysClr'), val='windowText', lastClr='000000')
    etree.SubElement(rPr, qn('a:latin'), typeface=name)
    etree.SubElement(rPr, qn('a:ea'), typeface=JP)
    etree.SubElement(r, qn('a:t')).text = t
    return r


def fill_comment(sh, lines, body=9, head=10, after=3):
    """テンプレートのComment本文ボックスに書く。≪…≫・【…】・✓ で始まる行は見出し（10pt太字）、
    それ以外は本文（9pt）、空文字は空行。bodyPr（余白・自動調整）はテンプレートのまま。"""
    txBody = sh.text_frame._txBody
    for p in txBody.findall(qn('a:p')):
        txBody.remove(p)
    for ln in lines:
        p = etree.SubElement(txBody, qn('a:p'))
        pPr = etree.SubElement(p, qn('a:pPr'))
        sa = etree.SubElement(etree.SubElement(pPr, qn('a:spcAft')), qn('a:spcPts')); sa.set('val', str(after * 100))
        if ln == '':
            e = etree.SubElement(p, qn('a:endParaRPr'), lang='ja-JP', sz=str(int(body * 50)))
            continue
        hd = ln[0] in '≪【✓'
        _run(p, ln, head if hd else body, hd)
    return sh


def set_lead(slide, text_):
    """字幕（リード文）プレースホルダー。サイズ・太字はレイアウト既定、フォントは游ゴシック（鳥塚と同じ）。"""
    for sh in slide.placeholders:
        if sh.placeholder_format.idx == 1:
            txBody = sh.text_frame._txBody
            for p in txBody.findall(qn('a:p')):
                txBody.remove(p)
            p = etree.SubElement(txBody, qn('a:p'))
            r = etree.SubElement(p, qn('a:r'))
            rPr = etree.SubElement(r, qn('a:rPr'), lang='ja-JP', altLang='en-US', dirty='0')
            etree.SubElement(rPr, qn('a:latin'), typeface=JP)
            etree.SubElement(rPr, qn('a:ea'), typeface=JP)
            etree.SubElement(r, qn('a:t')).text = text_
            return sh


def note(slide, x, y, w, lines, size=8):
    """表の下の注記（※…）。黒・游ゴシック、余白0のテキストボックス。"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.18 * len(lines)))
    tf = tb.text_frame
    tf.word_wrap = True
    for a in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, a, 0)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.name = JP; r.font.color.rgb = RGBColor(0, 0, 0)
        rPr = r._r.get_or_add_rPr()
        etree.SubElement(rPr, qn('a:ea'), typeface=JP)
    return tb


def clone(src, dst_slide, x, y, w, h):
    """テンプレートの図形を複製して配置（書式をそのまま使う）。"""
    el = copy.deepcopy(src._element)
    dst_slide.shapes._spTree.append(el)
    sh = dst_slide.shapes[-1]
    sh.left, sh.top, sh.width, sh.height = Inches(x), Inches(y), Inches(w), Inches(h)
    ids = [s.shape_id for s in dst_slide.shapes]
    el.find('.//' + qn('p:cNvPr')).set('id', str(max(ids) + 1))
    return sh


def set_box(sh, lines, size=None, size1=None):
    """1行目は先頭段落の書式、2行目以降は2段落目の書式で差し替える（組織図の箱など）。"""
    txBody = sh.text_frame._txBody
    ps = txBody.findall(qn('a:p'))
    t1 = copy.deepcopy(ps[0]); t2 = copy.deepcopy(ps[1] if len(ps) > 1 else ps[0])
    for p in ps:
        txBody.remove(p)
    for i, ln in enumerate(lines):
        p = copy.deepcopy(t1 if i == 0 else t2)
        rs = p.findall(qn('a:r'))
        for r in rs[1:]:
            p.remove(r)
        for e in p.findall(qn('a:br')):
            p.remove(e)
        rs[0].find(qn('a:t')).text = ln
        sz = size1 if (i == 0 and size1) else size
        if sz:
            rs[0].find(qn('a:rPr')).set('sz', str(int(sz * 100)))
            e = p.find(qn('a:endParaRPr'))
            if e is not None:
                e.set('sz', str(int(sz * 100)))
        txBody.append(p)
    return sh


def range_height(wb, ref, upto_row=None):
    """貼付範囲の高さ（インチ）。upto_row指定時は範囲先頭からその行の直前まで。"""
    from openpyxl.utils import range_boundaries
    sheet, rng = ref.split('!')
    ws = wb[sheet.strip("'")]
    c1, r1, c2, r2 = range_boundaries(rng)
    end = upto_row if upto_row else r2 + 1
    return sum((ws.row_dimensions[r].height or 15) / 72 for r in range(r1, end))



def comment(slide, src_head, src_body, x, y, w, h, lines, head='Comment', body_dy=0.58):
    """テンプレートのComment枠（グレー枠＋見出し：src_head、本文テキストボックス：src_body）を複製して
    任意のページ・位置に置き、本文を鳥塚の書式で書く。既存のComment枠があるページはそれを place して使う。"""
    box = clone(src_head, slide, x, y, w, h)
    set_text(box, head)
    body = clone(src_body, slide, x, y + body_dy, w, h - body_dy - 0.05)
    fill_comment(body, lines)
    return box, body


def place(sh, x, y, w=None, h=None):
    """図形の位置・サイズ（インチ）を変える。"""
    sh.left, sh.top = Inches(x), Inches(y)
    if w is not None:
        sh.width = Inches(w)
    if h is not None:
        sh.height = Inches(h)


def paste_ref(wb, key):
    """マスターの「ルール」シートではなく、各IMシート A2 の貼付範囲表記から範囲を引く補助。
    例：paste_ref(wb, 'PLハイライト') -> 'PLハイライト!B4:F25'（A2に複数範囲がある場合は先頭）"""
    a2 = wb[key]['A2'].value or ''
    m_ = re.search(r'([A-Z]+\d+:[A-Z]+\d+)', a2)
    if not m_:
        raise KeyError(f'{key}!A2 に貼付範囲の表記がない')
    return f'{key}!{m_.group(1)}'


def remove_watermark(prs, marker='御中'):
    """テンプレートのスライドマスターにある宛先透かし（「株式会社●●御中」）を外す。鳥塚の概要書と同じ扱い。"""
    for m_ in prs.slide_masters:
        for sh in list(m_.shapes):
            if sh.has_text_frame and marker in sh.text_frame.text:
                sh._element.getparent().remove(sh._element)

"""openpyxlで作った案件マスターに、LibreOfficeで再計算した値（キャッシュ値）を書き戻す。

なぜ必要か：
  - openpyxlで保存したxlsxは数式のキャッシュ値を持たない → data_only=True で読むと全て None になり、
    概要書へ貼れない。
  - かといって LibreOffice で再計算して保存したファイルをそのまま使うと、游ゴシックが代替フォント
    （WenQuanYi等）に置き換わり、書式（＝貼付範囲の見た目）が崩れる。
  → 書式は元ファイル、値は再計算ファイルから取り、元ファイルの <c><f> に <v> を書き足す。

使い方::

    soffice --headless --norestore -env:UserInstallation=file:///tmp/lo-calc \
        --convert-to xlsx --outdir recalc 案件マスター_raw.xlsx
    python inject_cache.py 案件マスター_raw.xlsx recalc/案件マスター_raw.xlsx 案件マスター.xlsx

出力ファイルは書式・数式・キャッシュ値の全てを持つ（Excelで開いても再計算される）。
"""
import sys, zipfile, re
from lxml import etree
from openpyxl import load_workbook
src, calc, out = sys.argv[1:4]
vals = load_workbook(calc, data_only=True)
wbx = load_workbook(src)  # sheet order/names
NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
zin = zipfile.ZipFile(src)
wbxml = etree.fromstring(zin.read('xl/workbook.xml'))
rels = etree.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
rid2t = {r.get('Id'): r.get('Target') for r in rels}
path2name = {}
for s in wbxml.find(f'{{{NS}}}sheets'):
    t = rid2t[s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
    path2name['xl/' + t.lstrip('/').replace('xl/', '')] = s.get('name')
zout = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
n = 0
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename in path2name:
        ws = vals[path2name[item.filename]]
        root = etree.fromstring(data)
        for c in root.iter(f'{{{NS}}}c'):
            f = c.find(f'{{{NS}}}f')
            if f is None:
                continue
            v = ws[c.get('r')].value
            old = c.find(f'{{{NS}}}v')
            if old is not None:
                c.remove(old)
            ve = etree.SubElement(c, f'{{{NS}}}v')
            if isinstance(v, bool):
                c.set('t', 'b'); ve.text = '1' if v else '0'
            elif isinstance(v, (int, float)):
                c.attrib.pop('t', None); ve.text = repr(float(v)) if isinstance(v, float) else str(v)
            else:
                c.set('t', 'str'); ve.text = '' if v is None else str(v)
            n += 1
        data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
    zout.writestr(item, data)
zout.close()
print('injected', n)

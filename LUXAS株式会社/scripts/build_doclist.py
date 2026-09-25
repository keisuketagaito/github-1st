# -*- coding: utf-8 -*-
"""必要資料一覧（LUXAS株式会社／売主様ご提出用）を当社テンプレートから作成する。
   テンプレート：85fc53f6（当社標準の必要資料一覧）
   ブロック構成 Ⅰ3件/Ⅱ5件/Ⅲ5件/Ⅳ4件/Ⅴ2件（計19行）をそのまま使う。
   株式価値評価書（P.12・P.24）で「要確認」とした論点を解消するための資料を含める。"""
import openpyxl
from openpyxl.styles import Alignment

SRC = '/root/.claude/uploads/fe5370cf-c5f7-5843-870c-82796c774b08/85fc53f6-____________________.xlsx'
OUT = '/home/user/github-1st/LUXAS株式会社/必要資料一覧_LUXAS株式会社_20260923.xlsx'

wb = openpyxl.load_workbook(SRC)
ws = wb['Sheet1']

ws['B1'] = '必要資料一覧　　LUXAS株式会社 御中'
ws['B2'] = ('いずれもコピーまたはPDF等のデータにて、ご用意ください。コピーやPDF化が困難な場合は、原本にてご準備いただければ弊社で対応致しますのでご相談ください。\n'
            '※企業概要書の更新と買手候補へのご提示に必要な最低限の資料に絞っています。'
            '※2026年9月23日時点。「受領」欄に✔のある資料は受領済です。')

ITEMS = [
 # (行, カテゴリ記号, カテゴリ名, 資料名, 受領済)
 (3,  'Ⅰ', '概要', '定款（最新のもの）', False),
 (4,  None, None,  '株主名簿（株主名・持株数・持株比率がわかるもの）', False),
 (5,  None, None,  '駐車場用地（89.4㎡）の固定資産税課税明細書または不動産登記簿謄本　※土地の時価評価に使用します', False),
 (6,  'Ⅱ', '財務', '決算書・勘定科目内訳明細書・法人税／消費税申告書（直近3期分）', True),
 (7,  None, None,  '月次試算表（2026年7月分以降）　※2026年6月分までは受領済', False),
 (8,  None, None,  '借入金の返済予定表（金融機関別・直近のもの）', False),
 (9,  None, None,  '減価償却資産台帳（2025年11月期末時点）　※償却の過不足額の確認に使用します', False),
 (10, None, None,  '差入保証金（27,371千円）・長期前払費用（1,712千円）の内訳がわかる資料', False),
 (11, 'Ⅲ', '事業', '店舗別の売上・粗利がわかる資料（2025年12月〜2026年6月）', True),
 (12, None, None,  '撤退2店舗（アクアウォーク大垣店・ペットプラザ岐阜店）の原状回復費用・違約金がわかる資料', False),
 (13, None, None,  '棚卸資産の内訳（2025年11月期末）　※2026年3月計上の廃棄損44,000千円の処分方法・証憑もあわせてご教示ください', True),
 (14, None, None,  'ペット保険の代理店手数料がわかる資料（直近1期分の精算書等）', False),
 (15, None, None,  '事業用車両3台の車検証または査定書　※車両運搬具の時価評価に使用します', False),
 (16, 'Ⅳ', '人事', '従業員名簿（氏名・生年月日・入社年月日・雇用形態・勤務店舗がわかるもの）', True),
 (17, None, None,  '給与台帳（直近1期分｜直近の年末調整の資料等）', True),
 (18, None, None,  '退職金規程・役員退職慰労金規程　※制度がない場合はその旨をご教示ください', False),
 (19, None, None,  '賞与の支給実績・支給対象期間がわかる資料　※制度がない場合はその旨をご教示ください', False),
 (20, 'Ⅴ', '契約', 'フランチャイズ契約書（有限会社ワンラブ）', False),
 (21, None, None,  '店舗の賃貸借契約書（稼働6店舗分。FC本部からの転借契約を含む）　※敷金の償却条項・返還条件の確認に使用します', False),
]
for r, mark, cat, name, got in ITEMS:
    ws.cell(row=r, column=2).value = mark          # B列：ローマ数字
    ws.cell(row=r, column=3).value = cat           # C列：カテゴリ
    ws.cell(row=r, column=5).value = name          # E列：資料名
    # F列：チェックボックス機能は openpyxl で保持できないため「✔／空欄」に置き換える
    ws.cell(row=r, column=6).value = '✔' if got else None
    ws.cell(row=r, column=6).alignment = Alignment(horizontal='center', vertical='center')

# テンプレートに残っていた不規則な塗りを一旦すべて消し、受領済の行だけ淡いグレーを掛ける
from openpyxl.styles import PatternFill
NOFILL = PatternFill(fill_type=None)
GOTFILL = PatternFill('solid', fgColor='F2F2F2')
for r, _m, _c, _n, got in ITEMS:
    for col in (2, 3, 4, 5, 6):
        ws.cell(row=r, column=col).fill = GOTFILL if got else NOFILL

last = ITEMS[-1][0]
ws.print_area = f'B1:F{last}'
ws.row_dimensions[2].height = 51           # 前書き3行分
for _r, _i in ((r, n) for r, _m, _c, n, _g in ITEMS):
    ws.row_dimensions[_r].height = 30.75 if len(_i) > 38 else 20.5
wb.save(OUT)
print('saved:', OUT)

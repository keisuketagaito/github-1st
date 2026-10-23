# -*- coding: utf-8 -*-
"""案件マスターの体裁を整える。

1. 出所注記に残っていた「他社仲介」の表記を「対象会社提供の既存会社説明資料」へ置換する。
   概要書側は置換済みだったが、マスターが元の表記のままだと貼り直した瞬間に
   買手向け資料へ戻ってしまう。『その他』シートの社内確認事項（プロセス関連）は
   隔離しておくべき情報なのでそのまま残す。
2. 空の雛形行に振られた№の連番が崩れている箇所を直す。

冪等。
"""
import openpyxl

MASTER = '/home/user/github-1st/LUXAS株式会社/案件マスター_LUXAS株式会社_20260923.xlsx'
KEEP = {('その他', 'D22'), ('その他', 'E22')}      # 社内確認事項として隔離している行
REPL = [('他社仲介作成の企業概要書', '対象会社提供の既存会社説明資料'),
        ('他社仲介資料', '対象会社提供の既存会社説明資料'),
        ('他社仲介の企業概要書', '対象会社提供の既存会社説明資料')]
SEQ = [('株主', 5, 14, 2), ('取締役', 5, 14, 2), ('仕入先', 5, 14, 2),
       ('取引先', 5, 14, 2), ('外注先', 5, 14, 2)]


def main():
    wb = openpyxl.load_workbook(MASTER)
    n = 0
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str) or '他社仲介' not in cell.value:
                    continue
                if (ws.title, cell.coordinate) in KEEP:
                    continue
                new = cell.value
                for a, b in REPL:
                    new = new.replace(a, b)
                if new != cell.value:
                    cell.value = new
                    n += 1
    m = 0
    for name, r0, r1, col in SEQ:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        for i, r in enumerate(range(r0, r1 + 1), start=1):
            if ws.cell(row=r, column=col).value != i:
                ws.cell(row=r, column=col).value = i
                m += 1
    wb.save(MASTER)
    print(f'出所注記の置換 {n} 件／№の連番是正 {m} 件')


if __name__ == '__main__':
    main()

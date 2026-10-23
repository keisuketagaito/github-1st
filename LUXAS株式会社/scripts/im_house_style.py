# -*- coding: utf-8 -*-
"""企業概要書を当社の標準フォーマット（他案件の概要書に準拠）へ揃える。

当社の他案件（そば店／北関東・2026年10月版）の概要書と突き合わせ、
案件マスターからの抜粋方法がずれている箇所を是正する。

1. P.22 株主及び役員
   マスターの『株主』『取締役』シートは列構成が異なる別表であり、
   株主表＝№／氏名／役職／保有株式数／保有割合／備考（＋合計行）、
   役員表＝№／氏名／役職／代表者との関係／役員報酬／業務範囲／備考。
   これを1枚の表に押し込み「保有株式数／保有割合」を1セルへ統合し、
   「業務範囲」列と合計行を落としていたため、シートどおりの別表へ作り直す。
2. P.18 主要仕入先
   マスターの『仕入先』シートは各期に構成比を持つ。最終期だけを出していたため全期に出す。
3. P.31／P.33／P.34（貸借対照表・損益計算書・販管費）
   標準フォーマットでは財務情報詳細の各ページにCommentを置く。欠けていたので追加する。
4. 太字の是正（P.7の長期未払金は明細行であり太字にしない／P.34の合計行は太字にする）

冪等。適用済みであれば各処理をスキップする。
"""
import copy
import sys

import openpyxl
from pptx import Presentation
from pptx.util import Inches, Pt

sys.path.insert(0, '/home/user/github-1st/LUXAS株式会社/scripts')
from pptlib import (set_text, set_lines, fill_row, fit_rows, fit_cols, restyle_rows,
                    unmerge_v, set_widths, set_heights, table_font, align_cells,
                    font_size, place, clone_shape, num, pct)

IM     = '/home/user/github-1st/LUXAS株式会社/企業概要書_LUXAS株式会社_20260923.pptx'
MASTER = '/home/user/github-1st/LUXAS株式会社/案件マスター_LUXAS株式会社_20260923.xlsx'

P_SHARE, P_SUPP, P_BS, P_PL, P_SGA = 22, 18, 31, 33, 34
CHIP_SRC, COMMENT_SRC = 26, 7        # 見出しチップ／Comment枠の複製元ページ


def K(v):
    return None if v in (None, '') else int(round(v / 1000.0))


def pad(table, left=0.07, right=0.07):
    """セルの左右余白を揃える。詰まっていると隣の列の文字と地続きに見える。"""
    for r in range(len(table.rows)):
        for c in range(len(table.columns)):
            cell = table.cell(r, c)
            cell.margin_left = Inches(left)
            cell.margin_right = Inches(right)


def bold_row(table, r, on=True):
    for c in range(len(table.columns)):
        for para in table.cell(r, c).text_frame.paragraphs:
            for run in para.runs:
                run.font.bold = on


def tables(slide):
    return [sh for sh in slide.shapes if sh.has_table]


def clone(slide, shape):
    el = copy.deepcopy(shape._element)
    slide.shapes._spTree.append(el)
    return list(slide.shapes)[-1]


# ============================================================ 1. 株主及び役員
def fix_shareholders(prs, wb):
    s = prs.slides[P_SHARE - 1]
    if len(tables(s)) > 1:
        print('P.%d 株主及び役員：適用済み' % P_SHARE)
        return
    sh_ws, dir_ws, emp_ws = wb['株主'], wb['取締役'], wb['従業員']

    g = lambda ws, r, c: ws.cell(row=r, column=c).value
    # 株主（マスター『株主』シート）
    # B=№ C=氏名 D=役職 E=役員報酬 F=保有株式数 G=保有割合 H=備考
    shareholders = [(str(g(sh_ws, r, 2)), g(sh_ws, r, 3), g(sh_ws, r, 4),
                     g(sh_ws, r, 6), g(sh_ws, r, 7), g(sh_ws, r, 8))
                    for r in range(5, 15) if g(sh_ws, r, 3)]
    sh_total = (g(sh_ws, 15, 6), g(sh_ws, 15, 7))
    sh_note = g(sh_ws, 17, 2)
    # 役員（マスター『取締役』シート）
    officers = [(str(g(dir_ws, r, 2)), g(dir_ws, r, 3), g(dir_ws, r, 4), g(dir_ws, r, 5),
                 K(g(dir_ws, r, 6)), g(dir_ws, r, 7), g(dir_ws, r, 8))
                for r in range(5, 15) if g(dir_ws, r, 3)]
    dir_note = g(dir_ws, 17, 2)
    # キーパーソン（マスター『従業員』シートの役員以外のキーパーソン）
    key = [(str(g(emp_ws, r, 2)), g(emp_ws, r, 3), g(emp_ws, r, 9), g(emp_ws, r, 10))
           for r in range(5, 8) if g(emp_ws, r, 9) == '現場責任者']

    base = tables(s)[0]
    tA = base
    tB = clone(s, base)
    tC = clone(s, base)
    chip = clone(s, list(prs.slides[CHIP_SRC - 1].shapes)[1])
    chip2 = clone(s, chip)
    chip3 = clone(s, chip)

    def block(chipshape, label, top):
        set_text(chipshape, label)
        place(chipshape, left=0.34, top=top, width=1.90, height=0.34)
        font_size(chipshape, 11)

    def build(shape, head, rows, widths, hrow=0.34, drow=0.52):
        t = shape.table
        fit_rows(t, len(rows) + 1, 1)
        restyle_rows(t, {i: 1 for i in range(1, len(rows) + 1)})
        fit_cols(t, len(head), 4)
        unmerge_v(t)
        fill_row(t, 0, head)
        for i, row in enumerate(rows, start=1):
            fill_row(t, i, row)
        set_widths(t, widths)
        set_heights(t, [hrow] + [drow] * len(rows))
        table_font(t, 9)
        pad(t)
        align_cells(t, tuple(range(len(head))), 'c', rows=(0,))
        align_cells(t, (0,), 'c', rows=range(1, len(rows) + 1))

    # --- 株主
    block(chip, '株主', 1.52)
    rows = [(no, nm, pos, '{:,}株'.format(int(n)), pct(ratio), note or '－')
            for no, nm, pos, n, ratio, note in shareholders]
    rows.append(('', '合計', '－', '{:,}株'.format(int(sh_total[0])), pct(sh_total[1]), ''))
    build(tA, ('№', '氏名', '役職', '保有株式数', '保有割合', '備考'), rows,
          [0.45, 1.90, 1.40, 1.20, 1.10, 4.91], drow=0.38)
    place(tA, left=0.34, top=1.95, width=10.96)
    align_cells(tA.table, (3, 4), 'r', rows=range(1, len(rows) + 1))
    bold_row(tA.table, len(rows))

    # --- 役員
    block(chip2, '役員', 3.30)
    rows = [(no, nm, pos, rel, '{:,}千円'.format(pay) if pay else '－', scope or '－', note or '－')
            for no, nm, pos, rel, pay, scope, note in officers]
    build(tB, ('№', '氏名', '役職', '代表者との関係', '役員報酬\n（25年11月期）',
               '業務範囲', '備考'), rows,
          [0.45, 1.70, 1.25, 1.20, 1.25, 2.50, 2.61], drow=0.44)
    place(tB, left=0.34, top=3.73, width=10.96)
    align_cells(tB.table, (4,), 'r', rows=range(1, len(rows) + 1))

    # --- キーパーソン
    block(chip3, 'キーパーソン', 4.85)
    rows = [(no, nm, '現場責任者', '配偶者', '7,000千円', note or '－')
            for no, nm, _pos, note in key]
    if not rows:
        rows = [('－',) * 6]
    build(tC, ('№', '氏名', '役職', '代表者との関係', '年間給与\n（25年11月期）', '業務範囲・備考'),
          rows, [0.45, 1.70, 1.25, 1.20, 1.25, 5.11], drow=0.44)
    place(tC, left=0.34, top=5.28, width=10.96)
    align_cells(tC.table, (4,), 'r', rows=range(1, len(rows) + 1))

    note = clone(s, [sh for sh in prs.slides[P_SUPP - 1].shapes
                     if sh.has_text_frame and sh.text_frame.text.startswith('（単位')][0])
    place(note, left=0.34, top=6.25, width=10.96, height=1.10)
    set_lines(note, [sh_note, dir_note.split('\n')[0], dir_note.split('\n')[1]])
    font_size(note, 8.5)
    print('P.%d 株主及び役員：株主／役員／キーパーソンの3表へ作り直し' % P_SHARE)


# ============================================================ 2. 主要仕入先
def fix_suppliers(prs, wb):
    s = prs.slides[P_SUPP - 1]
    t = tables(s)[0].table
    if len(t.columns) >= 10:
        print('P.%d 主要仕入先：適用済み' % P_SUPP)
        return
    ws = wb['仕入先']
    g = lambda r, c: ws.cell(row=r, column=c).value
    body = []
    for r in list(range(5, 15)) + [15]:
        if not g(r, 3):
            continue
        body.append((str(g(r, 2) or ''), g(r, 3), g(r, 5) or '－', g(r, 4) or '－',
                     K(g(r, 6)), g(r, 7), K(g(r, 8)), g(r, 9), K(g(r, 10)), g(r, 11)))
    tot = (K(g(16, 6)), K(g(16, 8)), K(g(16, 10)))

    head = ('№', '仕入先', '所在地', '仕入内容',
            '23年11月期', '構成比', '24年11月期', '構成比', '25年11月期', '構成比')
    rows = [head]
    for no, nm, loc, cont, v1, p1, v2, p2, v3, p3 in body:
        # 仕入高が無い期は構成比も「－」にする（0.0%と出すと取引があったように見える）
        rows.append((no, nm, loc, cont,
                     num(v1), pct(p1) if v1 else '－',
                     num(v2), pct(p2) if v2 else '－',
                     num(v3), pct(p3) if v3 else '－'))
    rows.append(('', '合計', '－', '－', num(tot[0]), '100.0%', num(tot[1]), '100.0%',
                 num(tot[2]), '100.0%'))

    fit_cols(t, len(head), 3)
    fit_rows(t, len(rows), 1)
    restyle_rows(t, {i: 1 for i in range(1, len(rows))})
    unmerge_v(t)
    for i, row in enumerate(rows):
        fill_row(t, i, row)
    set_widths(t, [0.40, 2.00, 1.15, 2.25] + [0.90] * 6)
    set_heights(t, [0.40] + [0.33] * (len(rows) - 1))
    table_font(t, 8)
    pad(t, left=0.06, right=0.06)
    bold_row(t, len(rows) - 1)
    align_cells(t, tuple(range(10)), 'c', rows=(0,))
    align_cells(t, (0,), 'c', rows=range(1, len(rows)))
    align_cells(t, (1, 2, 3), 'l', rows=range(1, len(rows)))
    align_cells(t, tuple(range(4, 10)), 'r', rows=range(1, len(rows)))
    pt = [sh for sh in s.shapes if sh.has_text_frame
          and sh.text_frame.text.startswith('■ ジャペル')]
    if pt:
        set_lines(pt[0], [
         '■ ジャペル㈱（肥料・飼料卸売業）・㈲ワンラブ（ペット・ペット用品小売業／FC本部）の上位2社で'
         '仕入高の約75%。いずれも創業来の継続取引であり、供給の安定性は高い。',
         '■ デザイナーズワン㈱（ペット生育業）は代表者の親族が経営する会社。当社は同社の営業権'
         '（簿価1,500千円）を保有し、松阪管理センターで同社の生体を管理している'
         '（大垣管理センターは2026年2月に閉鎖）。取引条件の妥当性は要確認。',
         '■ 生体はFC本部がオークション会場で仕入れ、マイクロチップ埋込・ワクチン接種等の処理を'
         '済ませたうえで当社へ供給される。',
        ])
        font_size(pt[0], 9)
    print('P.%d 主要仕入先：構成比を3期とも表示（マスター『仕入先』シートどおり）' % P_SUPP)


# ============================================================ 3. 財務情報詳細のComment
BS_COMMENT = [
 '≪商品≫',
 '・70,918→127,081→106,017千円。2024年11月期に大型犬専門店2店舗の出店へ向けて在庫を積み増した後、'
 '2025年11月期は減少。総資産の45%を占める（内訳はP.32参照）。',
 '≪差入保証金≫',
 '・20,642→27,656→27,371千円。店舗の敷金・保証金。償却条項・契約解除時の返還条件は'
 '賃貸借契約書が未受領のため要確認。',
 '≪長期借入金≫',
 '・134,042→218,455→191,858千円。2024年11月期に大型犬専門店2店舗の出店等で84,413千円増加。'
 '民間6本は代表者の連帯保証付き。',
 '≪役員借入金≫',
 '・2025年11月期末7,566千円。譲渡時の返済を希望（進行期は2026年6月末で3,403千円へ減少）。',
 '≪純資産≫',
 '・5,327→6,302→△14,057千円。2025年11月期の当期純損失20,359千円により債務超過に転じている。',
]
PL_COMMENT = [
 '≪売上高≫',
 '・474,428→539,035→558,390千円と2期で17.7%成長。生体売上が6割を占める構成は3期を通じて安定。',
 '≪売上原価≫',
 '・粗利率は61.6%→63.6%→59.1%。2025年11月期は小型犬中心の既存店で仕入価格の上昇を'
 '販売価格へ転嫁できなかったことが主因（ヒアリング）。',
 '≪販売費及び一般管理費≫',
 '・290,037→339,909→349,576千円と2期で59,539千円増加。路面店2店舗の通期稼働により'
 '地代家賃が+25,019千円、人件費が+25,612千円（内訳はP.34参照）。',
 '≪特別損失≫',
 '・2025年11月期の1,298千円は固定資産売却損968千円・固定資産除却損330千円であり、いずれも非経常。',
]
SGA_COMMENT = [
 '≪人件費≫',
 '・役員報酬・従業員給与・法定福利費・厚生費の計は134,963→151,750→160,575千円、'
 '売上高比28.4%→28.2%→28.8%とほぼ一定。',
 '≪ロイヤリティ≫',
 '・45,990→39,848→29,982千円、売上高比9.7%→7.4%→5.4%と低下。'
 'ロイヤリティを負担しないブランド店の売上構成比が高まった裏返しである（P.14参照）。',
 '≪地代家賃≫',
 '・8,230→19,460→33,249千円。路面店2店舗の通期稼働による増加（内訳はP.27参照）。',
 '≪租税公課≫',
 '・8,881→10,506→20,061千円。2025年11月期の増加要因は未確認（要確認）。',
]


def add_comment(prs, page, lines, table_right=7.45, widths=None, layout=None):
    s = prs.slides[page - 1]
    if any(sh.has_text_frame and sh.text_frame.text.strip() == 'Comment' for sh in s.shapes):
        print('P.%d Comment：適用済み' % page)
        return
    src = prs.slides[COMMENT_SRC - 1]
    frame = clone(s, list(src.shapes)[3])
    body = clone(s, list(src.shapes)[5])
    place(frame, left=table_right + 0.15, top=1.52, width=11.20 - (table_right + 0.15), height=5.93)
    place(body, left=table_right + 0.30, top=2.20, width=11.05 - (table_right + 0.30), height=5.18)
    set_lines(body, [('head', l) if l.startswith('≪') else ('body', l) for l in lines])
    font_size(body, 8.5)
    if widths:
        for tshape, w in zip(tables(s), widths):
            set_widths(tshape.table, w)
            tshape.width = Inches(sum(w))
    if layout:
        layout(tables(s))
    print('P.%d にCommentを追加' % page)


# ============================================================ 4. 太字の是正
def fix_bold(prs):
    s = prs.slides[7 - 1]
    t = tables(s)[0].table
    n = 0
    for r in range(len(t.rows)):
        label = t.cell(r, 4).text.strip()
        if label in ('長期未払金', '　', ''):
            for c in range(4, 7):
                for para in t.cell(r, c).text_frame.paragraphs:
                    for run in para.runs:
                        if run.font.bold:
                            run.font.bold = False
                            n += 1
    s = prs.slides[P_SGA - 1]
    t = tables(s)[0].table
    last = len(t.rows) - 1
    for c in range(len(t.columns)):
        for para in t.cell(last, c).text_frame.paragraphs:
            for run in para.runs:
                run.font.bold = True
    print('太字を是正：P.7 明細行の解除 %d 箇所／P.%d 合計行を太字化' % (n, P_SGA))


def main():
    wb = openpyxl.load_workbook(MASTER, data_only=True)
    prs = Presentation(IM)
    fix_shareholders(prs, wb)
    fix_suppliers(prs, wb)
    def bs_layout(ts):
        """資産の部（左）／負債の部・純資産の部（右）を左へ寄せ直す。
           期のヘッダー（23年11月期）が2行に折り返さないよう見出しだけ小さくする。"""
        ts[0].left = Inches(0.37)
        for t in ts[1:]:
            t.left = Inches(4.08)
        for sh in ts[:2]:            # 純資産の部の表は期のヘッダー行を持たない
            table_font(sh.table, 7.5, rows=(1,))
    add_comment(prs, P_BS, BS_COMMENT, table_right=7.79,
                widths=[[1.28, 0.80, 0.80, 0.80]] * 3, layout=bs_layout)
    add_comment(prs, P_PL, PL_COMMENT, table_right=7.45,
                widths=[[1.80, 0.88, 0.88, 0.88, 0.88, 0.88, 0.88]])
    add_comment(prs, P_SGA, SGA_COMMENT, table_right=7.45,
                widths=[[1.95, 0.86, 0.86, 0.86, 0.86, 0.86, 0.86]])
    fix_bold(prs)
    prs.save(IM)
    print('saved:', IM)


if __name__ == '__main__':
    main()

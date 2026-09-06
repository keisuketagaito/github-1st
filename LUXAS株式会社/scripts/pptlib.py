# -*- coding: utf-8 -*-
"""企業概要書のテンプレート（前回案件）を壊さずに編集するためのユーティリティ。"""
from copy import deepcopy
from pptx.oxml.ns import qn
from pptx.util import Inches, Emu

BLANK = '　'

# ---------------------------------------------------------------- テキスト
def _runs(p_el):
    return p_el.findall(qn('a:r'))

def _set_p_text(p_el, text):
    """段落の書式を保ったまま本文だけ差し替える（ランは作り直さない）。"""
    runs = _runs(p_el)
    for br in p_el.findall(qn('a:br')):
        p_el.remove(br)
    if not runs:
        return False
    t = runs[0].find(qn('a:t'))
    if t is None:
        return False
    t.text = text
    for r in runs[1:]:
        p_el.remove(r)
    return True

def set_text(shape, text):
    """単一段落のシェイプ／プレースホルダーのテキストを差し替える。"""
    tf = shape.text_frame
    ps = list(tf.paragraphs)
    for p in ps[1:]:
        p._p.getparent().remove(p._p)
    if not _set_p_text(ps[0]._p, text):
        tf.text = text

def _templates(tf):
    """見出し（≪ 【 で始まる）用と本文用の段落テンプレートを拾う。"""
    head = body = None
    for p in tf.paragraphs:
        t = p.text.strip()
        if not t or not _runs(p._p):
            continue
        if t[0] in '≪【' and head is None:
            head = p._p
        elif body is None:
            body = p._p
    if head is None:
        head = body
    if body is None:
        body = head
    return head, body

def set_lines(shape, lines):
    """複数段落のテキストボックスを差し替える。
       lines は文字列、または ('head'|'body', 文字列) のタプル。"""
    tf = shape.text_frame
    head, body = _templates(tf)
    if head is None:
        set_text(shape, '\n'.join(l if isinstance(l, str) else l[1] for l in lines))
        return
    head, body = deepcopy(head), deepcopy(body)
    txBody = tf._txBody
    for p in txBody.findall(qn('a:p')):
        txBody.remove(p)
    for item in lines:
        role, text = ('body', item) if isinstance(item, str) else item
        if role == 'auto':
            role = 'head' if (text[:1] in '≪【') else 'body'
        new = deepcopy(head if role == 'head' else body)
        _set_p_text(new, text)
        txBody.append(new)

# ---------------------------------------------------------------- 表
def _has_run(cell):
    return any(_runs(p._p) for p in cell.text_frame.paragraphs)

def set_cell(cell, text, donor=None):
    """空セル（ランを持たない）は既定18ptで描画されるため、
       同じ列の書式を持つセルから段落ごと複製してから差し替える。"""
    tf = cell.text_frame
    if not _has_run(cell) and donor is not None:
        dps = [p for p in donor.text_frame.paragraphs if _runs(p._p)]
        if dps:
            txBody = tf._txBody
            for p in txBody.findall(qn('a:p')):
                txBody.remove(p)
            txBody.append(deepcopy(dps[0]._p))
    ps = list(tf.paragraphs)
    for p in ps[1:]:
        p._p.getparent().remove(p._p)
    txt = text if text not in ('', None) else BLANK
    if not _set_p_text(ps[0]._p, str(txt)):
        tf.text = str(txt)

def _donor(table, r, c):
    n = len(table.rows)
    for d in range(1, n):
        for rr in (r + d, r - d):
            if 0 <= rr < n:
                cand = table.cell(rr, c)
                if _has_run(cand):
                    return cand
    return None

def fill_row(table, r, values):
    for c, v in enumerate(values):
        if c < len(table.columns):
            set_cell(table.cell(r, c), v, _donor(table, r, c))

def trs(table):
    return table._tbl.findall(qn('a:tr'))

def _clean_tr(tr):
    for tc in tr.findall(qn('a:tc')):
        for a in ('rowSpan', 'vMerge'):
            tc.attrib.pop(a, None)
    for h in list(tr.iter(qn('a:highlight'))):
        h.getparent().remove(h)

def add_rows(table, template_idx, n, after_idx=None):
    """template_idx 行を複製して after_idx の直後に n 行追加する。"""
    rows = trs(table)
    src = rows[template_idx]
    ref = rows[template_idx if after_idx is None else after_idx]
    for _ in range(n):
        new = deepcopy(src)
        _clean_tr(new)
        ref.addnext(new)
        ref = new

def restyle_rows(table, spec):
    """spec: {行番号: 書式を借りる行番号}。<a:tr> ごと差し替えて行スタイルを揃える。"""
    rows = trs(table)
    src = {k: deepcopy(rows[v]) for k, v in spec.items()}
    for k, new in src.items():
        _clean_tr(new)
        rows[k].addprevious(new)
        rows[k].getparent().remove(rows[k])

def align_cells(table, cols, how, rows=None):
    from pptx.enum.text import PP_ALIGN
    m = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}
    rng = rows if rows is not None else range(len(table.rows))
    for r in rng:
        for c in cols:
            for para in table.cell(r, c).text_frame.paragraphs:
                para.alignment = m[how]

def font_size(shape, pt):
    from pptx.util import Pt as _Pt
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = _Pt(pt)

def del_rows(table, idxs):
    rows = trs(table)
    for i in sorted(set(idxs), reverse=True):
        rows[i].getparent().remove(rows[i])

def fit_rows(table, want, template_idx, tail_keep=0):
    """行数を want に合わせる。tail_keep は末尾に残す行数。"""
    cur = len(trs(table))
    if cur < want:
        after = cur - 1 - tail_keep
        add_rows(table, template_idx, want - cur, after_idx=after)
    elif cur > want:
        start = cur - tail_keep - (cur - want)
        del_rows(table, range(start, start + (cur - want)))

def set_heights(table, heights):
    for i, h in enumerate(heights):
        if i < len(table.rows) and h is not None:
            table.rows[i].height = Inches(h)

def set_widths(table, widths):
    for i, w in enumerate(widths):
        if i < len(table.columns) and w is not None:
            table.columns[i].width = Inches(w)

# ---------------------------------------------------------------- シェイプ
def drop(slide, *idxs):
    shapes = list(slide.shapes)
    for i in sorted(set(idxs), reverse=True):
        el = shapes[i]._element
        el.getparent().remove(el)

def place(shape, left=None, top=None, width=None, height=None):
    if left is not None:   shape.left = Inches(left)
    if top is not None:    shape.top = Inches(top)
    if width is not None:  shape.width = Inches(width)
    if height is not None: shape.height = Inches(height)

# ---------------------------------------------------------------- スライド
def prune_and_order(prs, keep):
    """keep（1始まりの元スライド番号のリスト）だけを残し、その順に並べ替える。"""
    lst = prs.slides._sldIdLst
    ids = list(lst)
    keep_els = [ids[i - 1] for i in keep]
    for el in ids:
        if el not in keep_els:
            prs.part.drop_rel(el.get(qn('r:id')))
            lst.remove(el)
    for el in keep_els:
        lst.remove(el)
    for el in keep_els:
        lst.append(el)

def num(v, dash='－'):
    """千円単位の数値を表示用文字列にする（負値は△、0/None は dash）。"""
    if v is None:
        return dash
    if v == 0:
        return dash
    return ('△ ' if v < 0 else '') + '{:,}'.format(abs(int(round(v))))

def pct(v, dash='－'):
    if v is None:
        return dash
    return '{:.1f}%'.format(v * 100)

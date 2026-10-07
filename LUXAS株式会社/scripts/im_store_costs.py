# -*- coding: utf-8 -*-
"""企業概要書に『店舗別の費用構造（進行期7か月）』ページを追加する（P.13の直後＝新P.14）。

数値はすべて案件マスターの『店舗別費用』シート由来（master_store_costs.py → recalc 済み
であること）。ページ追加に伴いP.14以降が1つずつ繰り下がるため、デッキ内の「P.NN参照」
のうち14以上のものを+1する。

冪等性：同名ページが既にあれば何もせずに終了する。作り直す場合は概要書をgitから戻すこと。
"""
import re
import sys

import openpyxl
from pptx import Presentation
from pptx.util import Inches, Pt

sys.path.insert(0, '/home/user/github-1st/LUXAS株式会社/scripts')
from pptlib import (set_text, set_lines, fill_row, fit_rows, fit_cols, restyle_rows,
                    unmerge_v, set_widths, set_heights, table_font, align_cells,
                    font_size, place, dup_slide, num, pct)

IM     = '/home/user/github-1st/LUXAS株式会社/企業概要書_LUXAS株式会社_20260923.pptx'
MASTER = '/home/user/github-1st/LUXAS株式会社/案件マスター_LUXAS株式会社_20260923.xlsx'
TITLE  = '店舗別の費用構造（進行期7か月）'
SRC    = 13          # 複製元：店舗別損益ページ
AT     = 14          # 挿入後のページ番号

HEAD = ('部門', '売上高', '従業員給与', '対売上', '法定福利費', 'ロイヤリティ',
        '地代家賃', '広告宣伝費', 'その他', '販管費計', '対売上')


def K(v):
    """円→千円（四捨五入）。"""
    return None if v in (None, '') else int(round(v / 1000.0))


def load_rows():
    """マスターの『店舗別費用』から行を読む。
       スライドの「その他」は マスターのその他＋役員報酬＋支払手数料。"""
    wb = openpyxl.load_workbook(MASTER, data_only=True)
    ws = wb['店舗別費用']
    out = []
    r = 5
    while ws.cell(row=r, column=3).value:
        g = lambda c: ws.cell(row=r, column=c).value
        name, state = g(3), g(4)
        sales, pay, ins = K(g(5)), K(g(6)), K(g(7))
        exe, roy, rent, ad, fee, oth, sga = (K(g(10)), K(g(11)), K(g(12)),
                                             K(g(13)), K(g(14)), K(g(15)), K(g(16)))
        label = name
        if state and state not in ('－', '稼働中'):
            label += '（%s）' % state.replace('に撤退', '撤退').replace('に閉鎖', '閉鎖')
        # 「その他」は役員報酬・支払手数料を含む差額。千円に丸めたあとの差額とすることで
        # スライド上の各行が必ず販管費計に合う。
        other = (sga or 0) - sum(x or 0 for x in (pay, ins, roy, rent, ad))
        out.append(dict(label=label, bold=name.startswith('【'), sales=sales, pay=pay,
                        ins=ins, roy=roy, rent=rent, ad=ad, other=other, sga=sga,
                        pay_r=(pay / sales if sales and pay else None),
                        sga_r=(sga / sales if sales and sga else None)))
        r += 1
    return out


def main():
    prs = Presentation(IM)
    for s in prs.slides:
        if s.shapes[0].has_text_frame and s.shapes[0].text_frame.text.strip() == TITLE:
            print('すでに存在するため何もしません:', TITLE)
            return

    rows = load_rows()
    s = dup_slide(prs, SRC)

    # --- 並べ替え（末尾 → AT 番目）
    lst = prs.slides._sldIdLst
    el = list(lst)[-1]
    lst.remove(el)
    lst.insert(AT - 1, el)

    set_text(s.shapes[0], TITLE)
    set_text(s.shapes[3],
             '販管費率は稼働6店舗で41.4%。明和店の75.6%は従業員給与11,071千円'
             '（売上高比44.1%）が主因であり、松阪管理センターの人件費が同店に'
             '計上されているとの推察と整合する。')

    t = s.shapes[2].table
    body = [HEAD]
    for d in rows:
        body.append((d['label'], num(d['sales']), num(d['pay']), pct(d['pay_r']),
                     num(d['ins']), num(d['roy']), num(d['rent']), num(d['ad']),
                     num(d['other']), num(d['sga']), pct(d['sga_r'])))
    fit_cols(t, len(HEAD), 1)
    fit_rows(t, len(body), 2)
    unmerge_v(t)
    for i, row in enumerate(body):
        fill_row(t, i, row)
    place(s.shapes[2], left=0.37, top=1.72, width=10.83)
    set_widths(t, [2.43] + [0.84] * 10)
    set_heights(t, [0.40] + [0.30] * (len(body) - 1))
    table_font(t, 8)
    align_cells(t, (0,), 'l', rows=range(1, len(body)))
    align_cells(t, tuple(range(1, 11)), 'r', rows=range(1, len(body)))
    align_cells(t, tuple(range(0, 11)), 'c', rows=(0,))

    set_lines(s.shapes[4], [
     '（単位：千円）※出所：店舗別（部門別）勘定科目残高推移表 2025年12月〜2026年6月'
     '（2026年9月18日出力）。従業員給与・法定福利費・ロイヤリティ・地代家賃・広告宣伝費の'
     '各科目について、11部門の合計が全社の試算表と一致することを確認済。',
     '※法定福利費は本社に一括計上したうえで各店舗へマイナス計上（振替）する処理がとられており、'
     '店舗側は各店ともマイナス残高となる。店舗別の販管費は社会保険料の負担分だけ過少に'
     '表示されている（稼働6店舗計△6,449千円）。',
     '※「その他」は役員報酬・支払手数料・旅費交通費・厚生費・リース料・通信費・水道光熱費・'
     '租税公課・接待交際費・保険料・備品消耗品費・管理諸費・雑費等の合計。',
     '※ロイヤリティはワンラブFC加盟店のうち明和・伊勢・久居（および撤退2店）が負担しており'
     '（稼働3店で売上高比12.0〜13.8%）、地代家賃の計上はない。一方ドン・キホーテ緑店は'
     'ロイヤリティの負担がなく地代家賃5,407千円（同14.6%）を負担しており、'
     '同店のみ契約形態が異なる理由は要確認。',
     '※広告宣伝費は定額計上（明和330千円／月、伊勢・久居220千円／月、緑110千円／月、'
     '路面店2店は280〜360千円／月）。',
    ])
    font_size(s.shapes[4], 8)

    # --- ページ番号参照の繰り下げ（14以上を+1）
    bumped = 0

    def bump(text):
        nonlocal bumped

        def rep(m):
            nonlocal bumped
            n = int(m.group(1))
            if n >= AT:
                bumped += 1
                return 'P.%d' % (n + 1)
            return m.group(0)
        return re.sub(r'P\.(\d+)', rep, text)

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:
                walk(sh.shapes)
                continue
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for run in para.runs:
                        if 'P.' in run.text:
                            run.text = bump(run.text)
            if sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        for para in cell.text_frame.paragraphs:
                            for run in para.runs:
                                if 'P.' in run.text:
                                    run.text = bump(run.text)

    for sl in prs.slides:
        walk(sl.shapes)

    prs.save(IM)
    print(f'saved: {IM} / slides: {len(prs.slides._sldIdLst)} / P.{AT} に追加 / '
          f'ページ参照の繰り下げ {bumped} 件')


if __name__ == '__main__':
    main()
